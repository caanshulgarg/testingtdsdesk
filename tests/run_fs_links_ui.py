"""python3 run_fs_links_ui.py - every figure linked to the trial balance (owner, 29-Oct-2026).
On the React build (app/dist-test, or TDSDESK_SITE), the made-up books of tests/fixtures/fin-statements:
  - every balance sheet and P&L figure, and every note line, opens the trial balance filtered to exactly its ledgers,
    and their total is the figure (the surplus with the year's profit);
  - the Excel (full year and H1, thousands): every number on the Balance Sheet, Profit-Loss and Notes sheets is a formula
    (no typed figure), and LibreOffice recalculates each to the screen's figure;
  - a phone (390 wide): a figure opens its ledgers, no sideways scroll.
Screenshots go to TDSDESK_SHOTS (default tests/out/links-shots)."""
import os, re, sys, json, base64, shutil, zipfile, subprocess, threading, functools, http.server
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
from playwright.sync_api import sync_playwright
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.environ.get("TDSDESK_SITE", os.path.join(HERE, "..", "app", "dist-test"))
OUT = os.path.join(os.environ.get("TDSDESK_OUT", os.path.join(HERE, "out")), "fs-links"); shutil.rmtree(OUT, ignore_errors=True); os.makedirs(OUT)
SHOTS = os.environ.get("TDSDESK_SHOTS", os.path.join(HERE, "out", "links-shots")); os.makedirs(SHOTS, exist_ok=True)
fails, errors = [], []
def ok(c, w):
    print(("  ok   " if c else "  FAIL ") + w)
    if not c: fails.append(w)
class Q(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(("localhost", 0), functools.partial(Q, directory=SITE)); PORT = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()
BOOKS = json.loads(subprocess.check_output(["node", "-e", r"""
const F = require('./fixtures/fin-statements/company.json'), from = '20240401', r2 = v => Math.round(v * 100) / 100;
const under = Object.assign({}, F.groups), led = {};
Object.entries(F.o).forEach(([n, v]) => { led[n] = {open: -v, close: 0, parent: under[n]}; });
const vouchers = [], bank = 'Northbank Current A/c', shift = e => String(Number(e.slice(0, 4)) - 1) + e.slice(4);
F.additions.forEach(([l, d, a], i) => vouchers.push({id: 'a' + i, date: d, no: 'A' + i, type: 'Journal', narr: 'Purchase of ' + l, ent: [{l, a: -a}, {l: bank, a}]}));
const W = [0.2, 0.3, 0.25], QE = ['0630', '0930', '1231', '0331'], parts = v => { const a = W.map(w => r2(v * w)); return a.concat([r2(v - a[0] - a[1] - a[2])]); };
[['20250331', F.p, F.plp, F.o], ['20260331', F.c, F.plc, F.p]].forEach(([end, bs, pl, prev]) => { const y = Number(end.slice(0, 4)) - 1, mv = {};
  Object.keys(bs).forEach(n => { if (n === bank) return; const added = F.additions.filter(([l, d]) => l === n && d <= end && d > shift(end)).reduce((s, z) => s + z[2], 0);
    const m = r2((bs[n] || 0) - (prev[n] || 0) - added); if (m) mv[n] = m; });
  Object.entries(pl).forEach(([n, v]) => { if (v) mv[n] = v; });
  [0, 1, 2, 3].forEach(q => { const ent = []; Object.entries(mv).forEach(([n, v]) => { const a = parts(v)[q]; if (a) ent.push({l: n, a: -a}); });
    ent.push({l: bank, a: r2(-ent.reduce((s, e) => s + e.a, 0))}); vouchers.push({id: 'm' + end + q, date: (q === 3 ? y + 1 : y) + QE[q], no: 'M' + end + q, type: 'Journal', narr: 'Movements (fixture) SECRET-NARRATION', ent}); }); });
vouchers.push({id: 'n1', date: '20260630', no: 'S1', type: 'Sales', narr: 'SECRET-NARRATION 2026', ent: [{l: 'Software Services - Export', a: 500000}, {l: 'Fresh Client Ltd', a: -500000}]});
under['Fresh Client Ltd'] = 'Sundry Debtors';
process.stdout.write(JSON.stringify({under, groups: {}, vouchers, meta: {from, to: '20270331'}, tb: {from, to: '20270331', at: '2027-04-30T10:00:00Z', src: 'copy', led}}));
"""], cwd=HERE))
SETUP = """(bk) => {
  S.firm = Object.assign({}, S.firm, {firmName: "Made-up & Associates", frn: "000000S", partner: "A. Sample", mno: "000001", place: "Bengaluru"});
  const c = newCompany({name: "Larkspur Analytics Private Limited (made-up)", gstin: "29AABCL1234F1Z5"}); c.pan = "AABCL1234F"; c.cin = "U72900KA2015PTC000001";
  S.companies[c.id] = c; Store.saveCompany(c); S.coId = c.id; S.view = "company"; S.tab = "books"; S.loadingCo = false;
  const b = Object.assign({loading: false, cid: c.id, challans: [], alloc: {}, ledInfo: {}, gstins: {}, pans: {}, states: {}, misCfg: {freq: "off"}, auditCfg: {freq: "off"},
    msme: {"Bluebell Print Works": "Micro"}, fs: {taxDep: {"2024": {c40: 220000, pm15: 140000, f10: 80000}}, inputs: {unit: 1, dec: 2}}}, bk);
  b.map = Books.mapLedgers(b.vouchers, {});
  S.books = b; window.__bk = b; S.fsFy = "2025"; S.fin = null; saveBooks(); goClient('books:fs'); return c.id; }"""
OPEN = "() => new Promise((r) => { const go = () => (typeof doAct === 'function' && typeof signInNeeded === 'function' && typeof S === 'object' && S && S.firm) ? (() => { try { if (signInNeeded()) doAct('useOffline'); } finally { r(); } })() : setTimeout(go, 50); go(); })"
MAKE = r"""async (a) => {
  const m = await window.loadFinExport();
  const x = await m.FinXlsx.make("2025", {unit: a.unit, dec: a.dec, period: a.period});
  const u8 = new Uint8Array(x.buffer); let s = ""; for (let i = 0; i < u8.length; i += 0x8000) s += String.fromCharCode.apply(null, u8.subarray(i, i + 0x8000));
  const d = FinSt.build("2025", {unit: a.unit, dec: a.dec, period: a.period}); d.ratios = FinSt.ratios(d); d.notesByKey = {}; d.notes.forEach(n => { d.notesByKey[n.key] = n; });
  return {b64: btoa(s), cells: x.cells, d: JSON.parse(JSON.stringify(d))};
}"""
def resolve(d, src):
    if src[0] == "ppeRow": _, key, col, i, c = src; return (d["ppe"] if key == "8" else d["intang"])[col]["rows"][i][c]
    if src[0] == "ppeTot": _, key, col, c = src; return (d["ppe"] if key == "8" else d["intang"])[col]["total"][c]
    x = d
    for k in src:
        if isinstance(x, dict) and k not in x: return None
        x = x[k]
    return x
def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)): return abs(a - b) < 1e-6 * max(1, abs(b) * 1e-3)
    if a in (None, "") and b in (None, ""): return True
    if isinstance(a, (int, float)) and b in (None, ""): return a == 0
    if isinstance(b, (int, float)) and a in (None, ""): return b == 0
    return a == b
def strip_values(src, dst):
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.startswith("xl/worksheets/sheet"):
                s = re.sub(r'<c ([^>]*?)>(<f>.*?</f>)<v>.*?</v></c>', lambda m: "<c " + re.sub(r'\s*t="[^"]*"', "", m.group(1)) + ">" + m.group(2) + "</c>", data.decode("utf-8"), flags=re.S)
                data = s.encode("utf-8")
            zo.writestr(it, data)

made = []
with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(viewport={"width": 1440, "height": 900}, color_scheme="light", accept_downloads=True); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    pg.goto("http://localhost:%d/" % PORT); pg.wait_for_timeout(2500); pg.evaluate(OPEN); pg.wait_for_timeout(1000)
    pg.evaluate(SETUP, BOOKS); pg.wait_for_timeout(2500)
    tab = lambda k: (pg.click('#app nav.fin-tabs [data-fin-tab="%s"]' % k), pg.wait_for_timeout(600))
    total = lambda: pg.locator("#app [data-fin-tbtotal]")
    print("Every figure opens exactly its ledgers")
    for t, table in (("bs", "bs"), ("pl", "pl")):
        tab(t)
        ids = pg.locator('#app table[data-fin-table="%s"] button[data-fin-fig]' % table).evaluate_all("xs => xs.map(x => x.dataset.finFig)")
        bad = []
        for i in ids:
            tab(t); pg.click('#app table[data-fin-table="%s"] button[data-fin-fig="%s"]' % (table, i)); pg.wait_for_timeout(500)
            if total().count() != 1 or total().get_attribute("data-same") != "1": bad.append((i, total().inner_text() if total().count() else "no total"))
        ok(len(ids) >= (12 if t == "bs" else 10) and not bad, "%s: %d figures, each the total of the ledgers it opens %s" % (t, len(ids), bad[:3]))
    tab("bs"); pg.click('#app table[data-fin-table="bs"] button[data-fin-fig="tr"]'); pg.wait_for_timeout(500)
    keys = pg.locator('#app table[data-fin-table="tb"] tbody tr[data-key]').evaluate_all("xs => xs.map(x => x.dataset.key)")
    ok(sorted(keys) == ["Maple Holdings Inc", "Willow Labs"] and "= the figure 31,90,000.75" in total().inner_text(), "Trade receivables: the sundry debtors, %s: %s" % (keys, total().inner_text()))
    pg.screenshot(path=os.path.join(SHOTS, "links-desktop-light-tb.png"))
    tab("bs"); pg.click('#app table[data-fin-table="bs"] button[data-fin-fig="rs"]'); pg.wait_for_timeout(500)
    ok("the year’s profit" in total().inner_text() and total().get_attribute("data-same") == "1", "Reserves and surplus: its ledgers + the year's profit = the figure (%s)" % total().inner_text())
    tab("pl"); pg.click('#app table[data-fin-table="pl"] button[data-fin-fig="exp"]'); pg.wait_for_timeout(500)
    ok(pg.locator('#app table[data-fin-table="tb"] tbody tr[data-key]').count() >= 10 and total().get_attribute("data-same") == "1", "Other expenses: its %d ledgers, total = the figure" % pg.locator('#app table[data-fin-table="tb"] tbody tr[data-key]').count())
    tab("notes")
    lines = pg.locator("#app button[data-fin-tbline]").evaluate_all("xs => xs.map((x, i) => i)")
    bad = []
    for i in lines:
        tab("notes"); b2 = pg.locator("#app button[data-fin-tbline]").nth(i); lab = b2.inner_text(); b2.click(); pg.wait_for_timeout(400)
        if total().count() and total().get_attribute("data-same") != "1": bad.append((lab, total().inner_text()))
    ok(len(lines) >= 30 and not bad, "every note line (%d): its ledgers' total = the line %s" % (len(lines), bad[:3]))
    pg.click('#app [data-fin-seg="Unit"] button:text-is("Thousands")'); pg.wait_for_timeout(800)
    tab("bs"); pg.click('#app table[data-fin-table="bs"] button[data-fin-fig="cash"]'); pg.wait_for_timeout(500)
    ok(total().get_attribute("data-same") == "1" and "(thousands)" in total().inner_text(), "in thousands too: " + total().inner_text())

    print("\nThe Excel: every figure a formula")
    for nm, per in (("fy", None), ("h1", {"preset": "custom", "from": "20250401", "to": "20250930"})):
        x = pg.evaluate(MAKE, {"unit": 1000, "dec": 2, "period": per})
        path = os.path.join(OUT, nm + ".xlsx"); open(path, "wb").write(base64.b64decode(x["b64"])); made.append((nm, path, x))

    ctx.close()
    for nm, vw, scheme, mob in (("desktop-dark", 1440, "dark", False), ("phone", 390, "light", True)):
        c2 = br.new_context(viewport={"width": vw, "height": 900 if not mob else 844}, color_scheme=scheme, is_mobile=mob, has_touch=mob, device_scale_factor=2 if mob else 1); p2 = c2.new_page()
        p2.goto("http://localhost:%d/" % PORT); p2.wait_for_timeout(2500); p2.evaluate(OPEN); p2.wait_for_timeout(1000); p2.evaluate(SETUP, BOOKS); p2.wait_for_timeout(2500)
        p2.evaluate("() => { S.fin.tab = 'bs'; render(); }"); p2.wait_for_timeout(600)
        p2.click('#app table[data-fin-table="bs"] button[data-fin-fig="tr"]'); p2.wait_for_timeout(800)
        p2.screenshot(path=os.path.join(SHOTS, "links-%s-tb.png" % nm))
        if mob: ok(p2.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth") <= 1, "a phone: no sideways scroll")
        c2.close()
    br.close()

for nm, path, x in made:
    wb = openpyxl.load_workbook(path)
    typed = [(ws.title, c.coordinate, c.value) for ws in (wb["Balance Sheet"], wb["Profit-Loss"], wb["Notes"]) for row in ws.iter_rows() for c in row
             if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]
    nform = sum(1 for ws in (wb["Balance Sheet"], wb["Profit-Loss"], wb["Notes"]) for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    ok(not typed and nform > 150, "%s: every number on the Balance Sheet, Profit-Loss and Notes sheets is a formula (%d formulas; typed: %s)" % (nm, nform, typed[:5]))
soffice = shutil.which("soffice") or shutil.which("libreoffice")
if not soffice: ok(os.environ.get("FIN_XLSX_NO_SOFFICE") == "1", "LibreOffice is installed")
else:
    bare = os.path.join(OUT, "bare"); outd = os.path.join(OUT, "recalc"); os.makedirs(bare); os.makedirs(outd)
    for nm, path, x in made: strip_values(path, os.path.join(bare, os.path.basename(path)))
    subprocess.run([soffice, "-env:UserInstallation=file://" + os.path.join(OUT, "lo"), "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outd] + [os.path.join(bare, os.path.basename(m[1])) for m in made],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    for nm, path, x in made:
        lv = openpyxl.load_workbook(os.path.join(outd, os.path.basename(path)), data_only=True)
        cells = [c for c in x["cells"] if c["sheet"] in ("Balance Sheet", "Profit-Loss", "Notes")]
        bad = [(c["sheet"] + "!" + c["addr"], lv[c["sheet"]][c["addr"]].value, c["v"]) for c in cells if not same(lv[c["sheet"]][c["addr"]].value, c["v"])]
        scr = [(c["sheet"] + "!" + c["addr"], lv[c["sheet"]][c["addr"]].value, resolve(x["d"], c["src"])) for c in cells if c["src"] and not same(lv[c["sheet"]][c["addr"]].value, resolve(x["d"], c["src"]))]
        ok(not bad and not scr and len(cells) > 150, "%s: each of %d formulas recalculates in LibreOffice to the screen's figure %s" % (nm, len(cells), (bad + scr)[:4]))
    shutil.rmtree(os.path.join(OUT, "lo"), ignore_errors=True)
errs = [e for e in errors if "supabase" not in e.lower() and "Failed to load" not in e]
ok(not errs, "no page errors" + ("" if not errs else ": " + " | ".join(errs[:4])))
print("\n" + ("%d FAILED" % len(fails) if fails else "all passed")); sys.exit(1 if fails else 0)
