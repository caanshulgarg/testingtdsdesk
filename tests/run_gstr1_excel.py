"""python3 run_gstr1_excel.py - GSTR-1 as an Excel to download, change and bring back (next-gstr1-excel; the owner,
24-Oct-2026: "prepare this type of download and upload return in GSTR-1 .. user can add or delete or modify the details"
and "it should be dynamic.. add, delete or modify"). Made-up books (gstr1_books.py, worked out by hand):
  1. the GSTR-1 screen as it was, with Download Excel and Upload Excel beside its actions
  2. Download: a sheet a portal table (the offline tool's sheet names and column heads), each an Excel Table; the hidden
     Key column; lists (place of supply, rate, invoice type, reverse charge, UQC...) to the last of 2,000 spare rows;
     the spare rows work their tax out; Summary and About; LibreOffice works the file out to FinCom's tables to the rupee
  3. edits in openpyxl: S/101 at 12% (tax changed), S/102 deleted, B2CS's place of supply 09 -> 05 (IGST); added only
     in Excel, typed in the first empty rows with no tax typed: a B2B invoice, a B2CL invoice, a CDNR note; and an HSN line.
     LibreOffice: the added rows' tax and Check worked out, the Summary takes them all
  4. Upload: the preview lists exactly those changes, no error; the totals; a file with errors is refused (Apply off,
     the errors in words); warnings need a tick
  5. Apply: the GSTR-1 screen's tables and the JSON equal the Excel's Summary to the rupee; 3B 3.1(a) and 3.2 follow
  6. the same file uploaded without being worked out (a program that does not calculate): the same preview
  7. Undo: as before
Screenshots (1440 and 390) to TDSDESK_SHOTS (default tests/out/g1x-shots), and the Excel as LibreOffice draws it."""
import os, sys, json, threading, functools, http.server, base64, shutil, subprocess, tempfile
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import gstr1_books as GB
import openpyxl
from playwright.sync_api import sync_playwright
SITE = os.environ.get("TDSDESK_SITE", os.path.join(HERE, "..", "app", "dist-test"))
PORT = int(os.environ.get("G1X_PORT") or 8461)
SHOTS = os.environ.get("TDSDESK_SHOTS") or os.path.join(os.environ.get("TDSDESK_OUT", os.path.join(HERE, "out")), "g1x-shots")
os.makedirs(SHOTS, exist_ok=True)
fails, errors = [], []
def ok(c, w):
    print(("  ok   " if c else "  FAIL ") + w)
    if not c: fails.append(w)
class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(("localhost", PORT), functools.partial(Quiet, directory=SITE)); threading.Thread(target=srv.serve_forever, daemon=True).start()
H5 = ["taxable", "igst", "cgst", "sgst", "cess"]
def shot(pg, name, full=False): pg.screenshot(path=os.path.join(SHOTS, name + ".png"), full_page=full)
def W(pg, ms=400): pg.wait_for_timeout(ms)
CAPTURE = """() => { window.__files = []; window.saveFile = async (name, data) => { const a = new Uint8Array(await data.arrayBuffer()); let s = ''; for (let i = 0; i < a.length; i += 8192) s += String.fromCharCode.apply(null, a.subarray(i, i + 8192)); window.__files.push({name, b64: btoa(s)}); return true; }; }"""
# the GSTR-1 screen's tables (portal.js r1Portal over the JSON FinCom makes), each table's totals
SCREEN = """() => { const T = GSTPortal.r1Portal(GSTR.toJson('202604', '09'), GSTR.toJson('202604', '09', {plain: true}), {}, []); const o = {};
  Object.keys(T).forEach(k => { o[k] = Object.assign({rows: T[k].rows.length}, T[k].tot); }); return o; }"""
def lo(path, outdir, to="xlsx:Calc MS Excel 2007 XML"):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice: return None
    os.makedirs(outdir, exist_ok=True)
    prof = "file://" + tempfile.mkdtemp(prefix="g1x-lo-")
    subprocess.run([soffice, "-env:UserInstallation=" + prof, "--headless", "--calc", "--convert-to", to, "--outdir", outdir, path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    shutil.rmtree(prof[7:], ignore_errors=True)
    return os.path.join(outdir, os.path.splitext(os.path.basename(path))[0] + "." + to.split(":")[0])
def summary(path):
    ws = openpyxl.load_workbook(path, data_only=True)["Summary"]; out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[8]: out[r[8]] = {"rows": r[2], **{h: round(float(r[3 + i] or 0), 2) for i, h in enumerate(H5)}}
    return out
def cols(ws): return {c.value: c.column for c in ws[1] if c.value}
def first_empty(ws, c):
    for r in range(2, ws.max_row + 2):
        if ws.cell(r, c).value in (None, ""): return r

with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(viewport={"width": 1440, "height": 950}, accept_downloads=True); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    GB.setup(pg, PORT)
    tmp = tempfile.mkdtemp(prefix="g1x-")

    print("1. the GSTR-1 screen")
    GB.open_r1(pg)
    acts = pg.evaluate("[...document.querySelectorAll('#app [data-return-card] .ret-acts button')].map(b => b.innerText.trim()).filter(Boolean)")
    ok("Download Excel" in acts and "Upload Excel" in acts and any("GSTR-1 JSON" in a for a in acts), "Download Excel and Upload Excel beside the return's actions (%s)" % acts)
    ok(pg.locator("#app [data-r1-summary]").count() == 1, "the GSTR-1 summary as before")
    base = pg.evaluate(SCREEN)
    ok(abs(base["b2b"]["taxable"] - 150000) < 0.01 and abs(base["b2cl"]["taxable"] - 200000) < 0.01 and abs(base["cdnr"]["taxable"] - 10000) < 0.01, "the books: B2B 1,50,000, B2CL 2,00,000, CDNR 10,000 (%s)" % {k: base[k]["taxable"] for k in ("b2b", "b2cl", "cdnr")})
    shot(pg, "1-gstr1-buttons-desktop")
    pg.set_viewport_size({"width": 390, "height": 860}); W(pg, 600); shot(pg, "1-gstr1-buttons-phone")
    ok(pg.evaluate("document.documentElement.scrollWidth <= window.innerWidth + 1"), "phone: no sideways scroll")
    pg.set_viewport_size({"width": 1440, "height": 950}); W(pg, 400)

    print("2. download")
    pg.evaluate(CAPTURE); pg.click("[data-r1x=down]")
    for _ in range(120):
        if pg.evaluate("window.__files.length"): break
        W(pg, 250)
    f = pg.evaluate("window.__files[0]")
    ok(f and f["name"].endswith(".xlsx") and "GSTR-1" in f["name"], "Download Excel: %s" % (f and f["name"]))
    xp = os.path.join(tmp, "gstr1.xlsx"); open(xp, "wb").write(base64.b64decode(f["b64"]))
    wb = openpyxl.load_workbook(xp)
    want = ["b2b,sez,de", "b2cl", "exp", "b2cs", "exemp", "cdnr", "cdnur", "b2ba", "b2cla", "expa", "cdnra", "cdnura", "b2csa", "hsn(b2b)", "hsn(b2c)", "docs", "Summary", "About"]
    ok([s for s in wb.sheetnames if s != "Lists"] == want, "a sheet a table, the offline tool's names, Summary and About (%s)" % wb.sheetnames)
    ws = wb["b2b,sez,de"]; C = cols(ws)
    heads = [c.value for c in ws[1]]
    ok(heads[1:6] == ["GSTIN/UIN of Recipient", "State (from GSTIN)", "Receiver Name", "Invoice Number", "Invoice date"] and "Place Of Supply" in heads and "Invoice Type" in heads and "Integrated Tax" in heads and "Cess Amount" in heads, "B2B's column heads as the portal's (%s)" % heads)
    ok(ws.column_dimensions["A"].hidden and str(ws["A2"].value).startswith("b2b|"), "the Key column hidden, a books row keyed (%s)" % ws["A2"].value)
    tables = {s: list(wb[s].tables.keys()) for s in want[:-2]}
    tref = wb["b2b,sez,de"].tables["T_b2b"].ref
    ok(all(len(v) == 1 for v in tables.values()) and tref.endswith(str(2 + 2 + 2000 - 1)), "each table sheet one Excel Table; B2B's reaches 2,000 spare rows (%s)" % tref)
    dv = {str(d.sqref): d.formula1 for d in ws.data_validations.dataValidation}
    def dv_for(col):
        from openpyxl.utils import get_column_letter as L
        return [(k, v) for k, v in dv.items() if k.startswith(L(col) + "2:")]
    pos_dv, rate_dv, typ_dv, rc_dv = dv_for(C["Place Of Supply"]), dv_for(C["Rate"]), dv_for(C["Invoice Type"]), dv_for(C["Reverse Charge"])
    ok(pos_dv and rate_dv and typ_dv and rc_dv and pos_dv[0][0].endswith("2003") and "Lists" in pos_dv[0][1], "lists for place of supply, rate, invoice type and reverse charge to the last spare row (%s)" % pos_dv)
    uq = [d for d in wb["hsn(b2b)"].data_validations.dataValidation if "Lists!$E" in (d.formula1 or "")]
    ok(len(uq) == 1, "UQC a list on the HSN sheets")
    r_sp = first_empty(ws, C["Invoice Number"])
    ok("[#This Row]" in str(ws.cell(r_sp, C["Integrated Tax"]).value) and "[#This Row]" in str(ws.cell(r_sp, C["Central Tax"]).value) and ws.cell(2, C["Central Tax"]).value == 9000, "a spare row works its tax out (structured formulas); a books row keeps the books' tax")
    ok("[#This Row]" in str(ws.cell(2, C["Check (worked out)"]).value) and "LEFT(" in str(ws.cell(2, C["State (from GSTIN)"]).value), "Check and the GSTIN's state are formulas on every row")
    sm = wb["Summary"]
    ok(any("SUM(T_b2b[" in str(c.value) for c in sm[2]) and wb["Summary"].protection.sheet and not ws.protection.sheet, "Summary: formulas over the tables (protected); the tables open")
    d0 = os.path.join(tmp, "lo0"); r0 = lo(xp, d0)
    LO = bool(r0)
    if LO:
        s0 = summary(r0)
        bad = [(t, s0[t], base[t]) for t in s0 if t in base and any(abs(s0[t][h] - base[t][h]) > 0.5 for h in H5)]
        ok(not bad and len(s0) >= 10, "LibreOffice works the Excel out to the GSTR-1 screen's totals, table by table (%d tables%s)" % (len(s0), ", off: %s" % bad if bad else ""))
        chk = {r[8]: r for r in openpyxl.load_workbook(r0, data_only=True)["Summary"].iter_rows(min_row=2, values_only=True) if r[8] and str(r[8]).startswith("chk_")}
        ok(abs(chk["chk_hsnb2b"][3] or 0) < 0.01 and abs(chk["chk_hsnb2c"][3] or 0) < 0.01 and chk["chk_docs"][2] == 5 and chk["chk_inv"][2] == 3, "Summary checks: HSN against tables 4 to 9 (nil), documents issued 5, invoices listed 3 (%s)" % {k: (v[2], v[3]) for k, v in chk.items()})
    else: print("  (LibreOffice not here: the recalculation checks are skipped)")

    print("3. edits in openpyxl")
    wb = openpyxl.load_workbook(xp); ws = wb["b2b,sez,de"]; C = cols(ws)
    rows = {ws.cell(r, C["Invoice Number"]).value: r for r in range(2, 10) if ws.cell(r, C["Invoice Number"]).value}
    r101 = rows["S/101"]; ws.cell(r101, C["Rate"]).value = 12; ws.cell(r101, C["Central Tax"]).value = 6000; ws.cell(r101, C["State/UT Tax"]).value = 6000
    ws.delete_rows(rows["S/102"])
    r = first_empty(ws, C["Invoice Number"])
    for k, v in {"GSTIN/UIN of Recipient": GB.GD, "Receiver Name": "Delta Stores", "Invoice Number": "S/150", "Invoice date": "25-04-2026", "Invoice Value": 11800, "Place Of Supply": "09-Uttar Pradesh",
                 "Reverse Charge": "N", "Invoice Type": "Regular B2B", "Rate": 18, "Taxable Value": 10000}.items(): ws.cell(r, C[k]).value = v
    new_b2b_row = r
    wl = wb["b2cl"]; CL = cols(wl); r = first_empty(wl, CL["Invoice Number"])
    for k, v in {"Invoice Number": "S/151", "Invoice date": "26-04-2026", "Invoice Value": 177000, "Place Of Supply": "07-Delhi", "Rate": 18, "Taxable Value": 150000}.items(): wl.cell(r, CL[k]).value = v
    wn = wb["cdnr"]; CN = cols(wn); r = first_empty(wn, CN["Note Number"])
    for k, v in {"GSTIN/UIN of Recipient": GB.GD, "Receiver Name": "Delta Stores", "Note Number": "CN/2", "Note Date": "28-04-2026", "Note Type": "C", "Place Of Supply": "09-Uttar Pradesh", "Reverse Charge": "N",
                 "Note Supply Type": "Regular B2B", "Note Value": 1180, "Rate": 18, "Taxable Value": 1000}.items(): wn.cell(r, CN[k]).value = v
    wh = wb["hsn(b2b)"]; CH = cols(wh); r = first_empty(wh, CH["HSN"])
    for k, v in {"HSN": "9983", "Description": "Consulting", "UQC": "NA", "Total Quantity": 0, "Rate": 18, "Taxable Value": 10000, "Integrated Tax": 0, "Central Tax": 900, "State/UT Tax": 900, "Cess Amount": 0}.items(): wh.cell(r, CH[k]).value = v
    wc = wb["b2cs"]; CS = cols(wc)
    wc.cell(2, CS["Place Of Supply"]).value = "05-Uttarakhand"; wc.cell(2, CS["Integrated Tax"]).value = 3600; wc.cell(2, CS["Central Tax"]).value = 0; wc.cell(2, CS["State/UT Tax"]).value = 0
    ed = os.path.join(tmp, "edited.xlsx"); wb.save(ed)
    s1 = None
    if LO:
        d1 = os.path.join(tmp, "lo1"); r1 = lo(ed, d1); s1 = summary(r1)
        v = openpyxl.load_workbook(r1, data_only=True)
        b = v["b2b,sez,de"]; rr = next(i for i in range(2, 40) if b.cell(i, C["Invoice Number"]).value == "S/150")
        cl = v["b2cl"]; rl = next(i for i in range(2, 40) if cl.cell(i, CL["Invoice Number"]).value == "S/151")
        cn = v["cdnr"]; rn = next(i for i in range(2, 40) if cn.cell(i, CN["Note Number"]).value == "CN/2")
        ok(b.cell(rr, C["Central Tax"]).value == 900 and b.cell(rr, C["Integrated Tax"]).value == 0 and b.cell(rr, C["State (from GSTIN)"]).value == "09" and cl.cell(rl, CL["Integrated Tax"]).value == 27000 and cn.cell(rn, CN["Central Tax"]).value == 90,
           "LibreOffice: the rows added with no tax typed work it out (S/150 CGST 900, S/151 IGST 27,000, CN/2 CGST 90) and the GSTIN's state")
        ok(b.cell(r101 if r101 < rows["S/102"] else r101 - 1, C["Check (worked out)"]).value in ("", None) and v["b2cs"].cell(2, CS["Check (worked out)"]).value in ("", None), "their Check is clear; S/101's too (rate and tax changed together)")
        ok(abs(s1["b2b"]["taxable"] - 110000) < 0.01 and abs(s1["b2b"]["cgst"] - 6900) < 0.01 and abs(s1["b2cl"]["taxable"] - 350000) < 0.01 and abs(s1["b2cl"]["igst"] - 63000) < 0.01 and abs(s1["cdnr"]["taxable"] - 11000) < 0.01
           and abs(s1["hsnb2b"]["taxable"] - 150000) < 0.01 and abs(s1["b2cs"]["igst"] - 3600) < 0.01, "the Summary takes the added, deleted and changed rows: B2B 1,10,000, B2CL 3,50,000, CDNR 11,000, HSN B2B 1,50,000, B2CS IGST 3,600")
        # the Excel as LibreOffice draws it: the B2B sheet with the new row and its worked-out tax
        # (LibreOffice's own figures, drawn by its HTML export: headless LibreOffice has no picture filter for a sheet here)
        try:
            htm = lo(r1, os.path.join(tmp, "html"), "html:HTML (StarCalc)")
            if htm and os.path.exists(htm):
                xp2 = br.new_page(viewport={"width": 1440, "height": 260}); xp2.goto("file://" + htm); xp2.wait_for_timeout(500)
                xp2.evaluate("""() => { const t = document.querySelector('table'); [...document.body.children].forEach(el => { if (el !== t && !el.contains(t) && el.tagName !== 'H1') el.remove(); }); document.querySelectorAll('h1').forEach((x, i) => { if (i) x.remove(); }); [...t.querySelectorAll('tr')].forEach(tr => { if (/S\/150/.test(tr.innerText)) tr.style.background = '#fff4bf'; else if (!tr.innerText.trim()) tr.remove(); });
                  const h = document.createElement('p'); h.textContent = 'b2b,sez,de as LibreOffice works it out: S/150 typed with no tax; its Integrated/Central/State tax, its state and Check worked out by the Table formulas'; h.style.cssText = 'font:13px sans-serif;margin:6px'; document.body.prepend(h); }""")
                xp2.screenshot(path=os.path.join(SHOTS, "4-excel-b2b-added-row-libreoffice.png")); xp2.close()
        except Exception as e: print("  (the Excel picture was not made: %s)" % str(e)[:120])

    print("4. upload: the preview")
    pv = '[data-pop="r1x-import"]'
    # first a file with errors: refused
    wbe = openpyxl.load_workbook(ed); we = wbe["b2b,sez,de"]
    we.cell(new_b2b_row, C["GSTIN/UIN of Recipient"]).value = GB.GD[:14] + ("A" if GB.GD[14] != "A" else "B")
    we.cell(new_b2b_row, C["Rate"]).value = 17; we.cell(new_b2b_row, C["Central Tax"]).value = 850; we.cell(new_b2b_row, C["State/UT Tax"]).value = 850
    bad = os.path.join(tmp, "bad.xlsx"); wbe.save(bad)
    pg.set_input_files("[data-r1x-upload]", bad); W(pg, 2500)
    ok(pg.locator(pv).count() == 1 and int(pg.get_attribute(pv + " [data-pv-errors]", "data-pv-errors") or 0) >= 2 and pg.locator(pv + " [data-pv-apply]").is_disabled(), "a file with errors: the preview lists them, Apply is off")
    et = pg.inner_text(pv + " [data-pv-errors]")
    ok("is not valid" in et and "17% is not a GST rate" in et, "the errors in words (%s)" % et[:160].replace("\n", " | "))
    shot(pg, "2-upload-preview-error-desktop")
    pg.set_viewport_size({"width": 390, "height": 860}); W(pg, 500); shot(pg, "2-upload-preview-error-phone"); pg.set_viewport_size({"width": 1440, "height": 950}); W(pg, 300)
    pg.keyboard.press("Escape"); W(pg, 400)
    # a warning needs a tick
    wbw = openpyxl.load_workbook(ed); ww = wbw["b2b,sez,de"]; ww.cell(new_b2b_row, C["Invoice date"]).value = "25-03-2026"; warn = os.path.join(tmp, "warn.xlsx"); wbw.save(warn)
    pg.set_input_files("[data-r1x-upload]", warn); W(pg, 2500)
    d_before = pg.locator(pv + " [data-pv-apply]").is_disabled()
    if pg.locator(pv + " [data-pv-confirm]").count(): pg.check(pv + " [data-pv-confirm]"); W(pg, 300)
    ok(d_before and pg.locator(pv + " [data-pv-warnings]").count() == 1 and not pg.locator(pv + " [data-pv-apply]").is_disabled(), "a date outside the month is a warning: Apply once ticked")
    pg.keyboard.press("Escape"); W(pg, 400)
    # the good file, as LibreOffice saved it
    good = r1 if LO else ed
    pg.set_input_files("[data-r1x-upload]", good); W(pg, 2500)
    counts = pg.evaluate("() => Object.fromEntries([...document.querySelectorAll('[data-pop=\"r1x-import\"] [data-pv-table]')].map(s => [s.dataset.pvTable, [s.dataset.pvAdded, s.dataset.pvDeleted, s.dataset.pvModified].join('/')]))")
    ok(counts == {"b2b": "1/1/1", "b2cl": "1/0/0", "cdnr": "1/0/0", "hsnb2b": "1/0/0", "b2cs": "0/0/1"}, "the preview: exactly B2B 1 added, 1 left out, 1 changed; B2CL, CDNR and HSN B2B 1 added each; B2CS 1 changed (%s)" % counts)
    flds = pg.evaluate("() => [...document.querySelectorAll('[data-pv-table=b2b] [data-pv-field]')].map(li => li.dataset.pvField).sort()")
    ok(flds == ["camt", "rt", "samt"] and pg.locator(pv + " [data-pv-errors]").count() == 0 and pg.locator(pv + " [data-pv-warnings]").count() == 0, "S/101's rate, CGST and SGST old → new; no error, no warning (%s)" % flds)
    ok("09-" not in pg.inner_text(pv + " [data-pv-table=b2cs]") and pg.locator(pv + " [data-pv-table=b2cs] [data-pv-field=pos]").count() == 1, "B2CS: the place of supply changed")
    ok(pg.locator(pv + " [data-pv-tot='b2b|excel']").count() == 1 and pg.locator(pv + " [data-pv-tot='b2b|after']").count() == 1, "the totals: the Excel's, FinCom's now, after Apply")
    shot(pg, "3-upload-preview-desktop")
    pg.set_viewport_size({"width": 390, "height": 860}); W(pg, 500); shot(pg, "3-upload-preview-phone")
    ok(pg.evaluate("() => { const b = document.querySelector('[data-pop=\"r1x-import\"] [data-r1x-preview]'); return b.scrollWidth <= b.clientWidth + 1 && b.getBoundingClientRect().right <= window.innerWidth + 1; }"), "phone: the preview fits the screen (its totals table scrolls on its own)")
    pg.set_viewport_size({"width": 1440, "height": 950}); W(pg, 300)

    print("5. Apply")
    pg.click(pv + " [data-pv-apply]"); W(pg, 1200)
    after = pg.evaluate(SCREEN)
    ref = s1 or {"b2b": {"taxable": 110000, "igst": 0, "cgst": 6900, "sgst": 6900, "cess": 0}}
    off = [(t, ref[t], after[t]) for t in ref if t in after and any(abs(ref[t][h] - after[t][h]) > 0.5 for h in H5)]
    ok(not off and len(ref) >= 1, "after Apply the GSTR-1 screen's tables equal the Excel's Summary to the rupee, table by table (%d%s)" % (len(ref), ", off: %s" % off if off else ""))
    js = pg.evaluate("() => { const j = GSTR.toJson('202604', '09'); const s = (a, f) => (a || []).reduce((x, g) => x + g.inv.reduce((y, i) => y + i.itms.reduce((z, t) => z + (t.itm_det[f] || 0), 0), 0), 0);"
                     " return {b2b: s(j.b2b, 'txval'), b2bc: s(j.b2b, 'camt'), b2cl: s(j.b2cl, 'txval'), has102: JSON.stringify(j).includes('S/102'), b2cs: j.b2cs[0].pos + '/' + j.b2cs[0].iamt, hsn: j.hsn.hsn_b2b.length}; }")
    ok(abs(js["b2b"] - 110000) < 0.01 and abs(js["b2bc"] - 6900) < 0.01 and abs(js["b2cl"] - 350000) < 0.01 and not js["has102"] and js["b2cs"] == "05/3600" and js["hsn"] == 2, "the JSON for the portal is the same (%s)" % js)
    t3 = pg.evaluate("() => { const t = GSTR.threeB('202604', '09'); return {a: t.net.taxable, ig: t.net.igst, cg: t.net.cgst, up: (t.unregPos || []).map(x => x.pos + ':' + x.taxable).join(',')}; }")
    ok(abs(t3["a"] - (360000 - 50000 + 10000 + 150000 - 1000)) < 0.01 and "05:20000" in t3["up"] and "07:350000" in t3["up"], "3B 3.1(a) and 3.2 follow (%s)" % t3)
    ok(pg.evaluate("S.books.vouchers.some(v => v.no === 'S/102')") and pg.evaluate("!!S.books.r1Fix['09|202604'].cur"), "S/102 stays in the books; the corrections kept with the books (r1Fix)")
    GB.open_r1(pg, sub="details"); W(pg, 600)
    ok("S/150" in pg.inner_text("#app [data-r1-portal]") and "S/102" not in pg.inner_text("#app [data-r1-portal]") and pg.locator("[data-r1x=undo]").count() == 1, "the screen's B2B table shows S/150, not S/102; Undo upload is offered")
    GB.open_r1(pg); W(pg, 500)
    ok(pg.locator("#app [data-r1x-note]").count() == 1 and "books are not changed" in pg.inner_text("#app [data-r1x-note]"), "the summary says the month's GSTR-1 has corrections made in FinCom")
    shot(pg, "5-after-apply-desktop")
    pg.set_viewport_size({"width": 390, "height": 860}); W(pg, 500); shot(pg, "5-after-apply-phone"); pg.set_viewport_size({"width": 1440, "height": 950}); W(pg, 300)

    print("6. the file not worked out, uploaded again")
    pg.set_input_files("[data-r1x-upload]", ed); W(pg, 2500)
    left = pg.evaluate("() => [...document.querySelectorAll('[data-pop=\"r1x-import\"] [data-pv-table] li')].map(li => li.innerText.replace(/\\s+/g, ' ')).slice(0, 8)")
    ok(pg.get_attribute(pv + " [data-r1x-preview]", "data-pv-n") == "0" and pg.locator(pv + " [data-pv-errors]").count() == 0, "(%s) " % left + "the same edits saved by a program that does not calculate: nothing more to apply (the tax of the added rows worked out on upload)")
    pg.keyboard.press("Escape"); W(pg, 400)

    print("7. Undo")
    pg.click("[data-r1x=undo]"); W(pg, 1000)
    back = pg.evaluate(SCREEN)
    ok(all(abs(back[t][h] - base[t][h]) < 0.01 for t in base for h in H5) and not pg.evaluate("S.books.r1Fix['09|202604'].cur"), "Undo: every table as before")
    ok(pg.locator("#app [data-r1x-note]").count() == 0 and pg.locator("[data-r1x=undo]").count() == 0, "and no note, nothing more to undo")
    ok(not errors, "no page errors (%s)" % errors[:3])
    br.close()
shutil.rmtree(tmp, ignore_errors=True)
print("\n%d FAILED" % len(fails) if fails else "\nall passed")
sys.exit(1 if fails else 0)
