"""Recalculate an .xlsx written by FinCom's reports (src/js/78-reports.js) independently of FinCom, and compare.

Every formula cell's cached value is removed first, so the calculator has to work each formula out itself. The calculator
is LibreOffice (headless) when it is installed; otherwise a small evaluator for the formulas the reports write (+ - * / ^,
names, cell and range references, SUM, SUMIFS, ROUND, MAX, MIN, ABS, IF, OR, ISTEXT), written here in Python, apart
from FinCom's own JavaScript.

    python3 reports_recalc.py file.xlsx      # prints the cells whose value differs from FinCom's
"""
import os, re, shutil, subprocess, sys, tempfile, zipfile, math

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def _strip_values(src, dst):
    """A copy of the workbook with no cached value on any formula cell."""
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.startswith("xl/worksheets/sheet"):
                s = data.decode("utf-8")
                s = re.sub(r'(<c [^>]*?)( t="str")?(>)(<f>.*?</f>)<v>.*?</v>(</c>)', r"\1\3\4\5", s)
                data = s.encode("utf-8")
            zo.writestr(it, data)


def cached(path):
    """FinCom's own values: {(sheet, 'B7'): value} for every formula cell."""
    import openpyxl
    wf = openpyxl.load_workbook(path, data_only=False)
    wv = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for ws in wf.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out[(ws.title, c.coordinate)] = wv[ws.title][c.coordinate].value
    return out


def recalc_libreoffice(path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    import openpyxl
    d = tempfile.mkdtemp(prefix="fx-recalc-")
    bare = os.path.join(d, "bare.xlsx")
    _strip_values(path, bare)
    outd = os.path.join(d, "out")
    os.makedirs(outd)
    prof = "file://" + os.path.join(d, "profile")
    subprocess.run([soffice, "-env:UserInstallation=" + prof, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outd, bare],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    wb = openpyxl.load_workbook(os.path.join(outd, "bare.xlsx"), data_only=True)
    keys = cached(path).keys()
    return {k: wb[k[0]][k[1]].value for k in keys}


# ---------------------------------------------------------------- the fallback: a small evaluator
def _col_index(s):
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


def recalc_python(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=False)
    names = {}
    for n, dn in wb.defined_names.items():
        names[n] = dn.attr_text
    memo = {}

    def cell_value(sheet, ref):
        key = (sheet, ref.replace("$", ""))
        if key in memo:
            return memo[key]
        v = wb[sheet][key[1]].value
        if isinstance(v, str) and v.startswith("="):
            memo[key] = None
            v = evaluate(v[1:], sheet)
        memo[key] = v
        return v

    def rng(sheet, a, b):
        ma, mb = re.match(r"\$?([A-Z]+)\$?(\d+)", a), re.match(r"\$?([A-Z]+)\$?(\d+)", b)
        c1, r1, c2, r2 = _col_index(ma.group(1)), int(ma.group(2)), _col_index(mb.group(1)), int(mb.group(2))
        out = []
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                out.append(cell_value(sheet, openpyxl.utils.get_column_letter(c) + str(r)))
        return out

    def ref_of(text):
        m = re.match(r"^'?(.*?)'?!(\$?[A-Z]+\$?\d+)(?::(\$?[A-Z]+\$?\d+))?$", text)
        return m.group(1).replace("''", "'"), m.group(2), m.group(3)

    def name_value(n, sheet):
        sh, a, b = ref_of(names[n])
        return rng(sh, a, b) if b else cell_value(sh, a)

    def evaluate(f, sheet):
        toks = re.findall(r'"(?:[^"]|"")*"|\d+\.?\d*(?:[eE]-?\d+)?|\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+|[A-Za-z_][A-Za-z0-9_.]*|<=|>=|<>|[-+*/^(),=<>]', f)
        pos = [0]

        def peek():
            return toks[pos[0]] if pos[0] < len(toks) else None

        def take():
            pos[0] += 1
            return toks[pos[0] - 1]

        def num(v):
            if v is None or v == "":
                return 0
            if isinstance(v, bool):
                return 1 if v else 0
            return v

        def cmp_():
            a = expr()
            while peek() in ("=", "<", ">", "<=", ">=", "<>"):
                o = take(); b = expr()
                a = {"=": a == b, "<": a < b, ">": a > b, "<=": a <= b, ">=": a >= b, "<>": a != b}[o]
            return a

        def expr():
            a = term()
            while peek() in ("+", "-"):
                o = take(); b = term()
                a = num(a) + num(b) if o == "+" else num(a) - num(b)
            return a

        def term():
            a = unary()
            while peek() in ("*", "/"):
                o = take(); b = unary()
                a = num(a) * num(b) if o == "*" else num(a) / num(b)
            return a

        def unary():
            if peek() == "-":
                take(); return -num(unary())
            if peek() == "+":
                take(); return num(unary())
            return power()

        def power():
            a = atom()
            while peek() == "^":
                take(); a = math.pow(num(a), num(atom()))
            return a

        def args():
            out = []
            if peek() == ")":
                take(); return out
            while True:
                out.append(cmp_())
                if peek() == ",":
                    take(); continue
                take(); return out

        def xround(x, d):
            k = 10 ** int(d)
            y = abs(x) * k
            r = math.floor(y + 0.5 + 1e-9)
            return math.copysign(r / k, x) if r else 0.0

        def atom():
            t = take()
            if t == "(":
                v = cmp_(); take(); return v
            if t.startswith('"'):
                return t[1:-1].replace('""', '"')
            if re.match(r"^\d", t):
                return float(t)
            if re.match(r"^\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?\d+$", t):
                a, b = t.split(":"); return rng(sheet, a, b)
            if peek() == "(":
                take(); fn = t.upper(); a = args()
                if fn == "SUM":
                    s = 0
                    for x in a:
                        s += sum(num(y) for y in x if not isinstance(y, str)) if isinstance(x, list) else num(x)
                    return s
                if fn == "SUMIFS":
                    vals, crit = a[0], a[1:]
                    s = 0
                    for i, v in enumerate(vals):
                        if all((crit[j][i] or "") == crit[j + 1] for j in range(0, len(crit), 2)):
                            s += num(v)
                    return s
                if fn == "ROUND":
                    return xround(num(a[0]), num(a[1]))
                if fn == "MAX":
                    return max(num(x) for x in a)
                if fn == "MIN":
                    return min(num(x) for x in a)
                if fn == "ABS":
                    return abs(num(a[0]))
                if fn == "IF":
                    return a[1] if a[0] else (a[2] if len(a) > 2 else False)
                if fn == "OR":
                    return any(a)
                if fn == "AND":
                    return all(a)
                if fn == "ISTEXT":
                    return isinstance(a[0], str)
                if fn == "POWER":
                    return math.pow(num(a[0]), num(a[1]))
                raise ValueError("function " + fn)
            if re.match(r"^\$?[A-Z]+\$?\d+$", t):
                return cell_value(sheet, t)
            if t in names:
                return name_value(t, sheet)
            raise ValueError("token " + t)

        return cmp_()

    out = {}
    for k in cached(path).keys():
        out[k] = cell_value(k[0], k[1])
    return out


def recalc(path, prefer="libreoffice"):
    """{(sheet, cell): value} worked out independently, and which calculator did it."""
    if prefer == "libreoffice":
        try:
            v = recalc_libreoffice(path)
            if v is not None:
                return v, "LibreOffice"
        except Exception as e:  # LibreOffice present but failing: fall back, and say so
            sys.stderr.write("LibreOffice recalculation failed (%s); using the Python evaluator\n" % e)
    return recalc_python(path), "Python evaluator"


def differences(path, prefer="libreoffice", tol=0.005):
    mine = cached(path)
    theirs, how = recalc(path, prefer)
    bad = []
    for k, v in mine.items():
        w = theirs.get(k)
        if isinstance(v, (int, float)) and isinstance(w, (int, float)):
            if abs(v - w) > 1e-6 * max(1.0, abs(v)):
                bad.append((k, v, w))
        elif (v if v not in ("", None) else None) != (w if w not in ("", None) else None):
            bad.append((k, v, w))
    return bad, how, len(mine)


if __name__ == "__main__":
    for p in sys.argv[1:]:
        for how in ("libreoffice", "python"):
            bad, used, n = differences(p, how)
            print(p, used, n, "formulas,", len(bad), "differ")
            for b in bad[:20]:
                print("  ", b)
