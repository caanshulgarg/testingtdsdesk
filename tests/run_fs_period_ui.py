"""python3 run_fs_period_ui.py - the Financial statements page for a period of the year, and last year's balance sheet
uploaded on the Trial balance & mapping tab (owner, 29-Oct-2026). On the React build (app/dist-test, or TDSDESK_SITE), with
the made-up books of tests/fixtures/fin-statements spread over four quarter-end journals (as tests/run_fs_period.js):
  - the header's Period: each preset (Full year, H1, Q1-Q4, 9 months, Year to date) and Custom From / To give the
    statements for those dates: the P&L and its headings "for the period 01-Apr-2025 to 30-Sep-2025", the balance sheet
    "as at 30-Sep-2025" (it ties), last year's same dates as the comparative; the choice is kept for the client;
  - the computation of income says it is for a full year; the checks read 0;
  - the Excel for H1, Q2 and a custom range, recalculated by LibreOffice from its formulas alone, equals the screen
    (every figure of FinSt.build), with the period in every heading; the PDF too (no computation page);
  - last year's balance sheet (a made-up Excel, and a text PDF made from the same lines) is read, shown as a table to
    confirm, kept; ledgers' codes are then suggested with their reason, a hand choice is kept, and the Checks tab shows
    the difference per head against FinCom's opening balances; a PDF with no text is said plainly not to be read;
  - a phone (390 wide): the period controls wrap, the page never scrolls sideways.
Samples (the sample balance sheet .xlsx / .pdf and the H1 workbook) go to TDSDESK_OUT (tests/out)/fs-period."""
import os, re, sys, json, shutil, zipfile, subprocess, threading, functools, http.server, base64
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
from playwright.sync_api import sync_playwright
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.environ.get("TDSDESK_SITE", os.path.join(HERE, "..", "app", "dist-test"))
OUT = os.path.join(os.environ.get("TDSDESK_OUT", os.path.join(HERE, "out")), "fs-period")
shutil.rmtree(OUT, ignore_errors=True); os.makedirs(OUT)
fails, errors = [], []
def ok(c, w):
    print(("  ok   " if c else "  FAIL ") + w)
    if not c: fails.append(w)
class Q(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(("localhost", 0), functools.partial(Q, directory=SITE)); PORT = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()
PYBS = json.load(open(os.path.join(HERE, "fixtures", "fin-statements", "py-bs-2025.json")))

# the books, each year's movements over four quarter-end journals (20% / 30% / 25% / the rest); from2025: a copy that
# starts on 1-Apr-2025 (last year's closing as its opening), so last year's balance sheet is the only PY there is
GEN = r"""
const F = require('./fixtures/fin-statements/company.json'), from2025 = process.argv[1] === '2025', from = from2025 ? '20250401' : '20240401';
const under = Object.assign({}, F.groups), led = {}, open = from2025 ? Object.assign({}, F.p, {'Profit & Loss A/c': -4531000}) : F.o;
Object.entries(open).forEach(([n, v]) => { led[n] = {open: -v, close: 0, parent: under[n]}; });
const r2 = v => Math.round(v * 100) / 100, W = [0.2, 0.3, 0.25], QE = ['0630', '0930', '1231', '0331'], parts = v => { const a = W.map(w => r2(v * w)); return a.concat([r2(v - a[0] - a[1] - a[2])]); };
const vouchers = [], bank = Object.keys(under).find(n => under[n] === 'Bank Accounts'), shift = e => String(Number(e.slice(0, 4)) - 1) + e.slice(4);
(F.additions || []).filter(([, d]) => d >= from).forEach(([l, d, a], i) => vouchers.push({id: 'a' + i, date: d, no: 'A' + i, type: 'Journal', narr: 'Purchase of ' + l, ent: [{l, a: -a}, {l: bank, a}]}));
const years = from2025 ? [['20260331', F.c, F.plc, F.p]] : [['20250331', F.p, F.plp, F.o], ['20260331', F.c, F.plc, F.p]];
years.forEach(([end, bs, pl, prev]) => {
  const y = Number(end.slice(0, 4)) - 1, mv = {};
  Object.keys(bs).forEach(n => { if (n === bank) return; const added = (F.additions || []).filter(([l, d]) => l === n && d <= end && d > shift(end)).reduce((s, z) => s + z[2], 0);
    const m = r2((bs[n] || 0) - (prev[n] || 0) - added); if (m) mv[n] = m; });
  Object.entries(pl).forEach(([n, v]) => { if (v) mv[n] = v; });
  [0, 1, 2, 3].forEach(q => { const ent = [];
    Object.entries(mv).forEach(([n, v]) => { const a = parts(v)[q]; if (a) ent.push({l: n, a: -a}); });
    ent.push({l: bank, a: r2(-ent.reduce((s, e) => s + e.a, 0))});
    vouchers.push({id: 'm' + end + q, date: (q === 3 ? y + 1 : y) + QE[q], no: 'M' + end + q, type: 'Journal', narr: 'Movements for the quarter (fixture)', ent}); });
});
process.stdout.write(JSON.stringify({under, groups: {}, vouchers, meta: {from, to: '20260331'}, tb: {from, to: '20260331', at: '2026-04-30T10:00:00Z', src: 'copy', led}}));
"""
BOOKS = json.loads(subprocess.check_output(["node", "-e", GEN, "2024"], cwd=HERE))
BOOKS25 = json.loads(subprocess.check_output(["node", "-e", GEN, "2025"], cwd=HERE))
SETUP = """(bk) => {
  S.firm = Object.assign({}, S.firm, {firmName: "Made-up & Associates", frn: "000000S", partner: "A. Sample", mno: "000001", place: "Bengaluru"});
  const c = newCompany({name: "Larkspur Analytics Private Limited (made-up)", gstin: "29AABCL1234F1Z5"}); c.pan = "AABCL1234F"; c.cin = "U72900KA2015PTC000001";
  S.companies[c.id] = c; Store.saveCompany(c); S.coId = c.id; S.view = "company"; S.tab = "books"; S.loadingCo = false;
  const b = Object.assign({loading: false, cid: c.id, challans: [], alloc: {}, ledInfo: {}, gstins: {}, pans: {}, states: {}, misCfg: {freq: "off"}, auditCfg: {freq: "off"},
    msme: {"Bluebell Print Works": "Micro"}, fs: {taxDep: {"2024": {c40: 220000, pm15: 140000, f10: 80000}}, inputs: {unit: 1, dec: 2}}}, bk);
  b.map = Books.mapLedgers(b.vouchers, {});
  S.books = b; window.__bk = b; S.fsFy = "2025"; S.fin = null; saveBooks(); goClient('books:fs'); return c.id; }"""
OPEN = "() => new Promise((r) => { const go = () => (typeof doAct === 'function' && typeof signInNeeded === 'function' && typeof S === 'object' && S && S.firm) ? (() => { try { if (signInNeeded()) doAct('useOffline'); } finally { r(); } })() : setTimeout(go, 50); go(); })"
MAKE = r"""async (a) => {
  const m = await window.loadFinExport();
  const x = await m.FinXlsx.make("2025", {unit: a.unit, dec: a.dec, period: a.period});
  const u8 = new Uint8Array(x.buffer); let s = ""; for (let i = 0; i < u8.length; i += 0x8000) s += String.fromCharCode.apply(null, u8.subarray(i, i + 0x8000));
  const d = FinSt.build("2025", {unit: a.unit, dec: a.dec, period: a.period}); d.ratios = FinSt.ratios(d); d.notesByKey = {}; d.notes.forEach(n => { d.notesByKey[n.key] = n; });
  return {b64: btoa(s), cells: x.cells, sheets: x.sheets, name: x.name, d: JSON.parse(JSON.stringify(d))};
}"""

def resolve(d, src):
    if src[0] == "ppeRow":
        _, key, col, i, c = src; return (d["ppe"] if key == "8" else d["intang"])[col]["rows"][i][c]
    if src[0] == "ppeTot":
        _, key, col, c = src; return (d["ppe"] if key == "8" else d["intang"])[col]["total"][c]
    x = d
    for k in src:
        if isinstance(x, dict) and k not in x: return None
        x = x[k]
    return x
def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)): return abs(a - b) < 1e-6 * max(1, abs(b) * 1e-3)
    if a in (None, "") and b in (None, ""): return True
    if isinstance(a, (int, float)) and b in (None, ""): return a == 0
    if isinstance(b, (int, float)) and a in (None, ""): return b == 0
    return a == b
def strip_values(src, dst):
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.startswith("xl/worksheets/sheet"):
                s = data.decode("utf-8")
                s = re.sub(r'<c ([^>]*?)>(<f>.*?</f>)<v>.*?</v></c>', lambda m: "<c " + re.sub(r'\s*t="[^"]*"', "", m.group(1)) + ">" + m.group(2) + "</c>", s, flags=re.S)
                data = s.encode("utf-8")
            zo.writestr(it, data)
def recalc(files):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice: return None
    bare = os.path.join(OUT, "bare"); outd = os.path.join(OUT, "recalc"); os.makedirs(bare, exist_ok=True); os.makedirs(outd, exist_ok=True)
    names = []
    for f in files:
        b = os.path.join(bare, os.path.basename(f)); strip_values(f, b); names.append(b)
    subprocess.run([soffice, "-env:UserInstallation=file://" + os.path.join(OUT, "lo-profile"), "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outd] + names,
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    return {f: os.path.join(outd, os.path.basename(f)) for f in files}

# the sample last-year balance sheet (made up): as an Excel, and as a PDF with a text layer, from the same lines
def ind(v):
    if v in ("", None): return ""
    s = "{:,.0f}".format(abs(v)) if v == int(v) else "{:,.2f}".format(abs(v))
    # Indian grouping: 12,34,567
    head, _, tail = s.replace(",", "").partition(".")
    h = head[-3:]; rest = head[:-3]
    while rest: h = rest[-2:] + "," + h; rest = rest[:-2]
    s = h + ("." + tail if tail else "")
    return "(" + s + ")" if v < 0 else s
def sample_xlsx(path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Balance Sheet"
    for r in [[PYBS["title"]], [PYBS["what"]], [PYBS["unit"]], PYBS["head"]] + PYBS["rows"]: ws.append(r)
    wb.save(path)
def sample_html():
    rows = "".join("<tr><td>%s</td><td>%s</td><td class=n>%s</td><td class=n>%s</td></tr>" % (r[0], r[1] if len(r) > 1 else "", ind(r[2]) if len(r) > 2 else "", ind(r[3]) if len(r) > 3 else "") for r in PYBS["rows"])
    h = PYBS["head"]
    return ("<html><head><style>body{font:11px Arial}td{padding:1px 14px}td.n{text-align:right}</style></head><body><h3>%s</h3><p>%s</p><p>%s</p><table><tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>%s</table></body></html>"
            % (PYBS["title"], PYBS["what"], PYBS["unit"], h[0], h[1], h[2], h[3], rows))

PERIODS = {"fy": ("20250401", "20260331"), "h1": ("20250401", "20250930"), "q1": ("20250401", "20250630"), "q2": ("20250701", "20250930"), "q3": ("20251001", "20251231"),
           "q4": ("20260101", "20260331"), "m9": ("20250401", "20251231")}
made = []
with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(viewport={"width": 1440, "height": 900}, color_scheme="light", accept_downloads=True); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    pg.goto("http://localhost:%d/" % PORT); pg.wait_for_timeout(2500); pg.evaluate(OPEN); pg.wait_for_timeout(1000)
    pg.evaluate(SETUP, BOOKS); pg.wait_for_timeout(2500)
    ok(pg.locator("#app [data-fin-head] [data-fin-period]").count() == 1, "the header has the Period")
    opts = pg.locator('#app [data-fin-preset] option').evaluate_all("xs => xs.map(x => x.textContent)")
    ok(opts == ["Full year", "H1 Apr–Sep", "Q1 Apr–Jun", "Q2 Jul–Sep", "Q3 Oct–Dec", "Q4 Jan–Mar", "9 months Apr–Dec", "Year to date", "Custom"], "the presets: %s" % opts)
    fin = lambda: pg.evaluate("() => { const e = document.querySelector('#app [data-fin]'); return [e.dataset.finFrom, e.dataset.finTo]; }")
    title = lambda: pg.locator("#app [data-fin-title]").first.inner_text().strip()
    cols = lambda: [pg.locator("#app th[data-fin-colcy]").first.inner_text().strip(), pg.locator("#app th[data-fin-colpy]").first.inner_text().strip()]
    tab = lambda k: (pg.click('#app nav.fin-tabs [data-fin-tab="%s"]' % k), pg.wait_for_timeout(600))
    tot = lambda r: pg.locator('#app table[data-fin-table="bs"] tr[data-fin-total="%s"] td' % r).nth(2).inner_text().strip()
    model = lambda per, js: pg.evaluate("(a) => { const d = FinSt.build('2025', {unit: 1, dec: 2, period: a}); return " + js + "; }", per)
    full_bs = tot("II")
    ok(full_bs == "85,87,800.55" and title() == "Balance sheet as at 31-Mar-2026", "Full year: as today, assets 85,87,800.55 (%s), %s" % (full_bs, title()))

    print("\nThe presets")
    for k, (f, t) in PERIODS.items():
        pg.select_option("#app [data-fin-preset]", k); pg.wait_for_timeout(900)
        ok(fin() == [f, t] and (k == "fy" or pg.evaluate("S.books.fs.period && S.books.fs.period.preset") == k), "%s: %s to %s, kept for the client (%s)" % (k, f, t, fin()))
        tab("bs")
        want = model({"preset": "custom", "from": f, "to": t}, "d.bs.total.assets.cy")
        ok(tot("I") == tot("II") and tot("II").replace(",", "") == "%.2f" % want, "... the balance sheet ties, as FinSt's: %s / %s" % (tot("I"), tot("II")))
    pg.select_option("#app [data-fin-preset]", "ytd"); pg.wait_for_timeout(900)
    today = pg.evaluate("(() => { const t = new Date(); return t.getFullYear() + String(t.getMonth() + 1).padStart(2, '0') + String(t.getDate()).padStart(2, '0'); })()")
    ok(fin() == ["20250401", min(max(today, "20250401"), "20260331") if today >= "20250401" else "20260331"], "Year to date: %s" % fin())

    print("\nH1: the headings, the figures, the comparative")
    pg.select_option("#app [data-fin-preset]", "h1"); pg.wait_for_timeout(900); tab("bs")
    ok(title() == "Balance sheet as at 30-Sep-2025" and cols() == ["As at 30-Sep-2025", "As at 30-Sep-2024"], "the balance sheet: %s, %s" % (title(), cols()))
    h1 = {"preset": "custom", "from": "20250401", "to": "20250930"}
    want = "{:,.2f}".format(model(h1, "d.bs.total.assets.cy"))
    ok(tot("II").replace(",", "") == want.replace(",", "") and tot("I") == tot("II"), "assets on screen = FinSt's for H1 (%s), equity and liabilities equal" % tot("II"))
    tab("pl")
    ok(title() == "Statement of profit and loss for the period 01-Apr-2025 to 30-Sep-2025" and cols() == ["01-Apr-2025 to 30-Sep-2025", "01-Apr-2024 to 30-Sep-2024"], "the P&L: %s, %s" % (title(), cols()))
    rev = pg.locator('#app table[data-fin-table="pl"] tr[data-fin-line="rev"] td').nth(2).inner_text().strip()
    ok(rev == "61,00,000.00", "revenue for April-September: 61,00,000.00 (20%% + 30%% of 1,22,00,000) (%s)" % rev)
    pat = pg.locator('#app table[data-fin-table="pl"] tr[data-fin-line="pat"] td').first.inner_text().strip()
    ok(pat.endswith("Profit for the period"), "the last line: %s" % pat)
    tab("cf")
    ok(title() == "Cash flow statement for the period 01-Apr-2025 to 30-Sep-2025", "the cash flow: " + title())
    tab("comp")
    ok(pg.locator("#app [data-fin-compnote]").count() == 1 and "full financial year" in pg.inner_text("#app [data-fin-compnote]"), "the computation of income: for a full year only (a line says so)")
    tab("checks")
    ok(pg.locator('#app [data-fin-check]').count() >= 8 and pg.locator('#app [data-fin-check]:not([data-v="0"])').count() == 0, "every check reads 0 (%d)" % pg.locator('#app [data-fin-check]').count())
    pg.evaluate("goClient('dash')"); pg.wait_for_timeout(600); pg.evaluate("goClient('books:fs')"); pg.wait_for_timeout(1500)
    ok(fin() == ["20250401", "20250930"], "back on the page: still H1 (kept for the client)")

    print("\nCompare with: the same period last year, or the last full year")
    tab("bs")
    ok(pg.locator("#app select[data-fin-cmp]").input_value() == "same" and "Comparative: the same period last year (01-Apr-2024 to 30-Sep-2024" in pg.inner_text("#app [data-fin-cmpnote]"), "a part of the year compares with the same period by default, and says so")
    pg.select_option("#app select[data-fin-cmp]", "fy"); pg.wait_for_timeout(900); tab("bs")
    ok(cols() == ["As at 30-Sep-2025", "As at 31-Mar-2025"] and "the last full year (the year ended 31-Mar-2025" in pg.inner_text("#app [data-fin-cmpnote]"), "Last full year: the balance sheet against 31-Mar-2025 (%s)" % cols())
    ok(pg.locator('#app table[data-fin-table="bs"] tr[data-fin-total="II"] td').nth(3).inner_text().strip() == "61,41,000.00", "... 61,41,000.00, last year's closing")
    tab("pl")
    ok(cols() == ["01-Apr-2025 to 30-Sep-2025", "Year ended 31-Mar-2025"] and pg.locator('#app table[data-fin-table="pl"] tr[data-fin-line="rev"] td').nth(3).inner_text().strip() == "1,00,00,000.00", "the P&L against the whole of 2024-25 (%s)" % cols())
    ok(pg.evaluate("S.books.fs.period.cmp") == "fy", "the comparative chosen is kept for the client")
    hy = pg.evaluate("(a) => window.loadFinExport().then(m => FinPdf.html('2025', {unit: 1, dec: 2, period: a}))", {"preset": "h1", "cmp": "fy"})
    ok("Year ended<br>31-Mar-2025" in hy and "the last full year" in hy, "the PDF names the comparative's dates")
    pg.select_option("#app select[data-fin-cmp]", "same"); pg.wait_for_timeout(900)

    print("\nQ2: the earlier part of the year in Note 4")
    pg.select_option("#app [data-fin-preset]", "q2"); pg.wait_for_timeout(900); tab("notes")
    n4 = pg.inner_text('#app section[data-fin-note="4"]')
    ok("Add: profit for 01-Apr-2025 to 30-Jun-2025 (earlier in the year)" in n4 and "Add: profit for the period" in n4 and "Balance at the end of the period" in n4, "Note 4 carries April-June's profit on its own line")

    print("\nCustom")
    pg.select_option("#app [data-fin-preset]", "custom"); pg.wait_for_timeout(900)
    ok(pg.locator("#app input[data-fin-from]").count() == 1 and pg.locator("#app input[data-fin-to]").count() == 1, "Custom shows From and To")
    pg.fill("#app input[data-fin-from]", "2025-08-15"); pg.wait_for_timeout(900)
    pg.fill("#app input[data-fin-to]", "2025-11-20"); pg.wait_for_timeout(900)
    ok(fin() == ["20250815", "20251120"], "15-Aug-2025 to 20-Nov-2025: %s" % fin())
    tab("bs")
    ok(title() == "Balance sheet as at 20-Nov-2025" and tot("I") == tot("II"), "as at 20-Nov-2025, ties (%s)" % tot("II"))
    ok(pg.evaluate("S.books.fs.period.preset") == "custom" and pg.evaluate("S.books.fs.period.from") == "20250815", "the custom dates are kept")

    print("\nA period across 31 March")
    pg.fill("#app input[data-fin-from]", "2025-01-01"); pg.wait_for_timeout(900)
    pg.fill("#app input[data-fin-to]", "2025-06-30"); pg.wait_for_timeout(900)
    tab("bs")
    ok(fin() == ["20250101", "20250630"] and "spans two financial years (2024-25 and 2025-26)" in pg.inner_text("#app [data-fin-spans]") and tot("I") == tot("II"), "01-Jan-2025 to 30-Jun-2025: said to span two years, ties (%s)" % tot("II"))
    tab("notes")
    ok("already in the balance at the beginning of the year" in pg.inner_text('#app section[data-fin-note="4"]'), "Note 4 takes out January-March's profit, already in the opening")

    print("\nStock, typed as on its dates")
    pg.select_option("#app [data-fin-preset]", "h1"); pg.wait_for_timeout(900); tab("tb")
    ok("No stock is known" in pg.inner_text("#app [data-fin-stockpanel]"), "no stock known: nothing assumed")
    box = lambda dt: pg.locator('#app [data-fin-stockpanel] input[aria-label="Stock as on %s"]' % dt)
    box("31-Mar-2025").fill("200000"); box("31-Mar-2025").press("Tab"); pg.wait_for_timeout(1000)
    tab("checks")
    ok(pg.locator('#app tr[data-fin-check="stockMissing"]').count() == 1 and "Closing stock (as on 30-Sep-2025)" in pg.inner_text('#app tr[data-fin-check="stockMissing"]'), "the opening typed, the closing asked for on the Checks tab")
    tab("tb")
    box("30-Sep-2025").fill("230000"); box("30-Sep-2025").press("Tab"); pg.wait_for_timeout(1000)
    src = pg.inner_text('#app [data-fin-stockpanel] tr[data-fin-stock="20250930"] [data-fin-stocksrc]')
    ok(src.startswith("typed by ") and pg.evaluate("S.books.fs.stockAt['20250930'].v") == 230000, "typed, kept per date, with who and when: " + src)
    tab("pl")
    chg = pg.locator('#app table[data-fin-table="pl"] tr[data-fin-line="chg"] td').nth(2).inner_text().strip()
    ok(chg == "(30,000.00)", "changes in inventories: 2,00,000 less 2,30,000 = (30,000.00) (%s)" % chg)
    tab("checks")
    ok(pg.locator('#app [data-fin-check]:not([data-v="0"])').count() == 0, "every check reads 0 with the stock")

    print("\nThe Excel and the PDF for a period")
    for nm, per in [("h1", h1), ("h1fy", {"preset": "custom", "from": "20250401", "to": "20250930", "cmp": "fy"}), ("cross", {"preset": "custom", "from": "20250101", "to": "20250630"}), ("q2", {"preset": "custom", "from": "20250701", "to": "20250930"}), ("custom", {"preset": "custom", "from": "20250815", "to": "20251120"})]:
        for unit, dec in ((1, 2), (1000, 0)):
            x = pg.evaluate(MAKE, {"unit": unit, "dec": dec, "period": per})
            path = os.path.join(OUT, "%s-u%d-d%d.xlsx" % (nm, unit, dec)); open(path, "wb").write(base64.b64decode(x["b64"])); made.append((nm, unit, dec, path, x))
    x = made[0][4]
    ok("01-Apr-2025 to 30-Sep-2025" in x["name"], "the file is named for the period: " + x["name"])
    html = pg.evaluate("(a) => window.loadFinExport().then(m => FinPdf.html('2025', {unit: 1, dec: 2, period: a}))", h1)
    txt = re.sub(r"<[^>]+>", " ", html)
    ok("Balance Sheet as at 30-Sep-2025" in txt and "Statement of Profit and Loss for the period 01-Apr-2025 to 30-Sep-2025" in txt and "Cash Flow Statement for the period 01-Apr-2025 to 30-Sep-2025" in txt
       and "Notes forming part of the financial statements for the period 01-Apr-2025 to 30-Sep-2025" in txt, "the PDF: the period in every heading")
    ok('data-fs-page="comp"' not in html and "For the period" in txt, "... no computation page for a part of the year; the columns 'For the period'")

    print("\nLast year's balance sheet, uploaded")
    pg.evaluate(SETUP, BOOKS25); pg.wait_for_timeout(2500)
    pg.select_option("#app [data-fin-preset]", "fy"); pg.wait_for_timeout(600)
    tab("tb")
    ok(pg.locator("#app [data-fin-pysrc]").get_attribute("data-fin-pysrc") == "opening", "before: PY from the opening balances")
    pg.click('#app button[data-fin-codebtn="Security Deposit - Office Rent"]'); pg.wait_for_timeout(300)
    pg.click('#app [data-fin-pick] button[data-fin-code="38D"]'); pg.wait_for_timeout(1200)
    ok(pg.locator('#app tr[data-key="Security Deposit - Office Rent"] [data-fin-srctag="hand"]').count() == 1, "a hand choice first: the deposit 38D, by hand")
    xl = os.path.join(OUT, "sample-last-year-bs.xlsx"); sample_xlsx(xl)
    pg.set_input_files("#app input[data-fin-pybs-file]", xl); pg.wait_for_timeout(2500)
    ok(pg.locator("#app [data-fin-pybs-confirm]").count() == 1, "the Excel is read and shown to confirm" + (": " + pg.inner_text("#app [data-fin-pybs-err]") if pg.locator("#app [data-fin-pybs-err]").count() else ""))
    rows = pg.locator('#app table[data-fin-table="pybs"] tr[data-fin-pyline]').evaluate_all("xs => xs.map(x => [x.cells[0].childNodes[0].textContent.trim(), x.dataset.code || '', x.dataset.use || ''])")
    rmap = {r[0]: r for r in rows}
    ok(len([r for r in rows if r[2]]) >= 25, "%d lines in the table, %d with a figure" % (len(rows), len([r for r in rows if r[2]])))
    ok(rmap.get("Loans and advances to related parties", [0, ""])[1] == "12H" and rmap.get("Security Deposit - Office Rent", [0, ""])[1] == "12A" and rmap.get("Bluebell Print Works", [0, ""])[1] == "6A",
       "codes from the lines' words: related parties 12H, the deposit 12A (from its note), Bluebell 6A (under the MSME dues)")
    ok(rmap.get("(c) Short-term loans and advances", [0, 0, "1"])[2] == "0" and rmap.get("Security Deposit - Office Rent", [0, 0, "0"])[2] == "1", "the note's lines are used, the face line they add up to is not")
    ok("(they tie)" in pg.inner_text("#app [data-fin-pybs-sum]"), "the lines used tie: " + pg.inner_text("#app [data-fin-pybs-sum]"))
    pg.click("#app [data-fin-pybs-save]"); pg.wait_for_timeout(1500)
    ok(pg.evaluate("S.books.fs.pyBs['2025'].versions.length") == 1 and "version 1" in pg.inner_text("#app [data-fin-pybs-state]"), "kept with the client's books: version 1 (%s)" % pg.inner_text("#app [data-fin-pybs-state]"))
    ok(pg.locator("#app [data-fin-pysrc]").get_attribute("data-fin-pysrc") == "upload", "the PY column is now last year's balance sheet")
    why = pg.locator('#app [data-fin-why="Advance to Staff"]')
    ok(why.count() == 1 and why.inner_text().strip() == "from last year's balance sheet: Loans and advances to related parties" and pg.locator('#app tr[data-key="Advance to Staff"] [data-fin-srctag="suggested"]').count() == 1,
       "a suggestion with its reason: Advance to Staff -> %s (%s)" % (pg.locator('#app tr[data-key="Advance to Staff"] [data-fin-codebtn] b').inner_text() if why.count() else "", why.inner_text() if why.count() else "none"))
    ok(pg.locator('#app tr[data-key="Advance to Staff"] [data-fin-codebtn] b').inner_text() == "12H", "... coded 12H as last year")
    ok("Bluebell Print Works" in pg.inner_text('#app [data-fin-why="Bluebell Print Works"]'), "the MSME creditor named in the note: " + pg.inner_text('#app [data-fin-why="Bluebell Print Works"]'))
    ok(pg.locator('#app tr[data-key="Security Deposit - Office Rent"] [data-fin-srctag="hand"]').count() == 1 and pg.locator('#app tr[data-key="Security Deposit - Office Rent"] [data-fin-codebtn] b').inner_text() == "38D"
       and pg.locator('#app [data-fin-why="Security Deposit - Office Rent"]').count() == 0, "the hand choice is kept: the deposit stays 38D by hand, no suggestion over it")
    tab("checks")
    head = lambda i: pg.locator('#app table[data-fin-table="pycheck"] tr[data-fin-pyhead="%s"] td' % i)
    ok(pg.locator("#app [data-fin-pycheck] h4").inner_text().strip() == "Last year’s balance sheet vs FinCom’s opening balances", "the Checks tab: last year's balance sheet vs FinCom's opening balances")
    ok(head("tr").nth(3).inner_text().strip() == "5,000.00" and head("cash").nth(3).inner_text().strip() == "(5,000.00)", "receivables +5,000, cash (5,000): the audit entry not in the books")
    ok(head("ltla").nth(3).inner_text().strip() == "(1,50,000.00)" and head("stla").nth(3).inner_text().strip() == "1,50,000.00", "the hand choice against last year: long-term (1,50,000), short-term 1,50,000")
    ok(head("sc").nth(3).inner_text().strip() == "0" and pg.locator('#app table[data-fin-table="pycheck"] td.fin-bad').count() == 4, "the other heads agree; four differ")
    badge = pg.locator('#app nav.fin-tabs [data-fin-tab="checks"]').inner_text()
    ok("4" in badge, "the Checks tab counts them: " + badge.replace("\n", " "))
    tab("bs")
    ok(cols() == ["As at 31-Mar-2026", "As at 31-Mar-2025"] and pg.locator('#app table[data-fin-table="bs"] tr[data-fin-total="I"] td').nth(3).inner_text().strip() == "61,41,000.00", "the balance sheet's PY: last year's audited 61,41,000.00")

    print("\nThe same balance sheet as a PDF (its text), and a scan")
    pdfp = ctx.new_page(); pdfp.set_content(sample_html()); pdf = os.path.join(OUT, "sample-last-year-bs.pdf"); pdfp.pdf(path=pdf, format="A4"); pdfp.close()
    tab("tb")
    pg.set_input_files("#app input[data-fin-pybs-file]", pdf); pg.wait_for_timeout(4000)
    ok(pg.locator("#app [data-fin-pybs-confirm]").count() == 1, "the PDF is read by its text" + (": " + pg.inner_text("#app [data-fin-pybs-err]") if pg.locator("#app [data-fin-pybs-err]").count() else ""))
    prow = pg.locator('#app table[data-fin-table="pybs"] tr[data-fin-pyline]').evaluate_all("xs => xs.map(x => [x.cells[0].childNodes[0].textContent.trim(), x.dataset.code || '', x.dataset.use || ''])")
    ok([r for r in prow if r[1]] == [r for r in rows if r[1]], "... the same lines and codes as the Excel (%d)" % len([r for r in prow if r[1]]))
    amt = pg.locator('#app table[data-fin-table="pybs"] tr[data-fin-pyline] input[aria-label="Amount: Juniper Travels"]').input_value() if pg.locator('#app input[aria-label="Amount: Juniper Travels"]').count() else ""
    ok(amt == "-5000", "(5,000) in brackets read as a negative: " + amt)
    pg.click("#app [data-fin-pybs-save]"); pg.wait_for_timeout(1500)
    ok(pg.evaluate("S.books.fs.pyBs['2025'].versions.map(v => v.v + ':' + v.src).join(',')") == "1:xlsx,2:pdf", "version 2 kept beside version 1 (add-only)")
    scan = ctx.new_page(); scan.set_content("<canvas id=c width=600 height=300></canvas><script>const x=document.getElementById('c').getContext('2d');x.fillRect(20,20,400,8);x.fillRect(20,60,300,8);</script>")
    spdf = os.path.join(OUT, "scan.pdf"); scan.pdf(path=spdf, format="A4"); scan.close()
    pg.set_input_files("#app input[data-fin-pybs-file]", spdf); pg.wait_for_timeout(3000)
    er = pg.inner_text("#app [data-fin-pybs-err]") if pg.locator("#app [data-fin-pybs-err]").count() else ""
    ok("could not read this PDF" in er and "not send it anywhere" in er, "a PDF without text: said plainly: " + er)
    ok(pg.evaluate("S.books.fs.pyBs['2025'].versions.length") == 2, "... and nothing kept")

    print("\nA phone, 390 wide")
    ph = br.new_context(viewport={"width": 390, "height": 844}, device_scale_factor=2, is_mobile=True, has_touch=True); pp = ph.new_page()
    pp.on("pageerror", lambda e: errors.append(str(e)))
    pp.goto("http://localhost:%d/" % PORT); pp.wait_for_timeout(2500); pp.evaluate(OPEN); pp.wait_for_timeout(1000)
    pp.evaluate(SETUP, BOOKS); pp.wait_for_timeout(2500)
    pp.select_option("#app [data-fin-preset]", "custom"); pp.wait_for_timeout(900)
    for k in ("bs", "tb", "checks"):
        pp.evaluate("(k) => { S.fin.tab = k; render(); }", k); pp.wait_for_timeout(700)
        w = pp.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth")
        ok(w <= 1, "%s: no sideways scroll (%d)" % (k, w))
    r = pp.evaluate("(() => { const e = document.querySelector('#app [data-fin-period]').getBoundingClientRect(); return [e.left, e.right]; })()")
    ok(r[0] >= 0 and r[1] <= 390, "the period controls fit (%s)" % r)
    ph.close()
    br.close()

print("\nThe workbooks for a period, recalculated by LibreOffice")
SHEETS = ["Inputs", "Balance Sheet", "Profit-Loss", "CF", "Notes", "PPE", "Computation of Income", "Trial Balance", "JV", "Accounting Ratios"]
for nm, unit, dec, path, x in made:
    wb = openpyxl.load_workbook(path)
    ok(wb.sheetnames == SHEETS, "%s: the listed sheets" % os.path.basename(path))
    d = x["d"]; P = d["period"]
    ok(P["range"] in wb["Profit-Loss"]["C3"].value and wb["Balance Sheet"]["C3"].value == "Balance Sheet as at " + P["asAt"] and P["range"] in wb["CF"]["C3"].value and P["range"] in wb["Notes"]["C3"].value,
       "%s: the period in the headings (%s / %s)" % (os.path.basename(path), wb["Balance Sheet"]["C3"].value, wb["Profit-Loss"]["C3"].value))
    ok("full financial year" in str(wb["Computation of Income"]["C5"].value), "%s: the computation sheet says it is for a full year" % os.path.basename(path))
    if nm == "h1fy": ok(wb["Profit-Loss"]["H5"].value == "Year ended\n31-Mar-2025" and wb["Balance Sheet"]["H6"].value == "As at\n31-Mar-2025", "%s: the comparative's dates in the heads" % os.path.basename(path))
    if nm in ("h1", "h1fy"):
        tbs = wb["Trial Balance"]; f = [tbs.cell(r, 4).value for r in range(5, tbs.max_row + 1) if str(tbs.cell(r, 1).value).startswith("Closing stock as on")]
        ok(f == ["=STOCK_CLOSE"] and wb["Inputs"]["E11"].value == 200000 and wb["Inputs"]["E12"].value == 230000, "%s: the stock rows read the Inputs' stock cells (%s)" % (os.path.basename(path), f))
    bad = [(c["sheet"] + "!" + c["addr"], c["v"], resolve(d, c["src"])) for c in x["cells"] if c["src"] and not same(c["v"], resolve(d, c["src"]))]
    ok(not bad, "%s: the values written are FinSt.build's %s" % (os.path.basename(path), bad[:4]))
rc = recalc([m[3] for m in made])
if rc is None: ok(os.environ.get("FIN_XLSX_NO_SOFFICE") == "1", "LibreOffice (soffice) is installed")
else:
    for nm, unit, dec, path, x in made:
        n = os.path.basename(path); lv = openpyxl.load_workbook(rc[path], data_only=True)
        errs = [(ws.title, c.coordinate, c.value) for ws in lv.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:"))]
        ok(not errs, "%s: no error cell %s" % (n, errs[:4]))
        diff = [(c["sheet"] + "!" + c["addr"], c["f"][:80], c["v"], lv[c["sheet"]][c["addr"]].value) for c in x["cells"] if not same(lv[c["sheet"]][c["addr"]].value, c["v"])]
        ok(not diff, "%s: every formula cell recalculates to FinCom's value (%d cells)%s" % (n, len(x["cells"]), "" if not diff else ": " + "; ".join(map(str, diff[:4]))))
        figs = [c for c in x["cells"] if c["src"]]
        bad = [(c["sheet"] + "!" + c["addr"], lv[c["sheet"]][c["addr"]].value, resolve(x["d"], c["src"])) for c in figs if not same(lv[c["sheet"]][c["addr"]].value, resolve(x["d"], c["src"]))]
        ok(not bad and len(figs) > 150, "%s: the %d figures on screen equal the formulas' values %s" % (n, len(figs), bad[:4]))
        ws = lv["Inputs"]; chk = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "Checks (all should be zero)"]
        vals = [(ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(chk[0] + 1, ws.max_row + 1) if ws.cell(r, 2).value] if chk else []
        # (a period whose stock on its first day was not typed: that check reads as on the screen, the rest 0)
        sm = [c["value"] for c in x["d"]["checks"] if c["key"] == "stockMissing"]
        ok(len(vals) >= 8 and all(v == 0 or (str(k).startswith("Stock not known") and sm and v == sm[0]) for k, v in vals), "%s: the %d checks on Inputs read as on the screen (0, but stock not typed %s) %s" % (n, len(vals), sm, [z for z in vals if z[1] != 0]))
        bs = lv["Balance Sheet"]; tots = [bs.cell(r, 6).value for r in range(1, bs.max_row + 1) if bs.cell(r, 3).value == "TOTAL"]
        ok(len(tots) == 2 and same(tots[0], tots[1]) and same(tots[1], x["d"]["bs"]["total"]["assets"]["cy"]), "%s: the balance sheet ties at %s, as on the screen" % (n, tots))
    shutil.copy(made[0][3], os.path.join(OUT, "sample-h1.xlsx"))
    shutil.rmtree(os.path.join(OUT, "bare"), ignore_errors=True); shutil.rmtree(os.path.join(OUT, "lo-profile"), ignore_errors=True)
errs = [e for e in errors if "supabase" not in e.lower() and "Failed to load" not in e]
ok(not errs, "no page errors" + ("" if not errs else ": " + " | ".join(errs[:4])))
print("\n" + ("%d FAILED" % len(fails) if fails else "all passed")); sys.exit(1 if fails else 0)
