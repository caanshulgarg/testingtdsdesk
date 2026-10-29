"""python3 run_fin_depr_ui.py - the Financial statements page's Fixed assets & deferred tax tab (owner, 29-Oct-2026:
"depreciation as per companies act and income tax act should be there so that deferred tax can easily be calculated ..
there should be an uploader also for itr 3,5,6,7"). On the React build (app/dist-test, or TDSDESK_SITE), with the made-up
books of tests/fixtures/fin-depr (the figures worked out by hand in tests/run_fin_depr.js):
  - the tab starts only when asked; last year's ITR-6 JSON is uploaded and read (the blocks' opening WDV), an ITR-4 is
    refused plainly; the Companies Act register shows by class and asset by asset, the income-tax blocks and the deferred
    tax; the Computation of income's Annexure 1 takes the blocks;
  - ticked into the statements: the P&L's depreciation is the register's, the PPE note its gross block, depreciation and
    net block, profit on sale in Note 14, the deferred tax line on the balance sheet and the P&L; every check reads 0;
  - the Excel has the sheets "FA - Companies Act", "FA - Income tax" and "Deferred tax"; LibreOffice recalculates the
    workbook from its formulas alone and every formula cell equals the value FinCom wrote, every FinSt / FinDepr figure
    its formula's value, the checks on Inputs 0; the PDF has the PPE note from the register and the deferred tax note;
  - a phone (390 wide): the page never scrolls sideways.
Screenshots: FS_DEPR_SHOTS (default tests/out/fs-depr/shots)."""
import os, re, sys, json, shutil, zipfile, subprocess, threading, functools, http.server, base64
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
from playwright.sync_api import sync_playwright
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.environ.get("TDSDESK_SITE", os.path.join(HERE, "..", "app", "dist-test"))
OUT = os.path.join(os.environ.get("TDSDESK_OUT", os.path.join(HERE, "out")), "fs-depr")
SHOTS = os.environ.get("FS_DEPR_SHOTS", os.path.join(OUT, "shots"))
shutil.rmtree(OUT, ignore_errors=True); os.makedirs(OUT); os.makedirs(SHOTS, exist_ok=True)
FIX = os.path.join(HERE, "fixtures", "fin-depr")
fails, errors = [], []
def ok(c, w):
    print(("  ok   " if c else "  FAIL ") + w)
    if not c: fails.append(w)
class Q(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(("localhost", 0), functools.partial(Q, directory=SITE)); PORT = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()
BOOKS = json.load(open(os.path.join(FIX, "books.json")))

SETUP = """(F) => {
  S.firm = Object.assign({}, S.firm, {firmName: "Made-up & Associates", frn: "000000S", partner: "A. Sample", mno: "000001", place: "Bengaluru"});
  const c = newCompany({name: "Wren Systems Private Limited (made-up)"}); c.pan = "AABCW1234F"; c.cin = "U72900KA2015PTC000002";
  S.companies[c.id] = c; Store.saveCompany(c); S.coId = c.id; S.view = "company"; S.tab = "books"; S.loadingCo = false;
  const under = Object.assign({}, F.groups), led = {};
  Object.entries(F.open).forEach(([n, v]) => { led[n] = {open: -v, close: 0, parent: under[n]}; });
  const vouchers = F.vouchers.map((v, i) => ({id: "v" + i, date: v[0], no: String(i + 1), type: v[1], narr: v[2], ent: v[3].map(([l, a]) => ({l, a: -a}))}));
  const b = {loading: false, cid: c.id, challans: [], alloc: {}, ledInfo: {}, gstins: {}, pans: {}, states: {}, misCfg: {freq: "off"}, auditCfg: {freq: "off"}, msme: {},
    under, groups: {}, vouchers, meta: {from: "20250401", to: "20260331"}, tb: {from: "20250401", to: "20260331", at: "2026-04-30T10:00:00Z", src: "copy", led}, fs: JSON.parse(JSON.stringify(F.fs))};
  b.map = Books.mapLedgers(b.vouchers, {});
  S.books = b; window.__bk = b; S.fsFy = "2025"; S.fin = null; saveBooks(); goClient('books:fs'); return c.id; }"""
OPEN = "() => new Promise((r) => { const go = () => (typeof doAct === 'function' && typeof signInNeeded === 'function' && typeof S === 'object' && S && S.firm) ? (() => { try { if (signInNeeded()) doAct('useOffline'); } finally { r(); } })() : setTimeout(go, 50); go(); })"
MAKE = r"""async (a) => {
  const m = await window.loadFinExport();
  const x = await m.FinXlsx.make("2025", {unit: a.unit, dec: a.dec, period: a.period});
  const u8 = new Uint8Array(x.buffer); let s = ""; for (let i = 0; i < u8.length; i += 0x8000) s += String.fromCharCode.apply(null, u8.subarray(i, i + 0x8000));
  const d = FinSt.build("2025", {unit: a.unit, dec: a.dec, period: a.period}); d.ratios = FinSt.ratios(d); d.notesByKey = {}; d.notes.forEach(n => { d.notesByKey[n.key] = n; });
  return {b64: btoa(s), cells: x.cells, sheets: x.sheets, name: x.name, d: JSON.parse(JSON.stringify(d))};
}"""
def resolve(d, src):
    if src[0] == "ppeRow":
        _, key, col, i, c = src; return (d["ppe"] if key == "8" else d["intang"])[col]["rows"][i][c]
    if src[0] == "ppeTot":
        _, key, col, c = src; return (d["ppe"] if key == "8" else d["intang"])[col]["total"][c]
    x = d
    for k in src:
        if isinstance(x, dict) and k not in x: return None
        x = x[k]
    return x
def same(a, b):
    # a date-formatted cell reads back as a datetime: its Excel serial
    import datetime as _dt
    a, b = [(z - _dt.datetime(1899, 12, 30)).days if isinstance(z, _dt.datetime) else z for z in (a, b)]
    if isinstance(a, bool) or isinstance(b, bool): return a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)): return abs(a - b) < 1e-6 * max(1, abs(b) * 1e-3)
    if a in (None, "") and b in (None, ""): return True
    if isinstance(a, (int, float)) and b in (None, ""): return a == 0
    if isinstance(b, (int, float)) and a in (None, ""): return b == 0
    return a == b
def strip_values(src, dst):
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.startswith("xl/worksheets/sheet"):
                s = data.decode("utf-8")
                s = re.sub(r'<c ([^>]*?)>(<f>.*?</f>)<v>.*?</v></c>', lambda m: "<c " + re.sub(r'\s*t="[^"]*"', "", m.group(1)) + ">" + m.group(2) + "</c>", s, flags=re.S)
                data = s.encode("utf-8")
            zo.writestr(it, data)
def recalc(files):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice: return None
    bare = os.path.join(OUT, "bare"); outd = os.path.join(OUT, "recalc"); os.makedirs(bare, exist_ok=True); os.makedirs(outd, exist_ok=True)
    names = []
    for f in files:
        b = os.path.join(bare, os.path.basename(f)); strip_values(f, b); names.append(b)
    subprocess.run([soffice, "-env:UserInstallation=file://" + os.path.join(OUT, "lo-profile"), "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outd] + names,
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    return {f: os.path.join(outd, os.path.basename(f)) for f in files}
def num(t):
    t = (t or "").strip()
    if t in ("", "-"): return 0.0
    neg = t.startswith("(")
    v = float(re.sub(r"[^\d.]", "", t) or 0)
    return -v if neg else v

made = []
with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(viewport={"width": 1440, "height": 900}, color_scheme="light", accept_downloads=True); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    pg.goto("http://localhost:%d/" % PORT); pg.wait_for_timeout(2500); pg.evaluate(OPEN); pg.wait_for_timeout(1000)
    pg.evaluate(SETUP, BOOKS); pg.wait_for_timeout(2500)
    app = "#app "
    ok(pg.locator(app + "[data-fin-head]").count() == 1, "the Financial statements page is open")
    d0 = pg.evaluate("(() => { const d = FinSt.build('2025', {unit: 1, dec: 2}); return {dep: d.pl.dep.cy, depr: !!d.depr, bad: d.checks.filter(c => c.value).map(c => c.key)}; })()")
    ok(d0["dep"] == 100000 and not d0["depr"] and not d0["bad"], "before the tab is used: the books' depreciation 1,00,000, nothing of FinDepr, checks 0 %s" % d0)

    print("\nThe tab")
    more = pg.locator(app + "[data-fin-head] ~ nav [data-tabs-more], " + app + "nav.fin-tabs [data-tabs-more]").first
    more.click(); pg.wait_for_timeout(300)
    pg.locator("[data-fin-tab=depr]").first.click(); pg.wait_for_timeout(1200)
    ok(pg.locator(app + "[data-fd-off]").count() == 1, "the tab opens from More with a word on what it does, nothing worked out yet")
    pg.locator(app + "[data-fd-start]").click(); pg.wait_for_timeout(1200)
    ok(pg.locator(app + "[data-fd-itr]").count() == 1 and pg.locator(app + "[data-fd-class=comp]").count() == 1, "started: the ITR uploader and the register")
    pg.locator(app + "[data-fd-itr-file]").set_input_files(os.path.join(FIX, "itr4-ay2025.json")); pg.wait_for_timeout(1200)
    ok("ITR-4 has no depreciation schedules" in (pg.locator(app + "[data-fd-itr-err]").inner_text() if pg.locator(app + "[data-fd-itr-err]").count() else ""), "an ITR-4: refused plainly")
    pg.locator(app + "[data-fd-itr-file]").set_input_files(os.path.join(FIX, "itr6-ay2025.json")); pg.wait_for_timeout(1500)
    st = pg.locator(app + "[data-fd-itr-state]").inner_text()
    ok("ITR-6" in st and "AY 2025-26" in st and "AABCW1234F" in st and "version 1" in st, "ITR-6 JSON uploaded and kept: " + st)
    blk = lambda k, c: num(pg.locator(app + "[data-fd-block=%s] [data-fd-v=%s]" % (k, c)).inner_text())
    ok(pg.locator(app + "[data-fd-block=c40] [data-fd-itsrc=itr]").count() == 1 and blk("c40", "dep") == 138000 and blk("c40", "close") == 252000, "the 40% block: opening from the ITR, depreciation 1,38,000, closing 2,52,000")
    ok(blk("pm15", "dep") == 0 and blk("pm15", "stcg") == -20000 and pg.locator(app + "[data-fd-ceased=pm15]").is_checked(), "the car's block ceases: no depreciation, STCL (20,000)")
    pg.locator(app + "[data-fd-method=furn]").select_option("SLM"); pg.wait_for_timeout(1000)
    ca = lambda k, c: num(pg.locator(app + "[data-fd-class=%s] [data-fd-v=%s]" % (k, c)).inner_text())
    ok(ca("comp", "dep") == 199152.00 and ca("furn", "dep") == 30868.49 and ca("veh", "dep") == 87798.61, "the register: computers 1,99,152.00; furniture (SLM) 30,868.49; car 87,798.61")
    ok(ca("comp", "gO") == 500000 and ca("comp", "dO") == 300000 and "last year's balance sheet" in pg.locator(app + "[data-fd-class=comp] [data-fd-opensrc]").inner_text(), "the opening from last year's audited balance sheet")
    ok(pg.locator(app + "[data-fd-item]").count() == 6 and num(pg.locator(app + "[data-fd-gain]").inner_text()) == 17798.61, "asset by asset (6: three openings, three additions; the car's opening sold), profit on the car's sale 17,798.61")
    ok("49,838.00" in pg.locator(app + "[data-fd-dt=closing]").inner_text() and "asset" in pg.locator(app + "[data-fd-dt=closing]").inner_text(), "deferred tax asset 49,838")
    ok("115JB(5A)" in pg.locator(app + "[data-fd-mat]").inner_text(), "MAT: not relevant under 115BAA, said")
    pg.screenshot(path=os.path.join(SHOTS, "1-tab-register.png"), full_page=True)
    pg.locator(app + "[data-fd-blockof=comp]").scroll_into_view_if_needed()
    pg.locator(app + "[data-fin-table=fd-it]").screenshot(path=os.path.join(SHOTS, "2-income-tax-blocks.png"))
    pg.locator(app + "[data-fin-table=fd-dt]").screenshot(path=os.path.join(SHOTS, "3-deferred-tax.png"))

    print("\nThe computation takes the blocks")
    pg.locator(app + "[data-fin-tab=comp]").click(); pg.wait_for_timeout(1200)
    ok(pg.locator(app + "[data-fin-table=itdep] [data-fin-block=c40]").count() == 1 and pg.locator(app + "[data-fin-table=itdep] [data-fin-block=s20]").count() == 0, "Annexure 1: the blocks with figures")
    ok("1,67,000.00" in pg.locator(app + "[data-fin-table=itdep]").inner_text(), "Annexure 1: section 32 depreciation 1,67,000")

    print("\nInto the statements")
    more.click(); pg.wait_for_timeout(300); pg.locator("[data-fin-tab=depr]").first.click(); pg.wait_for_timeout(1000)
    pg.locator(app + "[data-fd-use=ca]").check(); pg.wait_for_timeout(1000); pg.locator(app + "[data-fd-use=dt]").check(); pg.wait_for_timeout(1200)
    pg.locator(app + "[data-fin-tab=pl]").click(); pg.wait_for_timeout(1200)
    plt = pg.locator(app + "[data-fin-table=pl] [data-fin-line=dep]").inner_text()
    ok("3,17,819.10" in plt, "P&L depreciation: the register's 3,17,819.10 (" + plt.replace("\n", " ") + ")")
    ok("(61,838.00)" in pg.locator(app + "[data-fin-table=pl] [data-fin-line=def]").inner_text(), "P&L deferred tax: a credit of 61,838")
    pg.screenshot(path=os.path.join(SHOTS, "4-profit-loss.png"), full_page=True)
    pg.locator(app + "[data-fin-tab=bs]").click(); pg.wait_for_timeout(1000)
    ok("49,838.00" in pg.locator(app + "[data-fin-table=bs] [data-fin-line=dta]").inner_text(), "balance sheet: deferred tax assets (net) 49,838")
    pg.locator(app + "[data-fin-tab=notes]").click(); pg.wait_for_timeout(1000)
    ppe = pg.locator(app + "[data-fin-ppe=comp]").first.inner_text()
    ok(all(s in ppe for s in ("5,00,000.00", "2,10,000.00", "7,10,000.00", "1,99,152.00", "4,99,152.00", "2,10,848.00")), "the PPE note, computers: gross 7,10,000, depreciation 4,99,152, net 2,10,848")
    pg.locator(app + "[data-fin-ppe=comp]").first.scroll_into_view_if_needed(); pg.screenshot(path=os.path.join(SHOTS, "5-ppe-note.png"))
    dd = pg.evaluate("(() => { const d = FinSt.build('2025', {unit: 1, dec: 2}); return {bad: d.checks.filter(c => c.value).map(c => c.key + ' ' + c.value), gain: (d.notes.find(n => n.key === '14').lines.find(l => l.code === '14B') || {}).cy}; })()")
    ok(pg.locator(app + "nav.fin-tabs [data-fin-tab=checks] .sbar-n").count() == 0, "the Checks tab has no red count")
    ok(not dd["bad"] and dd["gain"] == 17798.61, "every check reads 0; Note 14 profit on sale 17,798.61 %s" % dd)

    print("\nThe Excel")
    for unit, dec in ((1, 2), (1000, 0)):
        x = pg.evaluate(MAKE, {"unit": unit, "dec": dec})
        path = os.path.join(OUT, "depr-u%d-d%d.xlsx" % (unit, dec)); open(path, "wb").write(base64.b64decode(x["b64"])); made.append((unit, dec, path, x))
    ok(all(s in made[0][3]["sheets"] for s in ("FA - Companies Act", "FA - Income tax", "Deferred tax")), "the workbook has FA - Companies Act, FA - Income tax and Deferred tax: %s" % made[0][3]["sheets"])
    print("\nThe PDF")
    h = pg.evaluate("window.loadFinExport().then(m => m.FinPdf.html('2025', {unit: 1, dec: 2}))")
    ok("data-fs-ppereg" in h and "Schedule II" in h, "the PDF's PPE note says it is the register's (Schedule II)")
    ok("data-fs-dtnote" in h and "49,838.00" in h and "61,838.00" in h, "the PDF's deferred tax note: components, 49,838, the credit of 61,838")
    with ctx.expect_page(timeout=30000) as pop:
        pg.evaluate("FinPdf.open('2025', {unit: 1, dec: 2})")
    w = pop.value; w.wait_for_load_state(); w.wait_for_timeout(600)
    w.locator("[data-fs-dtnote]").scroll_into_view_if_needed(); w.screenshot(path=os.path.join(SHOTS, "6-pdf-deferred-tax-note.png"))
    w.pdf(path=os.path.join(OUT, "depr-sample.pdf"), format="A4", prefer_css_page_size=True, print_background=True); w.close()

    print("\nA phone")
    ph = br.new_context(viewport={"width": 390, "height": 844}, color_scheme="light"); pp = ph.new_page(); pp.on("pageerror", lambda e: errors.append(str(e)))
    pp.goto("http://localhost:%d/" % PORT); pp.wait_for_timeout(2500); pp.evaluate(OPEN); pp.wait_for_timeout(800); pp.evaluate(SETUP, BOOKS); pp.wait_for_timeout(2000)
    pp.evaluate("() => { FinDepr.start(); S.fin = Object.assign({}, S.fin, {tab: 'depr'}); render(); }"); pp.wait_for_timeout(1500)
    sw = pp.evaluate("document.documentElement.scrollWidth")
    ok(pp.locator("#app [data-fin-depr]").count() == 1 and sw <= 392, "390 wide: the tab shows, the page does not scroll sideways (%d)" % sw)
    pp.screenshot(path=os.path.join(SHOTS, "7-phone.png"), full_page=False); ph.close()
    br.close()

print("\nRecalculated by LibreOffice from the formulas alone")
rc = recalc([m[2] for m in made])
if rc is None: ok(False, "LibreOffice (soffice) is installed")
else:
    for unit, dec, path, x in made:
        nm = os.path.basename(path); lv = openpyxl.load_workbook(rc[path], data_only=True)
        errs = [(ws.title, c.coordinate, c.value) for ws in lv.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:"))]
        ok(not errs, "%s: no #REF!, #NAME? or other error %s" % (nm, errs[:5]))
        diff = [(c["sheet"] + "!" + c["addr"], c["f"][:90], c["v"], lv[c["sheet"]][c["addr"]].value) for c in x["cells"] if c["v"] is not None and not same(lv[c["sheet"]][c["addr"]].value, c["v"])]
        ok(not diff, "%s: every formula cell recalculates to the value FinCom wrote (%d cells)%s" % (nm, len(x["cells"]), "" if not diff else ": " + "; ".join(map(str, diff[:6]))))
        mine = [c for c in x["cells"] if c["sheet"] in ("FA - Companies Act", "FA - Income tax", "Deferred tax")]
        figs = [c for c in x["cells"] if c["src"]]
        bad = [(c["sheet"] + "!" + c["addr"], lv[c["sheet"]][c["addr"]].value, resolve(x["d"], c["src"])) for c in figs if not same(lv[c["sheet"]][c["addr"]].value, resolve(x["d"], c["src"]))]
        ok(not bad and len(mine) > 100, "%s: %d formula cells on the three new sheets; %d figures of the screen equal their formulas' values %s" % (nm, len(mine), len(figs), bad[:4]))
        tbl = [c for c in x["cells"] if c["sheet"] == "Trial Balance" and ("FA - Companies Act" in c["f"] or "'Deferred tax'" in c["f"])]
        ok(len(tbl) >= 4, "%s: the Trial Balance's FinCom rows are linked to the new sheets (%d)" % (nm, len(tbl)))
        ws = lv["Inputs"]; chk = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "Checks (all should be zero)"]
        vals = [(ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(chk[0] + 1, ws.max_row + 1) if ws.cell(r, 2).value] if chk else []
        ok(len(vals) >= 9 and all(v == 0 for _, v in vals), "%s: the %d checks on Inputs all read 0 %s" % (nm, len(vals), [z for z in vals if z[1] != 0]))
        dts = lv["Deferred tax"]
        ok(same(dts["C15"].value, x["d"]["depr"]["dt"]["closing"]) and same(dts["C14"].value, 25.168), "%s: Deferred tax: closing %s at %s%% (the Inputs' regime)" % (nm, dts["C15"].value, dts["C14"].value))
    shutil.copy(made[0][2], os.path.join(SHOTS, "depr-sample.xlsx"))
errs = [e for e in errors if "supabase" not in e.lower() and "Failed to load" not in e]
ok(not errs, "no page errors" + ("" if not errs else ": " + " | ".join(errs[:4])))
print("\n" + ("%d FAILED" % len(fails) if fails else "all passed")); sys.exit(1 if fails else 0)
