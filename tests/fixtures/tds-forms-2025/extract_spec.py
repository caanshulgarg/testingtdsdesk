"""Extract Protean's file-format workbook (Form 138 / 140 / 143 / 144, version 1.1, dated 22-07-2026) to JSON, as it is.

    python3 -I extract_spec.py <workbook.xlsx> <form> > spec-<form>.json

Read-only (openpyxl, data_only). Nothing is interpreted: every record table (FH, BH, CD, DD) is kept field by field with
the workbook's own words and its sheet row number, and every annexure as its rows. The workbooks themselves are the
owner's upload of 17-Oct-2026 and are not committed; the JSON is what the tests check the files against."""
import json, re, sys
import openpyxl

def cell(v):
    if v is None: return ""
    return str(v).replace("\r", "").strip()

def main(path, form):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = [w for w in wb.worksheets if w.max_row > 5][0]
    rows = [(i, [cell(v) for v in r]) for i, r in enumerate(ws.iter_rows(values_only=True), 1)]
    out = {"form": form, "workbook": path.split("/")[-1].split("-", 1)[-1], "sheet": ws.title, "title": "", "notes": [],
           "records": {}, "annexures": []}
    rec, ann = None, None
    for i, r in rows:
        t = [x for x in r if x]
        if not t: continue
        line = " | ".join(t)
        if i <= 3 and not out["title"] and "File Format for" in line: out["title"] = line
        m = re.search(r"\((File Header|Batch Header|Challan Detail|Deductee Detail|Collectee Detail) Record\)", line)
        if m:
            rec = {"File Header": "FH", "Batch Header": "BH", "Challan Detail": "CD"}.get(m.group(1), "DD")
            out["records"][rec] = {"heading": line, "row": i, "fields": []}; ann = None; continue
        if re.match(r"^Annexure", t[0], re.I) and len(t) <= 2:
            ann = {"name": t[0], "row": i, "rows": []}; out["annexures"].append(ann); rec = None; continue
        if rec and r[0].isdigit() and len(r) > 6:
            out["records"][rec]["fields"].append({"no": int(r[0]), "name": r[1], "col": r[2], "type": r[3], "size": r[4],
                                                  "mo": r[5], "remarks": r[6], "row": i})
            continue
        if rec and r[0] in ("Sr. No.",): continue
        if ann is not None: ann["rows"].append({"row": i, "cells": t}); continue
        if not out["records"] and r[0].isdigit() and len(t) >= 2: out["notes"].append({"no": r[0], "text": t[1], "row": i})
        elif rec: out["records"][rec].setdefault("after", []).append({"row": i, "text": line})
    json.dump(out, sys.stdout, indent=1, ensure_ascii=False)

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
