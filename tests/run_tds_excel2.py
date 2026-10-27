"""python3 run_tds_excel2.py - every Excel of a TDS return is the smart workbook, and the workbook is "dynamic"
(next-tds-excel2, 24-Oct-2026; the owner: "TDS return Excel download is very basic.. it should be very dynamic"). On the
React build (app/dist-test or TDSDESK_SITE) with the fixture books a year on (tdsgst_snapshot.LOAD, CHALLANS):
  1. the return's tab row: Download Excel and Upload Excel at its right end (next-excel-buttons, 27-Oct-2026), not in the
     header or ⋯; at 390 still on screen
  2. every step of 26Q Q4 2025-26 (Deductees, Deductions by challan and as one table, Validate & file, …): the one
     Download Excel gives the return's workbook (About, Deductor, the upload's sheets), opened on the step's sheet; no
     second Excel in a step, no "working" or "Excel to correct" left
  3. 24Q Q1 2026-27 (the salary sheet): the same workbook, on Deductions and on Validate & file's Checks
  4. the workbook: frozen headers (and first columns), an Excel Table or a filter on every list, Summary's lines and
     Checks' rows link to their rows, red and amber conditional formatting, columns as wide as their data, "How to use"
     on top of About; LibreOffice works it out (Summary's Checks count) and renders Deductions, Challans, Summary, Checks
  5. All challans' Excel: an Excel Table with filters, totals by SUBTOTAL over the table, Total and Balance by formula,
     "Used in 26Q Q4" / "Used in all returns" / "Balance", red below nil; LibreOffice: the totals equal the page's foot,
     each Balance the page's Left (one challan made over-used: its Balance below nil)
Screenshots (1440) to TDSDESK_SHOTS (default tests/out/tdsx2-shots)."""
import os, sys, json, threading, functools, http.server, base64, shutil, subprocess, tempfile, glob
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from books_data import CACHE
from tdsgst_snapshot import LOAD, CHALLANS
from playwright.sync_api import sync_playwright
import openpyxl
SITE = os.environ.get("TDSDESK_SITE", os.path.join(HERE, "..", "app", "dist-test"))
PORT = int(os.environ.get("TDSX2_PORT") or 8479)
SHOTS = os.environ.get("TDSDESK_SHOTS") or os.path.join(HERE, "out", "tdsx2-shots")
WORK = tempfile.mkdtemp(prefix="tdsx2-")
os.makedirs(SHOTS, exist_ok=True)
fails, errors = [], []
def ok(c, w):
    print(("  ok   " if c else "  FAIL ") + w)
    if not c: fails.append(w)
class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(("localhost", PORT), functools.partial(Quiet, directory=SITE)); threading.Thread(target=srv.serve_forever, daemon=True).start()
books = json.load(open(CACHE))
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
CATCH = """() => { window.__files = []; window.saveFile = async (name, data) => { const a = new Uint8Array(await data.arrayBuffer()); let s = ''; for (let i = 0; i < a.length; i += 8192) s += String.fromCharCode.apply(null, a.subarray(i, i + 8192)); window.__files.push({name, b64: btoa(s)}); return true; }; }"""
DEDUCTOR = """() => { const co = CO(); co.tds26q = Object.assign({}, co.tds26q, {flat: '12 First Floor', road: 'MG Road', area: 'Karol Bagh', town: 'New Delhi', state: '07', pin: '110005', email: 'accounts@zz.example', phone: '9811111111',
  person: 'Ravi Kumar', personDesignation: 'Partner', personPan: 'ABCPK1234K', personFlat: '12 First Floor', personState: '07', personPin: '110005', personEmail: 'ravi@zz.example', personPhone: '9822222222', branch: 'Head office'}); co.pan = 'AAGCL4827M'; }"""
SMART = ["Deductor", "Challans", "Deductions", "Deductees", "Summary", "Checks", "About"]
EXCELISH = "/excel|working|export|xlsx/i"
NOT_DOWN = "/upload|import|paste/i"

def lo(path, fmt, sub):
    out = os.path.join(WORK, sub); os.makedirs(out, exist_ok=True)
    prof = "file://" + tempfile.mkdtemp(prefix="tdsx2-lo-", dir=WORK)
    subprocess.run([SOFFICE, "-env:UserInstallation=" + prof, "--headless", "--calc", "--convert-to", fmt, "--outdir", out, path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return out

def pdf_png(pdf, png):
    """the PDF's first page as a picture (pypdfium2), for the screenshot of LibreOffice's render"""
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return False
    from PIL import ImageChops, Image
    doc = pdfium.PdfDocument(pdf); im = doc[0].render(scale=3.2).to_pil().convert("RGB"); doc.close()
    box = ImageChops.difference(im, Image.new("RGB", im.size, "white")).getbbox()
    if box: im = im.crop((max(0, box[0] - 24), max(0, box[1] - 24), min(im.width, box[2] + 24), min(im.height, box[3] + 24)))
    if im.width > 1440: im = im.resize((1440, int(im.height * 1440 / im.width)))
    im.save(png)
    return True

with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(viewport={"width": 1440, "height": 950}); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    E, W = pg.evaluate, pg.wait_for_timeout
    pg.goto("http://localhost:%d/" % PORT); W(2500); pg.click('button[data-act="useOffline"]'); W(2000)
    E("window.__dueToday = '2026-10-24'")
    E(LOAD, books); W(1000); E("S.books = window.__bk; render();"); W(500); E(CHALLANS); W(300); E(DEDUCTOR)
    cid = E("S.coId")
    def go(h, wait=1500):
        E("(h) => { location.hash = h; }", "#/c/%s/books/tds/%s" % (cid, h)); W(wait)
    def menu_items():
        pg.click("#app [data-return-card] details.more > summary"); W(250)
        it = E("[...document.querySelectorAll('#app [data-return-card] .more-item')].map(b => b.textContent)")
        pg.keyboard.press("Escape"); W(150); E("document.querySelectorAll('#app details.more[open]').forEach(d => d.open = false)")
        return it
    def caught(click, wait=2500):
        """the file a click downloads (saveFile caught): (name, workbook) or (None, None)"""
        E(CATCH); click(); W(wait)
        f = E("window.__files[window.__files.length - 1] || null")
        if not f: return None, None
        path = os.path.join(WORK, "f%d-%s" % (len(os.listdir(WORK)), f["name"])); open(path, "wb").write(base64.b64decode(f["b64"]))
        return path, openpyxl.load_workbook(path)
    def smart(wb): return wb is not None and all(n in wb.sheetnames for n in SMART)
    def step_exports():
        """every Excel-ish button and ⋯ item in the step (not the header): [(kind, index, text)]"""
        return E("""([rx, no]) => { const R = new RegExp(rx.slice(1, rx.lastIndexOf('/')), 'i'), N = new RegExp(no.slice(1, no.lastIndexOf('/')), 'i'), out = [];
          const main = document.querySelector('#app .rt-main'); if (!main) return out;
          [...main.querySelectorAll('button')].filter(b => !b.closest('details.more') && b.offsetParent && R.test(b.textContent) && !N.test(b.textContent)).forEach((b, i) => out.push(['button', [...main.querySelectorAll('button')].indexOf(b), b.textContent.trim()]));
          [...main.querySelectorAll('details.more')].forEach((d, i) => d.querySelectorAll('.more-item').forEach(b => { if (R.test(b.textContent) && !N.test(b.textContent)) out.push(['more', i, b.textContent.trim()]); }));
          return out; }""", [EXCELISH, NOT_DOWN])
    def click_export(kind, i, text):
        if kind == "button": return lambda: E("(i) => document.querySelectorAll('#app .rt-main button')[i].click()", i)
        def f():
            d = pg.locator("#app .rt-main details.more").nth(i); d.locator("summary").click(); W(200)
            d.locator(".more-item", has_text=text).first.click()
        return f
    def words():
        return E("() => [...document.querySelectorAll('#app .rt button, #app .rt .more-item')].map(b => b.textContent)")

    # next-excel-buttons (27-Oct-2026; the owner: "download and upload excel button should be proper"): the return's one
    # Download Excel and Upload Excel are at the right end of the steps' tab row, on every step and width; the header, ⋯
    # and the steps' own ⋯ and buttons no longer repeat them
    TAB = "#app [data-ret-tabrow] [data-tdsx-btn=%s]"
    print("1. the return's tab row: Download Excel and Upload Excel")
    FY, Q = "2025-26", "Q4"
    go("%s/%s/26Q/deductions" % (FY, Q))
    vis = E("() => ['down', 'up'].map(k => { const b = document.querySelector('#app [data-ret-tabrow] [data-tdsx-btn=' + k + ']'); return b ? [b.textContent, !!b.offsetParent] : null; })")
    ok(vis == [["Download Excel", True], ["Upload Excel", True]], "1440: Download Excel and Upload Excel in the tab row (%s)" % vis)
    ok(E("document.querySelectorAll('#app [data-return-card] [data-tdsx-btn]').length") == 0, "not in the header")
    it = menu_items()
    ok("Download Excel" not in it and "Upload Excel" not in it and not any("working" in x for x in it), "not in ⋯; no plain 'working' Excel (%s)" % it)
    E("() => document.activeElement && document.activeElement.blur()"); pg.mouse.move(5, 900); W(200)
    pg.locator("#app [data-return-card]").screenshot(path=os.path.join(SHOTS, "1-return-header-1440.png"))
    pg.screenshot(path=os.path.join(SHOTS, "1-return-page-1440.png"))
    pg.set_viewport_size({"width": 390, "height": 844}); W(600)
    vis = E("() => ['down', 'up'].map(k => { const b = document.querySelector('#app [data-ret-tabrow] [data-tdsx-btn=' + k + ']'); if (!b || !b.offsetParent) return false; const r = b.getBoundingClientRect(); return r.left >= 0 && r.right <= innerWidth + 1; })")
    ok(vis == [True, True], "390: still on screen (%s)" % vis)
    sw = E("document.documentElement.scrollWidth")
    ok(sw <= 391, "390: no sideways scroll (%s)" % sw)
    pg.screenshot(path=os.path.join(SHOTS, "1-return-header-390.png"))
    pg.set_viewport_size({"width": 1440, "height": 950}); W(500)
    path, wb = caught(lambda: pg.click(TAB % "down"))
    ok(smart(wb) and wb.active.title == "Deductions", "Download Excel on Deductions: the workbook, opened on Deductions (%s)" % (wb and wb.active.title))
    WB26 = path

    print("2. every step of 26Q Q4 2025-26: the one Download Excel, opened on the step's sheet")
    WANT = {"deductor": "Deductor", "deductees": "Deductees", "challans": "Challans", "deductions": "Deductions", "file": "Checks"}
    for step, sheet in WANT.items():
        views = [("", "")] if step != "deductions" else [("by challan", ""), ("one table", "table")]
        for vname, view in views:
            go("%s/%s/26Q/%s" % (FY, Q, step)); E("(v) => { S.tdsDedView = v; render(); }", view); W(600)
            _, wb = caught(lambda: pg.click(TAB % "down"))
            ok(smart(wb) and wb.active.title == sheet, "%s%s: Download Excel opens on %s (%s)" % (step, " (" + vname + ")" if vname else "", sheet, wb and wb.active.title))
            ex = step_exports()
            ok(not ex, "%s%s: no second Excel in the step (%s)" % (step, " (" + vname + ")" if vname else "", ex))
            w = [x for x in words() if "working" in x.lower() or "Excel to correct" in x or x.strip() == "Excel"]
            ok(not w, "%s%s: no plain Excel left (%s)" % (step, " (" + vname + ")" if vname else "", w))

    print("3. 24Q Q1 2026-27 (the salary sheet)")
    go("2026-27/Q1/24Q/deductions"); E("() => { S.tdsDedView = 'table'; render(); }"); W(700)
    ok(not step_exports(), "the Employees' step: no second Excel (%s)" % step_exports())
    _, wb = caught(lambda: pg.click(TAB % "down"))
    ok(smart(wb) and wb.active.title == "Deductions" and any(wb["About"]["A%d" % r].value == "Form" and wb["About"]["B%d" % r].value == "24Q" for r in range(1, 30)), "24Q Employees: Download Excel, the workbook on Deductions")
    go("2026-27/Q1/24Q/file"); W(400)
    ok(pg.locator("#app .rt-main [data-file-excel]").count() == 0, "24Q Validate & file: no second Download Excel")
    _, wb = caught(lambda: pg.click(TAB % "down"))
    ok(smart(wb) and wb.active.title == "Checks", "and the tab row's opens on Checks (%s)" % (wb and wb.active.title))
    ok(not [x for x in words() if "working" in x.lower()], "24Q: no 'working' Excel left")

    print("4. the workbook itself")
    wb = openpyxl.load_workbook(WB26)
    fz = {n: wb[n].freeze_panes for n in ("Deductor", "Challans", "Deductions", "Deductees", "Summary", "Checks")}
    ok(fz == {"Deductor": "A3", "Challans": "E3", "Deductions": "E3", "Deductees": "C3", "Summary": "A3", "Checks": "C3"}, "frozen: the headers, and the first columns of the lists (%s)" % fz)
    ok(all(wb[n].tables for n in ("Challans", "Deductions", "Deductees")) and (wb["Checks"].auto_filter.ref or "").startswith("A2:F"), "a filter on every list: Challans, Deductions, Deductees Excel Tables; Checks filtered (%s)" % wb["Checks"].auto_filter.ref)
    su, ck, wd = wb["Summary"], wb["Checks"], wb["Deductions"]
    lk = {su["A%d" % r].value: (su["B%d" % r].hyperlink.location if su["B%d" % r].hyperlink else None) for r in range(3, 11)}
    ok(lk.get("tds", "").startswith("#'Deductions'!") and lk.get("over", "").startswith("#'Challans'!") and all(lk.values()), "Summary's lines link to their sheets (%s)" % lk)
    nd = sum(1 for r in range(3, 400) if wd["A%d" % r].value)
    ok(ck["A3"].hyperlink and ck["A3"].hyperlink.location == "#'Deductions'!D3" and ck["A%d" % (2 + nd)].hyperlink.location == "#'Deductions'!D%d" % (2 + nd) and ck["B1"].hyperlink
       and "HYPERLINK(\"#'Deductions'!D%d\"" % (3 + nd) in str(ck["A%d" % (3 + nd)].value), "Checks' rows link to the deduction (A3 → Deductions!D3; a spare row once it is filled, by HYPERLINK)")
    def fills(ws):
        out = set()
        for rng in ws.conditional_formatting:
            for r in rng.rules:
                c = r.dxf and r.dxf.fill and (r.dxf.fill.bgColor.rgb if r.dxf.fill.bgColor is not None else None)
                if c: out.add(str(c)[-6:])
        return out
    ok({"FDE2E2", "FFF2CC"} <= fills(wd) and {"FDE2E2", "FFF2CC"} <= fills(ck) and "FDE2E2" in fills(wb["Challans"]) and {"FDE2E2", "FFF2CC"} <= fills(su), "red and amber: Deductions, Checks, Summary; red on Challans (%s)" % {n: sorted(fills(wb[n])) for n in ("Deductions", "Checks", "Challans", "Summary")})
    parties = [wd["D%d" % r].value for r in range(3, 300) if wd["D%d" % r].value]
    wD, wF = wd.column_dimensions["D"].width, wd.column_dimensions["F"].width
    ok(abs(wD - min(46, max(len(x) for x in parties) + 2)) <= 1 and wF < wD, "columns as wide as their data: Deductee %s (longest name %d), Section %s" % (wD, max(len(x) for x in parties), wF))
    ab = wb["About"]
    ok(ab["A1"].value == "How to use" and "yellow" in (ab["B1"].value or "") and any((ab["A%d" % r].value == "Format") for r in range(1, 30)), "About: How to use on top, the fields under it")
    if SOFFICE:
        out = lo(WB26, "xlsx:Calc MS Excel 2007 XML", "calc")
        v = openpyxl.load_workbook(os.path.join(out, os.path.basename(WB26)), data_only=True)
        row = next(r for r in range(3, 25) if v["Summary"]["B%d" % r].value and str(v["Summary"]["B%d" % r].value).startswith("Rows with something"))
        n = v["Summary"]["C%d" % row].value
        nn = sum(1 for r in range(3, 400) if any(v["Checks"]["%s%d" % (c, r)].value not in (None, "") for c in "CDEF"))
        ok(isinstance(n, (int, float)) and n == nn, "LibreOffice: Summary's count of rows to look at is Checks' (%s = %s)" % (n, nn))
        # the renders: one sheet a page, the step's sheets
        for sheet in ("Deductions", "Challans", "Summary", "Checks"):
            one = os.path.join(WORK, "one-%s.xlsx" % sheet)
            w2 = openpyxl.load_workbook(os.path.join(out, os.path.basename(WB26)))
            for n_ in w2.sheetnames:
                if n_ != sheet: w2[n_].sheet_state = "hidden"
            w2.active = w2.sheetnames.index(sheet)
            ws = w2[sheet]; last = 3
            while last < 400 and any(ws.cell(row=last, column=c).value not in (None, "") for c in (2, 4)): last += 1
            if sheet in ("Deductions", "Challans"): ws.print_area = "B1:%s%d" % ("T" if sheet == "Deductions" else "V", max(last, 6))
            elif sheet == "Checks": ws.print_area = "A1:F%d" % max(last, 6)
            ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
            w2.save(one)
            d = lo(one, "pdf", "pdf-" + sheet)
            pdf = os.path.join(d, "one-%s.pdf" % sheet)
            png = os.path.join(SHOTS, "4-workbook-%s-libreoffice.png" % sheet.lower())
            if pdf_png(pdf, png): ok(os.path.exists(png), "LibreOffice's render of %s: %s" % (sheet, os.path.basename(png)))
    else: print("  (LibreOffice is not installed: no recalculation, no renders)")

    print("5. All challans' Excel")
    go("%s/%s/26Q/challans" % (FY, Q))
    over = E("() => { const c = TDS.challans().find(c => Object.keys(TDSReg.ctx().uses[c.id] || {}).length); if (!c) return null; const b = S.books.challans.find(z => z.id === c.id); b.tax = 1; TDSReg.forget(); render(); return c.id; }")
    pg.click("#app [data-return-card] details.more > summary"); W(200); pg.click("#app [data-return-card] [data-ret-allch]"); W(1200)
    ok(pg.locator("#app [data-allch-excel]").count() == 1, "All challans: its Excel button (Download Excel at the filter row's end)")
    page = E("""() => ({rows: [...document.querySelectorAll('#app [data-allch-row]')].map(r => ({id: r.dataset.allchRow, kind: r.dataset.allchKind, left: r.querySelector('[data-allch-left]') ? +r.querySelector('[data-allch-left]').dataset.allchLeft : null,
      used: r.querySelector('[data-allch-used]') ? +r.querySelector('[data-allch-used]').dataset.allchUsed : null})), tax: document.querySelector('#app [data-allch-total=tax]').textContent, used: document.querySelector('#app [data-allch-total=used]').textContent})""")
    pg.screenshot(path=os.path.join(SHOTS, "5-all-challans-page-1440.png"))
    path, wb = caught(lambda: pg.click("#app [data-allch-excel]"))
    ok(wb is not None and len(wb.sheetnames) == 1, "downloaded: one sheet (%s)" % (wb and wb.sheetnames))
    if wb:
        ws = wb.worksheets[0]
        t = list(ws.tables.values())
        head = [ws.cell(row=3, column=c).value for c in range(1, 14)]
        ok(t and t[0].ref.startswith("A3:M") and t[0].autoFilter is not None, "an Excel Table from the header row, with a filter on each header (%s)" % (t and t[0].ref))
        ok(head[9] == "Used in 26Q Q4" and head[10] == "Used in all returns" and head[11] == "Balance" and head[8] == "Total", "the columns: … Total · Used in 26Q Q4 · Used in all returns · Balance · Status (%s)" % head)
        n = len(page["rows"])
        ok(all(str(ws.cell(row=2, column=c).value).startswith("=SUBTOTAL(109,") for c in (6, 7, 8, 9, 10, 11, 12)), "the totals above the table by SUBTOTAL (they follow the filter)")
        chr_ = [r for r in range(4, 4 + n) if ws["L%d" % r].value is not None]
        ok(chr_ and all(str(ws["L%d" % r].value) == "=N(F%d)-N(K%d)" % (r, r) for r in chr_) and all(str(ws["I%d" % r].value).startswith("=N(F%d)" % r) for r in range(4, 4 + n)), "Total and Balance by formula on each row")
        reds = [r for rng in ws.conditional_formatting for r in rng.rules if r.dxf and r.dxf.fill and str(r.dxf.fill.bgColor.rgb).endswith("FDE2E2") and "<-0.004" in (r.formula or [""])[0]]
        ok(reds, "red for a balance below nil (conditional formatting)")
        ok(ws.freeze_panes == "D4", "the header frozen, and the first columns (%s)" % ws.freeze_panes)
        if SOFFICE:
            out = lo(path, "xlsx:Calc MS Excel 2007 XML", "calc-ach")
            v = openpyxl.load_workbook(os.path.join(out, os.path.basename(path)), data_only=True).worksheets[0]
            money = lambda s: float(str(s).replace("₹", "").replace(",", ""))
            ok(abs((v["F2"].value or 0) - money(page["tax"])) < 0.01 and abs((v["K2"].value or 0) - money(page["used"])) < 0.01, "LibreOffice: Tax and Used in all returns total as the page's foot (%s, %s)" % (v["F2"].value, v["K2"].value))
            bal = [(v["L%d" % (4 + i)].value, r["left"]) for i, r in enumerate(page["rows"]) if r["kind"] == "ch"]
            ok(bal and all(abs((a or 0) - b) < 0.01 for a, b in bal), "each challan's Balance is the page's Left (%s)" % bal[:4])
            ok(over is None or any((a or 0) < -0.005 for a, _ in bal), "the challan made over-used: its Balance below nil (%s)" % [a for a, _ in bal if (a or 0) < 0])
            vv = openpyxl.load_workbook(os.path.join(out, os.path.basename(path)))
            if True:
                w2 = vv; ws2 = w2.worksheets[0]; ws2.page_setup.orientation = "landscape"; ws2.sheet_properties.pageSetUpPr.fitToPage = True; ws2.page_setup.fitToWidth = 1; ws2.page_setup.fitToHeight = 0
                one = os.path.join(WORK, "allch.xlsx"); w2.save(one); d = lo(one, "pdf", "pdf-allch")
                pdf_png(os.path.join(d, "allch.pdf"), os.path.join(SHOTS, "5-all-challans-excel-libreoffice.png"))
    br.close()

shutil.rmtree(WORK, ignore_errors=True)
ok(not errors, "no page errors (%s)" % errors[:3])
srv.shutdown()
print("\n%d failed" % len(fails) if fails else "\nall passed")
sys.exit(1 if fails else 0)
