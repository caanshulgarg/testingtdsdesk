"""python3 run_fin_xlsx.py - the financial statements' Excel and PDF (ROUND 99, stage 2B: app/src/lib/finExport.js).

On the React build (app/dist-test, or TDSDESK_SITE), with made-up books (tests/fixtures/fin-statements: a company shaped
like the owner's audit workbooks, and a firm):
  - FinXlsx.download saves an .xlsx; ExcelJS comes from this site (assets/exceljs.min.js, its checksum as the module
    says), only when a workbook is made, and nothing of it is in legacy.js;
  - the workbook has only the listed sheets (Inputs, Balance Sheet, Profit-Loss, CF, Notes, PPE, Computation of Income,
    Trial Balance, JV, and Accounting Ratios for a company), the names UNIT, DEC, TOL, SHARES, and a cached value on every
    formula cell (a viewer that does not recalculate still sees the figures);
  - LibreOffice (headless) recalculates every workbook from its formulas alone (the cached values taken out first): every
    formula cell's value equals the one FinCom wrote, and each figure of FinSt.build (notes, statements, cash flow, PPE,
    computation, checks, ratios) is the formula's value, at units 1 / 1,000 / 1,00,000 / 1,00,00,000 and 0 / 2 decimals;
    no #REF!, #NAME? or other error; every check on Inputs reads 0;
  - FinPdf.open shows exactly the statements asked for (Balance Sheet, P&L, Cash Flow, Notes, Computation; the trial
    balance only when asked), each on a new page, A4 portrait, with the signature blocks; a sample PDF is saved.
Samples: TDSDESK_OUT (tests/out)/fin-xlsx/sample.xlsx and sample.pdf (FIN_SAMPLES=<dir> copies them there too).
Needs LibreOffice (soffice) for the recalculation (CI installs libreoffice-calc on the share that runs this test)."""
import os, re, sys, json, base64, shutil, hashlib, zipfile, subprocess, threading, functools, http.server, math
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
from playwright.sync_api import sync_playwright
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.environ.get("TDSDESK_SITE", os.path.join(HERE, "..", "app", "dist-test"))
OUT = os.path.join(os.environ.get("TDSDESK_OUT", os.path.join(HERE, "out")), "fin-xlsx")
shutil.rmtree(OUT, ignore_errors=True); os.makedirs(OUT)
FIX = os.path.join(HERE, "fixtures", "fin-statements")
fails, errors = [], []
def ok(c, w):
    print(("  ok   " if c else "  FAIL ") + w)
    if not c: fails.append(w)

class Q(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(("localhost", 0), functools.partial(Q, directory=SITE)); PORT = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()

C = json.load(open(os.path.join(FIX, "company.json"))); F = json.load(open(os.path.join(FIX, "firm.json")))
CO = {"name": "Larkspur Analytics Private Limited", "pan": "AABCL1234F", "cin": "U72900KA2015PTC000001", "address": "12 Made-up Road, Bengaluru"}
FIRMCO = {"name": "Kestrel & Co", "pan": "AAAFK1234K"}
AUDITOR = {"firmName": "Made-up & Associates", "frn": "000000S", "partner": "A. Sample", "mno": "000001", "place": "Bengaluru"}

# the books FinCom would hold (as tests/run_fin_statements.js makes them): Tally's opening balances (a debit negative) and
# one journal a year carrying each ledger's movement, with the asset purchases on their own dates
SETUP = r"""(a) => {
  const MIS_shift = end => String(Number(end.slice(0, 4)) - 1) + end.slice(4);
  function books(F, from, years){
    const under = Object.assign({}, F.groups), led = {};
    Object.entries(F.o).forEach(([n, v]) => { led[n] = {open: -v, close: 0, parent: under[n]}; });
    const vouchers = [], bank = Object.keys(under).find(n => under[n] === "Bank Accounts");
    (F.additions || []).forEach(([l, d, x], i) => vouchers.push({id: "a" + i, date: d, type: "Journal", narr: "Purchase of " + l, ent: [{l, a: -x}, {l: bank, a: x}]}));
    years.forEach(([end, bs, pl, prev]) => {
      const ent = [];
      Object.keys(bs).forEach(n => { if (n === bank) return; const added = (F.additions || []).filter(([l, d]) => l === n && d <= end && d > MIS_shift(end)).reduce((s, z) => s + z[2], 0);
        const mv = Math.round(((bs[n] || 0) - (prev[n] || 0) - added) * 100) / 100; if (mv) ent.push({l: n, a: -mv}); });
      Object.entries(pl).forEach(([n, v]) => { if (v) ent.push({l: n, a: -v}); });
      ent.push({l: bank, a: Math.round(-ent.reduce((s, e) => s + e.a, 0) * 100) / 100});
      vouchers.push({id: "m" + end, date: end, type: "Journal", narr: "Movements for the year (fixture)", ent});
    });
    return {under, groups: {}, ledInfo: {}, vouchers, meta: {from}, tb: {from, to: "20260331", at: "2026-04-30T10:00:00Z", src: "copy", led}, msme: {}, fs: {}, map: {}, challans: [], alloc: {}, loading: false};
  }
  const F = a.fix, co = newCompany({name: a.co.name, pan: a.co.pan}); co.cin = a.co.cin || ""; co.address = a.co.address || "";
  S.companies[co.id] = co; S.data = S.data || {}; S.data[co.id] = {parties: {}, entries: {}, loaded: true}; S.coId = co.id; S.firm = a.auditor;
  const b = a.kind === "co" ? books(F, "20240401", [["20250331", F.p, F.plp, F.o], ["20260331", F.c, F.plc, F.p]]) : books(F, "20250401", [["20260331", F.c, F.plc, F.o]]);
  b.cid = co.id; b.fs = a.fs; if (a.kind === "co") b.msme = {"Bluebell Print Works": "Micro"};
  S.books = b; window.__fx = b;
  return FinSt.build("2025", {unit: 1, dec: 2}).error || "";
}"""
MAKE = r"""async (a) => {
  const m = await window.loadFinExport(); S.books = window.__fx;
  const x = await m.FinXlsx.make("2025", {unit: a.unit, dec: a.dec});
  const u8 = new Uint8Array(x.buffer); let s = ""; for (let i = 0; i < u8.length; i += 0x8000) s += String.fromCharCode.apply(null, u8.subarray(i, i + 0x8000));
  const d = FinSt.build("2025", {unit: a.unit, dec: a.dec}); d.ratios = FinSt.ratios(d); d.notesByKey = {}; d.notes.forEach(n => { d.notesByKey[n.key] = n; });
  return {b64: btoa(s), cells: x.cells, sheets: x.sheets, name: x.name, d: JSON.parse(JSON.stringify(d))};
}"""

def resolve(d, src):
    if src[0] == "ppeRow":
        _, key, col, i, c = src; return (d["ppe"] if key == "8" else d["intang"])[col]["rows"][i][c]
    if src[0] == "ppeTot":
        _, key, col, c = src; return (d["ppe"] if key == "8" else d["intang"])[col]["total"][c]
    x = d
    for k in src: x = x[k]
    return x

def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)): return abs(a - b) < 1e-6 * max(1, abs(b) * 1e-3)
    if a in (None, "") and b in (None, ""): return True
    if isinstance(a, (int, float)) and b in (None, ""): return a == 0
    if isinstance(b, (int, float)) and a in (None, ""): return b == 0
    return a == b

def strip_values(src, dst):
    """A copy with no cached value on any formula cell, so LibreOffice works every formula out itself."""
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.startswith("xl/worksheets/sheet"):
                s = data.decode("utf-8")
                s = re.sub(r'<c ([^>]*?)>(<f>.*?</f>)<v>.*?</v></c>', lambda m: "<c " + re.sub(r'\s*t="[^"]*"', "", m.group(1)) + ">" + m.group(2) + "</c>", s, flags=re.S)
                data = s.encode("utf-8")
            zo.writestr(it, data)

def recalc(files):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice: return None
    bare = os.path.join(OUT, "bare"); outd = os.path.join(OUT, "recalc"); os.makedirs(bare, exist_ok=True); os.makedirs(outd, exist_ok=True)
    names = []
    for f in files:
        b = os.path.join(bare, os.path.basename(f)); strip_values(f, b); names.append(b)
    prof = "file://" + os.path.join(OUT, "lo-profile")
    subprocess.run([soffice, "-env:UserInstallation=" + prof, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outd] + names,
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)
    return {f: os.path.join(outd, os.path.basename(f)) for f in files}

CO_SHEETS = ["Inputs", "Balance Sheet", "Profit-Loss", "CF", "Notes", "PPE", "Computation of Income", "Trial Balance", "JV", "Accounting Ratios"]
NC_SHEETS = CO_SHEETS[:-1]
CO_FS = {"taxDep": {"2024": {"c40": 220000, "pm15": 140000, "f10": 80000}},
         "adj": {"2025": [{"dr": "Income Tax - Current Year", "cr": "Provision for Income Tax FY 2025-26", "amt": 683940, "narr": "Provision for income tax"}]},
         "inputs": {"shares": {"number": 10000, "face": 10, "authorised": 100000, "holders": [{"name": "Maple Holdings Inc", "n": 9999}]},
                    "directors": [{"name": "A. Director", "din": "00000001", "desig": "Director"}, {"name": "B. Director", "din": "00000002", "desig": "Director"}],
                    "place": "Bengaluru", "date": "15-09-2026", "udin": "26000001ABCDEF1234"}}
FM_FS = {"taxDep": {"2025": {"c40": 90000}}, "compEdits": {"2025": {"less_itdep": {"v": 30000, "why": "as per the tax audit report"}}},
         "inputs": {"partners": [{"name": "Asha Rao", "din": "", "desig": "Partner"}, {"name": "Vikram Sen", "din": "", "desig": "Partner"}], "place": "Bengaluru", "date": "15-09-2026"},
         "pyTb": {"2025": {"Consultancy Fees": -2000000, "Salaries": 410000, "Depreciation": 40000, "Profit & Loss A/c": 250000, "Asha Rao Capital": -400000, "Vikram Sen Capital": -300000,
                           "Asha Rao Current A/c": -50000, "Laptop": 150000, "Accumulated Depreciation - Laptop": -50000, "Heron Foods": 300000, "Ibis Supplies": -80000, "GST Payable": -20000, "Kestrel Bank": 1750000}}}
VARIANTS = [("co", u, dc) for u in (1, 1000, 100000, 10000000) for dc in (0, 2)] + [("nc", u, dc) for u in (1, 100000) for dc in (2, 0)]

made = []
with sync_playwright() as p:
    br = p.chromium.launch(); ctx = br.new_context(accept_downloads=True, viewport={"width": 1300, "height": 900}); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    asked = []
    pg.on("request", lambda r: asked.append(r.url))
    pg.goto("http://localhost:%d/" % PORT); pg.wait_for_timeout(2500)
    pg.click('button[data-act="useOffline"]'); pg.wait_for_timeout(1200)

    print("\nLoaded when asked for, from this site")
    ok(pg.evaluate("typeof window.loadFinExport === 'function' && typeof FinXlsx.download === 'function' && typeof FinPdf.open === 'function'"), "FinXlsx.download and FinPdf.open are there from the start (they load the module)")
    ok(pg.evaluate("!window.ExcelJS") and not any("exceljs" in u for u in asked), "ExcelJS is not loaded with the app")
    leg = open(os.path.join(SITE, "legacy.js"), encoding="utf8").read()
    ok("FinXlsx" not in leg and "ExcelJS" not in leg and "pdfHtml" not in leg, "nothing of the Excel / PDF writers is in legacy.js")
    ok("finExport" not in open(os.path.join(SITE, "index.html"), encoding="utf8").read(), "nor preloaded by index.html (a lazy chunk)")
    sha = hashlib.sha256(open(os.path.join(SITE, "assets", "exceljs.min.js"), "rb").read()).hexdigest()
    ok(os.path.exists(os.path.join(SITE, "assets", "exceljs-LICENSE.txt")), "ExcelJS's licence is served beside it")

    err = pg.evaluate(SETUP, {"fix": C, "co": CO, "auditor": AUDITOR, "kind": "co", "fs": CO_FS})
    ok(not err, "the company's statements are worked out" + (": " + err if err else ""))
    ok(pg.evaluate("window.loadFinExport().then(m => m.FinXlsx.sha256)") == sha, "assets/exceljs.min.js is the file the module names (sha256 %s...)" % sha[:16])

    print("\nFinXlsx.download")
    with pg.expect_download(timeout=60000) as dl:
        pg.evaluate("FinXlsx.download('2025', {unit: 1000, dec: 2})")
    sample = os.path.join(OUT, "sample.xlsx"); dl.value.save_as(sample)
    ok(dl.value.suggested_filename.endswith(".xlsx") and "Larkspur" in dl.value.suggested_filename, "saved as " + dl.value.suggested_filename)
    ok(pg.evaluate("!!window.ExcelJS && window.ExcelJS.Workbook != null") and any(u.endswith("/assets/exceljs.min.js") for u in asked), "ExcelJS loaded from assets/exceljs.min.js when the workbook was made")

    print("\nThe PDF")
    for tb in (False, True):
        with ctx.expect_page(timeout=30000) as pop:
            pg.evaluate("(tb) => { S.books = window.__fx; FinPdf.open('2025', {unit: 100000, dec: 2, includeTb: tb}); }", tb)
        w = pop.value; w.wait_for_load_state(); w.wait_for_timeout(600)
        pages = w.eval_on_selector_all("section.fs-page", "xs => xs.map(x => x.dataset.fsPage)")
        want = ["bs", "pl", "cf", "notes", "comp"] + (["tb"] if tb else [])
        ok(pages == want, "%s: the pages are exactly %s (%s)" % ("with the trial balance" if tb else "statements only", want, pages))
        brk = w.eval_on_selector_all("section.fs-page", "xs => xs.map(x => getComputedStyle(x).breakBefore)")
        ok(brk[0] != "page" and all(b == "page" for b in brk[1:]), "each statement starts on a new page (%s)" % brk)
        t = w.inner_text("body")
        ok(all(s in t for s in ("Balance Sheet as at March 31, 2026", "Statement of Profit and Loss for the year ended March 31, 2026", "Cash Flow Statement for the year ended March 31, 2026",
                                 "Notes forming part of the financial statements", "Computation of total income and tax liability", "(All amounts in Rupees lakhs, unless otherwise stated)")), "titles and the unit line on the pages")
        ok(t.count("Chartered Accountants") == 3 and t.count("Membership No. 000001") == 3 and t.count("For and on behalf of the Board of Directors of") == 3 and t.count("DIN: 00000001") == 3,
           "signature blocks on the balance sheet, P&L and cash flow: the auditor and the directors, from the Inputs")
        ok(("Trial balance for the year ended" in t) == tb, "the trial balance %s" % ("included, as asked" if tb else "left out"))
        ok("@page{size:A4 portrait" in w.content().replace(" ", "").replace("@page{size:A4portrait", "@page{size:A4 portrait"), "A4 portrait")
        if not tb:
            w.pdf(path=os.path.join(OUT, "sample.pdf"), format="A4", prefer_css_page_size=True, print_background=True)
            raw = open(os.path.join(OUT, "sample.pdf"), "rb").read()
            n = len(re.findall(rb"/Type\s*/Page[^s]", raw))
            ok(raw[:4] == b"%PDF" and n >= 5, "sample PDF saved (%d pages)" % n)
        w.close()
    for kind, unit, dec in VARIANTS:
        if kind == "nc" and not any(m[0] == "nc" for m in made):
            err = pg.evaluate(SETUP, {"fix": F, "co": FIRMCO, "auditor": AUDITOR, "kind": "nc", "fs": FM_FS})
            ok(not err, "the firm's statements are worked out" + (": " + err if err else ""))
            with ctx.expect_page(timeout=30000) as pop:
                pg.evaluate("() => { FinPdf.open('2025', {unit: 1, dec: 2}); }")
            w = pop.value; w.wait_for_load_state(); w.wait_for_timeout(600); t = w.inner_text("body")
            ok("Partners' funds" in t and t.count("For and on behalf of\nKestrel & Co") == 3 and "Board of Directors" not in t and "Asha Rao" in t and "Partner" in t,
               "a firm's PDF: partners' funds, signed for the firm by its partners")
            w.close()
        x = pg.evaluate(MAKE, {"unit": unit, "dec": dec})
        path = os.path.join(OUT, "%s-u%d-d%d.xlsx" % (kind, unit, dec))
        open(path, "wb").write(base64.b64decode(x["b64"]))
        made.append((kind, unit, dec, path, x))

    br.close()

print("\nThe workbooks, as written")
for kind, unit, dec, path, x in made:
    wb = openpyxl.load_workbook(path); wv = openpyxl.load_workbook(path, data_only=True)
    want = CO_SHEETS if kind == "co" else NC_SHEETS
    ok(wb.sheetnames == want, "%s: only the listed sheets %s" % (os.path.basename(path), "" if wb.sheetnames == want else wb.sheetnames))
    names = set(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else {n.name for n in wb.defined_names.definedName}
    ok({"UNIT", "DEC", "TOL", "SHARES"} <= names, "%s: names UNIT, DEC, TOL, SHARES" % os.path.basename(path))
    ok(wv["Inputs"]["E4"].value == unit and wv["Inputs"]["E5"].value == dec, "%s: Inputs!E4 = %s, E5 = %s" % (os.path.basename(path), unit, dec))
    nof = [c for c in x["cells"]]
    missing = [c["sheet"] + "!" + c["addr"] for c in nof if wv[c["sheet"]][c["addr"]].value is None and c["v"] not in (None, "")]
    ok(not missing and len(nof) > 300, "%s: %d formula cells, each with its value written (%s)" % (os.path.basename(path), len(nof), missing[:5]))
    nform = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    ok(nform == len(nof), "%s: every formula in the file is one FinCom listed (%d / %d)" % (os.path.basename(path), nform, len(nof)))
    bad = [(c["sheet"] + "!" + c["addr"], c["v"], resolve(x["d"], c["src"])) for c in nof if c["src"] and not same(c["v"], resolve(x["d"], c["src"]))]
    ok(not bad, "%s: the values written are FinSt.build's figures %s" % (os.path.basename(path), bad[:4]))
    if kind == "co" and unit == 1000 and dec == 2: shutil.copy(path, os.path.join(OUT, "co-sample-make.xlsx"))

print("\nRecalculated by LibreOffice from the formulas alone")
rc = recalc([m[3] for m in made] + [os.path.join(OUT, "sample.xlsx")])
if rc is None:
    ok(os.environ.get("FIN_XLSX_NO_SOFFICE") == "1", "LibreOffice (soffice) is installed (set FIN_XLSX_NO_SOFFICE=1 to skip this part)")
else:
    with zipfile.ZipFile(os.path.join(OUT, "bare", os.path.basename(made[0][3]))) as z:
        left = sum(z.read(n).decode("utf8").count("</f><v>") for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
    ok(left == 0, "LibreOffice was given the formulas without their values (%d left)" % left)
    for kind, unit, dec, path, x in made:
        nm = os.path.basename(path)
        lv = openpyxl.load_workbook(rc[path], data_only=True)
        errs = [(ws.title, c.coordinate, c.value) for ws in lv.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and (c.value.startswith("#") or c.value.startswith("Err:"))]
        ok(not errs, "%s: no #REF!, #NAME? or other error %s" % (nm, errs[:5]))
        diff = [(c["sheet"] + "!" + c["addr"], c["f"][:90], c["v"], lv[c["sheet"]][c["addr"]].value) for c in x["cells"] if not same(lv[c["sheet"]][c["addr"]].value, c["v"])]
        ok(not diff, "%s: every formula cell recalculates to the value FinCom wrote (%d cells)%s" % (nm, len(x["cells"]), "" if not diff else ": " + "; ".join(map(str, diff[:5]))))
        figs = [c for c in x["cells"] if c["src"]]
        bad = [(c["sheet"] + "!" + c["addr"], lv[c["sheet"]][c["addr"]].value, resolve(x["d"], c["src"])) for c in figs if not same(lv[c["sheet"]][c["addr"]].value, resolve(x["d"], c["src"]))]
        ok(not bad and len(figs) > 150, "%s: %d figures of FinSt.build equal their formulas' values %s" % (nm, len(figs), bad[:4]))
        ws = lv["Inputs"]; chk = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "Checks (all should be zero)"]
        vals = [(ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(chk[0] + 1, ws.max_row + 1) if ws.cell(r, 2).value] if chk else []
        ok(len(vals) >= 9 and all(v == 0 for _, v in vals), "%s: the %d checks on Inputs all read 0 %s" % (nm, len(vals), [z for z in vals if z[1] != 0]))
        d = x["d"]; bs = lv["Balance Sheet"]
        tots = [bs.cell(r, 6).value for r in range(1, bs.max_row + 1) if bs.cell(r, 3).value == "TOTAL"]
        ok(len(tots) == 2 and same(tots[0], d["bs"]["total"]["eqL"]["cy"]) and same(tots[1], d["bs"]["total"]["assets"]["cy"]) and same(tots[0], tots[1]),
           "%s: the balance sheet tallies at %s, as on the screen" % (nm, tots[:2]))
    # the unit changed in the workbook itself (Inputs!E4): every figure follows, the checks stay 0
    src = [m for m in made if m[:3] == ("co", 1000, 2)][0]; lakhs = [m for m in made if m[:3] == ("co", 100000, 2)][0][4]["d"]
    ch = os.path.join(OUT, "unit-changed.xlsx"); strip_values(src[3], ch)
    wbc = openpyxl.load_workbook(ch); wbc["Inputs"]["E4"] = 100000; wbc.save(ch)
    lv = openpyxl.load_workbook(recalc([ch])[ch], data_only=True); bs = lv["Balance Sheet"]; ws = lv["Inputs"]
    tots = [bs.cell(r, 6).value for r in range(1, bs.max_row + 1) if bs.cell(r, 3).value == "TOTAL"]
    chk = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "Checks (all should be zero)"][0]
    nz = [(ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(chk + 1, ws.max_row + 1) if ws.cell(r, 2).value and ws.cell(r, 3).value != 0]
    ok(ws["E6"].value == "(All amounts in Rupees lakhs, unless otherwise stated)" and len(tots) == 2 and same(tots[0], lakhs["bs"]["total"]["eqL"]["cy"]) and same(tots[1], tots[0]) and not nz,
       "the unit changed in the workbook (thousands -> lakhs on Inputs!E4): the balance sheet tallies at %s as FinSt's in lakhs, checks 0 %s" % (tots, nz))
    sv = openpyxl.load_workbook(rc[os.path.join(OUT, "sample.xlsx")], data_only=True)
    ok(sv["Inputs"]["E6"].value == "(All amounts in Rupees thousand, unless otherwise stated)", "the downloaded workbook recalculates too: " + str(sv["Inputs"]["E6"].value))

dest = os.environ.get("FIN_SAMPLES")
if dest:
    os.makedirs(dest, exist_ok=True)
    for f in ("sample.xlsx", "sample.pdf"):
        if os.path.exists(os.path.join(OUT, f)): shutil.copy(os.path.join(OUT, f), os.path.join(dest, f))
errs = [e for e in errors if "supabase" not in e.lower() and "Failed to load" not in e]
ok(not errs, "no page errors" + ("" if not errs else ": " + " | ".join(errs[:4])))
print("\n" + ("%d FAILED" % len(fails) if fails else "all passed")); sys.exit(1 if fails else 0)
