"""python3 run_reports_pack_ui.py - the reports pack (ROUND 100) on the Financial statements page, in a browser, on made-up
books (tests/fixtures/fin-statements/reports_books.js: the company of run_fin_statements.js with a director's loan, a loan
given, a director's remuneration and a TDS payable unpaid for more than six months):
  - the page keeps its seven tabs; a "Reports" menu in the header opens the Board's Report, the Auditor's Report, the
    management representation letter and the questionnaire;
  - each report previews, the answers still missing marked and counted in one chip; a mark opens its question;
  - answers given in the questionnaire are kept (b.fs.reports) and the chip reaches "All answers given";
  - the figures follow the page's unit;
  - Download Word: a real .docx (the `docx` library from assets/docx.iife.js, checked by its sha256), opened here with
    python-docx: headings, numbered sections, tables (financial summary, CARO), figures;
  - Download PDF: the print view holds the report; printed to PDF, its text (read with the site's own pdf.js) has it;
  - Download all: one PDF with the statements and the three reports; a zip of the Excel and the three Word files.
Samples (made-up company) go to REPORTS_SAMPLES (default tests/out/reports-samples): board-report.docx/.pdf,
auditor-report.docx/.pdf, mrl.docx, all.pdf, all.zip and screenshots.
Run on the React build: TDSDESK_SITE=../app/dist-test python3 run_reports_pack_ui.py"""
import base64, hashlib, io, json, os, subprocess, threading, functools, http.server, zipfile
import docx
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
from playwright.sync_api import sync_playwright
HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.environ.get("TDSDESK_SITE", os.path.join(HERE, "..", "app", "dist-test"))
SAMPLES = os.environ.get("REPORTS_SAMPLES", os.path.join(HERE, "out", "reports-samples")); os.makedirs(SAMPLES, exist_ok=True)
PORT = 8268
class Q(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(("localhost", PORT), functools.partial(Q, directory=SITE)); threading.Thread(target=srv.serve_forever, daemon=True).start()
fails, errors = [], []
def ok(c, w):
    print(("  ok   " if c else "  FAIL ") + w)
    if not c: fails.append(w)

BOOKS = json.loads(subprocess.check_output(["node", os.path.join(HERE, "fixtures", "fin-statements", "reports_books.js")]))
SETUP = """(bk) => {
  S.firm = Object.assign({}, S.firm, {firmName: "Made-up & Associates", frn: "000000S", partner: "A. Sample", mno: "000001", place: "Bengaluru"});
  const c = newCompany({name: "Larkspur Analytics Private Limited (made-up)", gstin: "29AABCL1234F1Z5"}); c.pan = "AABCL1234F"; c.cin = "U72900KA2015PTC000001";
  S.companies[c.id] = c; Store.saveCompany(c); S.coId = c.id; S.view = "company"; S.tab = "books"; S.loadingCo = false;
  const b = Object.assign({loading: false, cid: c.id, challans: [], alloc: {}, ledInfo: {}, gstins: {}, pans: {}, states: {}, misCfg: {freq: "off"}, auditCfg: {freq: "off"}}, bk, {
    msme: {"Bluebell Print Works": "Micro"}, fs: {taxDep: {"2024": {c40: 220000, pm15: 140000, f10: 80000}}, regime: "115BAA",
      inputs: {unit: 1000, dec: 2, address: "12 Made-up Road, Bengaluru", place: "Bengaluru", date: "20260905", udin: "26000001ABCDEF1234",
        directors: [{name: "Asha Rao", din: "00000001", desig: "Director"}, {name: "Vikram Sen", din: "00000002", desig: "Managing Director"}], shares: {number: 10000, face: 10, authorised: 500000, holders: []}}}});
  b.map = Books.mapLedgers(b.vouchers, {});
  S.books = b; window.__bk = b; S.fsFy = "2025"; S.booksTab = "dash"; saveBooks(); render(); return c.id; }"""
# the answers a careful auditor gives for the made-up company (the rest of them typed in the questionnaire below)
REST = {"e_agm_no": 11, "b_meetings": [{"date": "20250520"}, {"date": "20250905"}, {"date": "20251210"}, {"date": "20260310"}], "b_aud_agm": "8th Annual General Meeting held on 30th September, 2023",
        "b_aud_till": "13th Annual General Meeting", "b_posh_disp": 0, "b_posh_pend": 0, "b_emp_f": 9, "b_emp_m": 14, "m_engagement": "20250701", "b_loan_purpose": "working capital of the borrower",
        "rp_parties": [{"name": "Asha Rao", "relation": "Director"}, {"name": "Vikram Sen", "relation": "Managing Director"}]}
APP = lambda pg: pg.inner_text("#app")
def shot(pg, name): pg.screenshot(path=os.path.join(SAMPLES, name + ".png"), full_page=True)
def chip(pg): return pg.locator("#app .rp-meta .rp-chip").first
def menu(pg, k):
    pg.click("#app [data-fin-reports] > summary"); pg.wait_for_timeout(200)
    pg.click('#app [data-fin-reports] [data-fin-rep="%s"]' % k); pg.wait_for_timeout(1200)
def pick(pg, k): pg.click('#app [data-rp-pick="%s"]' % k); pg.wait_for_timeout(800)
def pdf_text(pg, path):
    b64 = base64.b64encode(open(path, "rb").read()).decode()
    return pg.evaluate("""async (b64) => { await ensurePdfJs(); const bin = atob(b64), u = new Uint8Array(bin.length); for (let i = 0; i < bin.length; i++) u[i] = bin.charCodeAt(i);
      const doc = await openPdfDoc(u); let t = ""; for (let p = 1; p <= doc.numPages; p++) { const c = await (await doc.getPage(p)).getTextContent(); t += c.items.map(x => x.str).join(" ") + "\\n"; } return {n: doc.numPages, t}; }""", b64)
def docx_of(path):
    d = docx.Document(path)
    paras = [p.text for p in d.paragraphs]
    heads = [p.text for p in d.paragraphs if p.style is not None and p.style.name.startswith("Heading")]
    return d, paras, heads, "\n".join(paras + [c.text for t in d.tables for r in t.rows for c in r.cells])

with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(viewport={"width": 1440, "height": 900}, color_scheme="light", accept_downloads=True); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    asked = []; pg.on("request", lambda r: asked.append(r.url))
    pg.goto("http://localhost:%d/" % PORT); pg.wait_for_timeout(2500); pg.click('button[data-act="useOffline"]'); pg.wait_for_timeout(1200)
    cid = pg.evaluate(SETUP, BOOKS); pg.wait_for_timeout(1200); pg.evaluate("S.books = window.__bk; goClient('books:fs');"); pg.wait_for_timeout(2500)

    print("The page and its Reports menu")
    tabs = pg.locator("#app nav.fin-tabs [data-fin-tab]").evaluate_all("xs => xs.map(x => x.innerText.trim())")
    ok(len(tabs) == 7 and tabs[0] == "Balance sheet", "still seven tabs (%s)" % tabs)
    pg.click("#app [data-fin-reports] > summary"); pg.wait_for_timeout(300)
    items = pg.locator("#app [data-fin-reports] .more-item").evaluate_all("xs => xs.map(x => x.innerText.trim())")
    ok(items == ["Board's Report", "Auditor's Report", "Management representation", "Questionnaire", "Download all as one PDF", "Download all as a zip"], "the Reports menu: %s" % items)
    pg.click("#app [data-fin-reports] > summary"); pg.wait_for_timeout(200)
    ok(pg.evaluate("!window.docx") and not any("docx.iife.js" in u for u in asked), "the Word writer is not loaded with the page")

    print("Board's Report: preview, the missing answers")
    menu(pg, "board")
    ok(pg.locator('#app article[data-rp-preview="board"]').count() == 1 and pg.locator('#app nav.fin-tabs [aria-selected="true"]').count() == 0, "the preview shows; no tab is selected")
    t = pg.inner_text('#app article[data-rp-preview="board"]')
    h1 = pg.locator("#app article h3.rp-h1").evaluate_all("xs => xs.map(x => x.textContent.trim())")
    ok("BOARD'S REPORT" in t.upper() and h1[0] == "1.Financial summary" and any(h.endswith("Directors' Responsibility Statement") for h in h1), "headings: Board's Report, 1. Financial summary, Directors' Responsibility Statement (%s ...)" % h1[:2])
    ok("Rule 8A" in t and "small company" in pg.inner_text("#app [data-rp-applies]"), "a small company: the abridged report, said in the applicability line (%s)" % pg.inner_text("#app [data-rp-applies]"))
    n0 = int(chip(pg).get_attribute("data-rp-missing"))
    marks = pg.locator("#app article mark.rp-miss").count()
    ok(n0 > 3 and marks >= n0 and "answers missing" in chip(pg).inner_text(), "one chip counts the missing answers (%d), each marked in the text (%d marks)" % (n0, marks))
    rev = pg.locator('#app article[data-rp-preview="board"] table.rp-t tbody tr').first.locator("td").nth(1).inner_text().strip()
    ok(rev == "12,200.00", "the financial summary in the page's unit (thousands): revenue %s" % rev)
    shot(pg, "board-missing")
    pg.click('#app article mark.rp-miss[data-miss="b_business"]'); pg.wait_for_timeout(1000)
    ok(pg.locator("#app [data-rp-questionnaire]").count() == 1 and pg.evaluate("document.activeElement && document.activeElement.closest('[data-q]') && document.activeElement.closest('[data-q]').dataset.q") == "b_business",
       "a mark opens the questionnaire at its question, the box in focus")

    print("The questionnaire")
    ok(pg.locator('#app [data-q="b_business"][data-q-need]').count() == 1 and pg.locator('#app [data-q="e_class"] [data-q-src="books"]').count() == 1, "the needed answers marked; the class of company from the CIN, said so")
    pg.fill('#app [data-q="b_business"] textarea', "software development and IT-enabled services for customers in India and abroad"); pg.press('#app [data-q="b_business"] textarea', "Tab"); pg.wait_for_timeout(600)
    pg.select_option('#app [data-q="b_maternity"] select', "yes"); pg.wait_for_timeout(500)
    pg.fill('#app [data-q="b_posh_recv"] input', "0"); pg.press('#app [data-q="b_posh_recv"] input', "Tab"); pg.wait_for_timeout(500)
    pg.fill('#app [data-q="e_udin"] input', "26000001ABCDEF1234"); pg.press('#app [data-q="e_udin"] input', "Tab"); pg.wait_for_timeout(500)
    a = pg.evaluate("S.books.fs.reports['2025'].a")
    ok(a.get("b_business", "").startswith("software development") and a.get("b_maternity") == "yes" and a.get("b_posh_recv") == 0, "answers kept in b.fs.reports['2025'] (%s)" % sorted(a.keys()))
    pg.evaluate("(r) => { const a = S.books.fs.reports['2025'].a; Object.assign(a, r); saveBooks(); render(); }", REST); pg.wait_for_timeout(1200)
    pg.check('#app input[aria-label="Only the answers missing"]'); pg.wait_for_timeout(500)
    left = pg.locator("#app [data-q]").evaluate_all("xs => xs.map(x => x.dataset.q)")
    ok(chip(pg).get_attribute("data-rp-missing") == "0" and not left, "every needed answer given: “All answers given”, nothing left under “only the answers missing” (%s)" % left)
    pg.uncheck('#app input[aria-label="Only the answers missing"]'); pg.wait_for_timeout(300)
    shot(pg, "questionnaire")

    print("Board's Report: complete, the unit, Word and PDF")
    pick(pg, "board")
    ok(chip(pg).inner_text().strip() == "All answers given" and pg.locator("#app article mark.rp-miss").count() == 0, "the Board's Report: no gap left")
    t = pg.inner_text('#app article[data-rp-preview="board"]')
    ok("9 female, 14 male and 0 transgender employees" in t and "received during the year: 0" in t, "POSH counts and employees by gender in the text")
    shot(pg, "board")
    pg.click('#app [data-fin-seg="Unit"] button:text-is("Lakhs")'); pg.wait_for_timeout(1200)
    rev = pg.locator('#app article[data-rp-preview="board"] table.rp-t tbody tr').first.locator("td").nth(1).inner_text().strip()
    ok(rev == "122.00" and "(All amounts in Rupees lakhs, unless otherwise stated)" in pg.inner_text('#app article[data-rp-preview="board"]'), "the unit switched to lakhs: revenue %s, the unit line" % rev)
    with pg.expect_download(timeout=90000) as dl:
        pg.click('#app [data-rp-dl="docx"]')
    path = os.path.join(SAMPLES, "board-report.docx"); dl.value.save_as(path)
    d, paras, heads, alltext = docx_of(path)
    sha = hashlib.sha256(open(os.path.join(SITE, "assets", "docx.iife.js"), "rb").read()).hexdigest()
    ok(pg.evaluate("window.ReportsPack && window.ReportsPack.sha256") == sha and os.path.exists(os.path.join(SITE, "assets", "docx-LICENSE.txt")) and any(u.endswith("/assets/docx.iife.js") for u in asked),
       "the Word writer loaded on first use from assets/docx.iife.js, the file the module names (sha256 %s...), its licence beside it" % sha[:16])
    ok(dl.value.suggested_filename.endswith("Board's Report FY 2025-26.docx"), "the file: %s" % dl.value.suggested_filename)
    ok(heads[:3] == ["1. Financial summary", "2. State of the Company's affairs", "3. Dividend"] and "Directors' Responsibility Statement" in " ".join(heads) and any(h.endswith("Acknowledgements") for h in heads), "Word headings, numbered: %s ..." % heads[:4])
    summ = d.tables[0]
    ok(summ.rows[0].cells[0].text == "Particulars" and summ.rows[1].cells[0].text == "Revenue from operations" and summ.rows[1].cells[1].text == "122.00", "the financial summary is a Word table, in lakhs (%s)" % summ.rows[1].cells[1].text)
    ok(any(p.startswith("(a)") for p in paras) and "UDIN" not in alltext and "Asha Rao" in alltext and "Place: Bengaluru" in alltext, "numbered clauses, the directors' signature block")
    ok("Page" in "".join(x.text for x in d.sections[0].footer.paragraphs) and len(d.sections[0].footer.paragraphs[0]._p.xpath(".//w:fldChar")) + len(d.sections[0].footer.paragraphs[0]._p.xpath(".//w:fldSimple")) + len(d.sections[0].footer.paragraphs[0]._p.xpath(".//w:instrText")) > 0, "page numbers in the footer")
    with ctx.expect_page(timeout=60000) as pop:
        pg.click('#app [data-rp-dl="pdf"]')
    w = pop.value; w.wait_for_load_state(); w.wait_for_timeout(800)
    ok(w.locator('section.rp-doc[data-rp-doc="board"]').count() == 1 and w.locator("mark.rp-miss").count() == 0, "Download PDF: the print view holds the Board's Report")
    pth = os.path.join(SAMPLES, "board-report.pdf"); w.pdf(path=pth, format="A4", prefer_css_page_size=True, print_background=True); w.close()
    pt = pdf_text(pg, pth)
    ok("BOARD" in pt["t"] and "Directors" in pt["t"] and "Responsibility" in pt["t"] and "122.00" in pt["t"], "the PDF (%d pages) has the report and its figures" % pt["n"])

    print("Auditor's Report")
    pick(pg, "audit")
    t = pg.inner_text('#app article[data-rp-preview="audit"]')
    ok("INDEPENDENT AUDITOR'S REPORT" in t.upper() and "Report on Other Legal and Regulatory Requirements" in t and "small company" in t and "UDIN: 26000001ABCDEF1234" in t, "the preview: SA 700 headings, CARO and IFC not applicable (small company), UDIN")
    shot(pg, "auditor")
    with pg.expect_download(timeout=60000) as dl:
        pg.click('#app [data-rp-dl="docx"]')
    path = os.path.join(SAMPLES, "auditor-report.docx"); dl.value.save_as(path)
    d, paras, heads, alltext = docx_of(path)
    ok("Opinion" in heads and "Basis for Opinion" in heads and "Report on Other Legal and Regulatory Requirements" in heads and "does not apply to the Company, the Company being a small company" in alltext, "Word: headings %s; CARO not applicable said" % heads[:3])
    with ctx.expect_page(timeout=60000) as pop:
        pg.click('#app [data-rp-dl="pdf"]')
    w = pop.value; w.wait_for_load_state(); w.wait_for_timeout(800)
    pth = os.path.join(SAMPLES, "auditor-report.pdf"); w.pdf(path=pth, format="A4", prefer_css_page_size=True, print_background=True); w.close()
    pt = pdf_text(pg, pth)
    ok("INDEPENDENT AUDITOR" in pt["t"].upper() and "Basis for Opinion" in pt["t"], "the auditor's report PDF (%d pages)" % pt["n"])
    # the same books answered as a public company: CARO (Annexure A) and IFC (Annexure B) come in
    pick(pg, "q")
    pg.select_option('#app [data-q="e_class"] select', "public"); pg.wait_for_timeout(800)
    pg.evaluate("() => { Object.assign(S.books.fs.reports['2025'].a, {b_kmp: 'Vikram Sen, Managing Director', b_rotation: 'Asha Rao (DIN 00000001)'}); saveBooks(); render(); }"); pg.wait_for_timeout(800)
    pick(pg, "audit")
    with pg.expect_download(timeout=60000) as dl:
        pg.click('#app [data-rp-dl="docx"]')
    path = os.path.join(SAMPLES, "auditor-report-public.docx"); dl.value.save_as(path)
    d, paras, heads, alltext = docx_of(path)
    caro = [p for p in paras if p.startswith("(") and "\t" in p and p.split("\t")[0] in ["(i)", "(ii)", "(iii)", "(iv)", "(v)", "(vi)", "(vii)", "(viii)", "(ix)", "(x)", "(xi)", "(xii)", "(xiii)", "(xiv)", "(xv)", "(xvi)", "(xvii)", "(xviii)", "(xix)", "(xx)", "(xxi)"]]
    ok("ANNEXURE A TO THE INDEPENDENT AUDITOR'S REPORT" in alltext and "ANNEXURE B TO THE INDEPENDENT AUDITOR'S REPORT" in alltext and "Meaning of Internal Financial Controls with Reference to Financial Statements" in heads, "a public company: Annexure A (CARO) and Annexure B (IFC) in Word")
    ok(len(d.tables) >= 3 and any("TDS Payable 194C" in c.text for t_ in d.tables for r in t_.rows for c in r.cells), "CARO tables in Word, the TDS outstanding more than six months among them (%d tables)" % len(d.tables))
    shot(pg, "auditor-public")
    pick(pg, "q"); pg.select_option('#app [data-q="e_class"] select', "private"); pg.wait_for_timeout(800)

    print("Management representation letter")
    pick(pg, "mrl")
    t = pg.inner_text('#app article[data-rp-preview="mrl"]')
    ok("Made-up & Associates" in t and "Representations for the audit" in t and "Vikram Sen" in t, "the preview: to the auditor, the subject, signed by the directors")
    shot(pg, "mrl")
    with pg.expect_download(timeout=60000) as dl:
        pg.click('#app [data-rp-dl="docx"]')
    path = os.path.join(SAMPLES, "mrl.docx"); dl.value.save_as(path)
    d, paras, heads, alltext = docx_of(path)
    ok("A. The financial statements" in heads and "E. Statutory confirmations" in heads and "Annexure 1: Summary of uncorrected misstatements" in heads and "₹35.00 lakh" not in alltext and "cash on hand of ₹" in alltext, "Word: sections A to E, the annexures, the cash from the books")
    with ctx.expect_page(timeout=60000) as pop:
        pg.click('#app [data-rp-dl="pdf"]')
    w = pop.value; w.wait_for_load_state(); w.wait_for_timeout(800)
    w.pdf(path=os.path.join(SAMPLES, "mrl.pdf"), format="A4", prefer_css_page_size=True, print_background=True); w.close()

    print("Download all")
    pg.click("#app [data-fin-reports] > summary"); pg.wait_for_timeout(200)
    with ctx.expect_page(timeout=60000) as pop:
        pg.click('#app [data-fin-reports] [data-fin-all="pdf"]')
    w = pop.value; w.wait_for_load_state(); w.wait_for_timeout(1500)
    pages = w.eval_on_selector_all("section.fs-page", "xs => xs.map(x => x.dataset.fsPage)"); docs = w.eval_on_selector_all("section.rp-doc", "xs => xs.map(x => x.dataset.rpDoc)")
    ok(pages == ["bs", "pl", "cf", "notes", "comp"] and docs == ["board", "audit", "mrl"], "one print view: the statements %s, then the reports %s" % (pages, docs))
    pth = os.path.join(SAMPLES, "all.pdf"); w.pdf(path=pth, format="A4", prefer_css_page_size=True, print_background=True); w.close()
    pt = pdf_text(pg, pth)
    ok(all(x in pt["t"].upper() for x in ["BALANCE SHEET", "BOARD'S REPORT", "INDEPENDENT AUDITOR'S REPORT", "REPRESENTATIONS FOR THE AUDIT"]), "the PDF (%d pages) has the statements and each report" % pt["n"])
    pg.click("#app [data-fin-reports] > summary"); pg.wait_for_timeout(200)
    with pg.expect_download(timeout=120000) as dl:
        pg.click('#app [data-fin-reports] [data-fin-all="zip"]')
    zp = os.path.join(SAMPLES, "all.zip"); dl.value.save_as(zp)
    z = zipfile.ZipFile(zp); names = z.namelist()
    ok(z.testzip() is None and len(names) == 4 and names[0].endswith("Financial statements FY 2025-26.xlsx") and sum(n.endswith(".docx") for n in names) == 3, "the zip: the Excel and the three Word files (%s)" % names)
    ok(all(len(docx.Document(io.BytesIO(z.read(n))).paragraphs) > 20 for n in names if n.endswith(".docx")), "... each Word file opens")

    print("Kept after a reload")
    pg.wait_for_timeout(1500); pg.reload(); pg.wait_for_timeout(2500)
    if pg.locator('button[data-act="useOffline"]').count(): pg.click('button[data-act="useOffline"]'); pg.wait_for_timeout(1000)
    pg.evaluate("(cid) => { S.fsFy = '2025'; openCompany(cid).then(() => { goClient('books:fs'); }); }", cid); pg.wait_for_timeout(4000)
    ok(str(pg.evaluate("((S.books.fs.reports || {})['2025'] || {a: {}}).a.b_business || ''")).startswith("software development"), "the answers are in the client's books after a reload")
    menu(pg, "board")
    ok(chip(pg).get_attribute("data-rp-missing") == "0", "... and the Board's Report still has no gap")
    ctx.close()

    print("A phone, 390 wide")
    ctx = br.new_context(viewport={"width": 390, "height": 844}, device_scale_factor=2, is_mobile=True, has_touch=True, color_scheme="light"); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    pg.goto("http://localhost:%d/" % PORT); pg.wait_for_timeout(2500); pg.click('button[data-act="useOffline"]'); pg.wait_for_timeout(1200)
    pg.evaluate(SETUP, BOOKS); pg.wait_for_timeout(1200); pg.evaluate("S.books = window.__bk; goClient('books:fs');"); pg.wait_for_timeout(2500)
    for k in ["board", "q"]:
        pg.evaluate("(k) => { S.fin = Object.assign(S.fin || {}, {rep: k}); render(); window.scrollTo(0, 0); }", k); pg.wait_for_timeout(1500)
        wdt = pg.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth")
        ok(wdt <= 0, "phone, %s: the page does not scroll sideways (%d px)" % (k, wdt))
        shot(pg, "phone-" + k)
    ctx.close()
    br.close()
errs = [e for e in errors if "supabase" not in e.lower() and "Failed to load" not in e]
ok(not errs, "no page errors" + ("" if not errs else ": " + " | ".join(errs[:4])))
print("\nsamples in " + SAMPLES)
print("\n" + ("%d FAILED" % len(fails) if fails else "all passed")); raise SystemExit(1 if fails else 0)
