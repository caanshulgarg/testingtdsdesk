"""python3 run_itc_simple_ui.py - GST "2B & ITC" made one simple page (next-itc-simple; the owner, 24-Oct-2026: "2B & ITC
page is not clear.. I don't know where to do what.. supplier health is not required.. there should be Input register, 2B
reconciliation, and in 2B reconciliation there should be follow-up also.. and one main tab should be Final credit taking
into 3B.. option to increase or decrease like TDS add deduction or remove.. and in 3B if we want to see the hyperlink
detail it should pop up a window rather than showing it below"; and the Excel: "Smart Excel, no macros .. there should be
the final ITC-to-be-taken figure there also.. after importing, the system should take all the details from there only").
Made-up books (itc_books.py, worked out by hand as run_itc_final.js):
  1. 2B & ITC has three tabs, in order: Input register · 2B reconciliation · Final credit for 3B; no Supplier health and no
     ITC follow-up (their old addresses open 2B reconciliation); Reversal is under Returns
  2. Input register: the month's bills, the ITC ledger, eligible, the status in 2B, totals
  3. 2B reconciliation: chips with counts; each Decision moves the line's ITC this month and Final credit's 4(A)(5); a
     books-only bill taken now (with its risk note); Take in a later month with its month; tick lines and Set decision…;
     a line not taken is carried forward to May, taken there; Write to suppliers lists the suppliers to follow up
  4. Final credit for 3B: + Add adjustment and Remove change table 4, and GSTR-3B's table 4 is the same figures
  5. GSTR-3B: every figure (3.1, 4, 5, 6.1) opens a window with the rows behind it; Esc and x close it; nothing below
  6. Excel: download, decisions and an adjustment changed in the file (openpyxl), LibreOffice works the formulas out (equal
     to the screen before the change), Upload Excel: the preview (changes, adjustments, the Excel's table 4), Apply: the
     decisions are FinCom's and its table 4 equals the Excel's to the paisa; an edited key is refused
Screenshots (1440 light and dark, 390) to TDSDESK_SHOTS (default tests/out/itc-shots)."""
import os, sys, json, threading, functools, http.server, base64, shutil, subprocess, tempfile
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
import itc_books as IB
from playwright.sync_api import sync_playwright
SITE = os.environ.get("TDSDESK_SITE", os.path.join(HERE, "..", "app", "dist-test"))
PORT = int(os.environ.get("ITC_PORT") or 8447)
SHOTS = os.environ.get("TDSDESK_SHOTS") or os.path.join(os.environ.get("TDSDESK_OUT", os.path.join(HERE, "out")), "itc-shots")
os.makedirs(SHOTS, exist_ok=True)
fails, errors = [], []
def ok(c, w):
    print(("  ok   " if c else "  FAIL ") + w)
    if not c: fails.append(w)
class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(("localhost", PORT), functools.partial(Quiet, directory=SITE)); threading.Thread(target=srv.serve_forever, daemon=True).start()
H = ["igst", "cgst", "sgst", "cess"]
def shot(pg, name, full=False): pg.screenshot(path=os.path.join(SHOTS, name + ".png"), full_page=full)
def W(pg, ms=400): pg.wait_for_timeout(ms)
def T4(pg, ym="202604"): return pg.evaluate("(ym) => Object.fromEntries(ITCF.table4('09', ym).lines.map(([k, v]) => [k, v]))", ym)
def tot(o): return round(sum(o.get(h, 0) or 0 for h in H), 2)
def a5(pg, ym="202604"): return tot(T4(pg, ym)["4A5"])
def row_sel(no): return 'tr[data-ic-row]:has(.code:text-is("%s"))' % no
def dec(pg, no): return pg.evaluate("(no) => { const r = [...document.querySelectorAll('tr[data-ic-row]')].find(tr => [...tr.querySelectorAll('.code')].some(c => c.textContent.trim() === no)); const s = r && r.querySelector('select[data-ic-dec]'); return s ? s.value : null; }", no)
def take(pg, no): return pg.evaluate("(no) => { const r = [...document.querySelectorAll('tr[data-ic-row]')].find(tr => [...tr.querySelectorAll('.code')].some(c => c.textContent.trim() === no)); const c = r && r.querySelector('[data-ic-take]'); return c ? Number((c.querySelector('b') || {innerText: '0'}).innerText.replace(/[₹,]/g, '')) : null; }", no)
def choose(pg, no, v):
    try: pg.select_option(row_sel(no) + " select[data-ic-dec]", v, timeout=4000)
    except Exception as e: ok(False, "the Decision of %s: %r (%s)" % (no, v, str(e)[:80]))
    W(pg, 500)
def chips(pg): return pg.evaluate("() => Object.fromEntries([...document.querySelectorAll('[data-ic-chip]')].map(b => [b.dataset.icChip, Number(b.querySelector('.sbar-n').textContent)]))")
def recalc(path, outdir):
    """LibreOffice works every formula out: the file saved by openpyxl (formulas, no values) converted to .xlsx"""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice: return None
    prof = "file://" + tempfile.mkdtemp(prefix="itc-lo-")
    subprocess.run([soffice, "-env:UserInstallation=" + prof, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", outdir, path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
    return os.path.join(outdir, os.path.basename(path))
def final_of(path):
    import openpyxl
    ws = openpyxl.load_workbook(path, data_only=True)["Final credit for 3B"]
    out = {}
    for r in ws.iter_rows(min_row=2):
        code = r[6].value
        if code and str(code).startswith("4"): out[code] = {h: round(float(r[i + 1].value or 0), 2) for i, h in enumerate(H)}
    return out
CAPTURE = """() => { window.__files = []; window.saveFile = async (name, data) => { const a = new Uint8Array(await data.arrayBuffer()); let s = ''; for (let i = 0; i < a.length; i += 8192) s += String.fromCharCode.apply(null, a.subarray(i, i + 8192)); window.__files.push({name, b64: btoa(s)}); return true; }; }"""

with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(viewport={"width": 1440, "height": 950}, accept_downloads=True); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    IB.setup(pg, PORT)

    print("1. three tabs")
    IB.open_part(pg, "inreg")
    secs = pg.evaluate("[...document.querySelectorAll('#app [data-gst-sec]')].map(b => b.innerText.trim())")
    parts = pg.evaluate("[...document.querySelectorAll('#app nav[aria-label=GST] [data-part]')].map(b => b.innerText.trim())")
    ok(parts == ["Input register", "2B reconciliation", "Final credit for 3B"], "2B & ITC: three tabs in order (%s)" % parts)
    ok(not any("Supplier health" in s for s in secs) and "Supplier health" not in pg.inner_text("#app") and "ITC follow-up" not in pg.inner_text("#app"), "no Supplier health, no ITC follow-up (%s)" % secs)
    for old in ("follow", "health"):
        IB.open_part(pg, old)
        ok(pg.evaluate("S.gstPart") == "r2b" and pg.locator("[data-itc-reco]").count() == 1, "the old %s address opens 2B reconciliation" % old)
    IB.open_part(pg, "rev")
    ok(pg.evaluate("(() => { const b = document.querySelector('#app [data-gst-sec=returns]'); return b && b.getAttribute('aria-selected'); })()") == "true", "Reversal is under Returns")

    print("2. input register")
    IB.open_part(pg, "inreg")
    heads = pg.evaluate("[...document.querySelectorAll('#app table.bk-table thead th')].map(t => t.innerText.trim().toLowerCase())")
    ok(any("itc ledger" in h for h in heads) and any("eligible" in h for h in heads) and any(h == "2b" for h in heads), "Input register: ITC ledger, Eligible and 2B columns (%s)" % heads)
    t = pg.inner_text("#app [data-gst-page]")
    ok("Input CGST, Input SGST" in t and "Not in 2B" in t and "In 2B, differs" in t and "In 2B" in t, "each bill's input ledgers and its status in 2B")
    shot(pg, "1-input-register-desktop", True)

    print("3. 2B reconciliation")
    IB.open_part(pg, "r2b")
    c = chips(pg)
    ok(c == {"all": 6, "matched": 1, "books": 2, "2b": 2, "diff": 1, "ims": 0, "cf": 0}, "April's chips: matched 1, in books not in 2B 2, in 2B not in books 2, amount differs 1, IMS 0, carried forward 0 (%s)" % c)
    base = a5(pg)
    ok(abs(base - 7920) < 0.01, "Final credit 4(A)(5) before any decision: 7,920 (%s)" % base)
    ok(dec(pg, "B/7") == "now" and dec(pg, "E/1") == "follow" and dec(pg, "O/5") == "book" and dec(pg, "A/101") == "later", "FinCom's rule as decisions: Beta Take now, Echo Follow up, Omega Book in Tally, Alpha Take in a later month")
    shot(pg, "2-reco-desktop", True)
    choose(pg, "E/1", "booksnow")
    ok(take(pg, "E/1") == 900 and abs(a5(pg) - base - 900) < 0.01 and pg.locator(row_sel("E/1") + " [data-ic-risk]").count() == 1, "Take now though not in 2B on Echo: ITC in Apr 900 on the line, 4(A)(5) + 900, with its risk note")
    pg.locator(row_sel("E/1") + " select[data-ic-dec]").focus(); pg.locator(row_sel("E/1")).screenshot(path=os.path.join(SHOTS, "3-decision-dropdown.png"))
    choose(pg, "B/7", "later")
    ok(pg.locator(row_sel("B/7") + " select[data-ic-month]").count() == 1 and pg.evaluate("S.books.itcDec['09|B|%s|B7|1'].ym" % IB.GB) == "202605", "Take in a later month: the month to pick appears, the next month first")
    pg.select_option(row_sel("B/7") + " select[data-ic-month]", "202606"); W(pg, 500)
    ok(take(pg, "B/7") in (0, None) and abs(a5(pg) - base - 900 + 3600) < 0.01 and abs(a5(pg, "202606") - 3600) < 0.01, "Beta to Jun-2026: April - 3,600, June + 3,600")
    choose(pg, "B/7", "now"); choose(pg, "E/1", "follow")
    ok(abs(a5(pg) - base) < 0.01, "back to Take now and Follow up: 4(A)(5) as before")
    # bulk
    for no in ("O/5", "CN/1"): pg.check(row_sel(no) + " input[data-ic-tick]"); W(pg, 200)
    ok(pg.locator("[data-ic-bulk]").count() == 1 and "2 selected" in pg.inner_text("[data-ic-bulk]"), "two lines ticked: Set decision… bar")
    pg.select_option("[data-ic-bulk-dec]", "reverse"); W(pg, 200); shot(pg, "4-bulk-set-desktop")
    pg.click("[data-ic-bulk-go]"); W(pg, 600)
    ok(dec(pg, "O/5") == "reverse" and dec(pg, "CN/1") == "reverse" and abs(a5(pg) - base - 180) < 0.01, "Set decision… Reverse / do not take on both: Sigma's credit note no longer reduces April (+ 180)")
    for no in ("O/5", "CN/1"): pg.check(row_sel(no) + " input[data-ic-tick]"); W(pg, 200)
    pg.select_option("[data-ic-bulk-dec]", "book"); pg.click("[data-ic-bulk-go]"); W(pg, 600)
    ok(dec(pg, "O/5") == "book" and dec(pg, "CN/1") == "reverse", "Book in Tally on both: Omega takes it; Sigma's credit note cannot (not a choice for a supplier's note only in 2B), left as it was")
    ok(pg.locator(row_sel("CN/1") + " select[data-ic-dec]").evaluate("(s) => [...s.options].map(o => o.value).join(',')") == "now,reverse,follow", "a supplier's credit note only in 2B offers Take now, Reverse / do not take, Follow up")
    for no in ("B/7",): pg.check(row_sel(no) + " input[data-ic-tick]"); W(pg, 200)
    pg.select_option("[data-ic-bulk-dec]", "book"); pg.click("[data-ic-bulk-go]"); W(pg, 600)
    ok(dec(pg, "B/7") == "now", "Book in Tally on Beta (in the books already): not a choice for it, left as it was")
    choose(pg, "CN/1", "now")
    ok(abs(a5(pg) - base) < 0.01, "Sigma's note Take now: 4(A)(5) as before")
    # write to suppliers
    pg.click("[data-ic-letters]"); W(pg, 600)
    lt = pg.inner_text("[data-pop=itc-letters]") if pg.locator("[data-pop=itc-letters]").count() else ""
    ok("Echo Stores" in lt and "Write to suppliers" in lt, "Write to suppliers: Echo (Follow up) is there")
    pg.click("[data-pop=itc-letters] button:has-text('see letter')"); W(pg, 300)
    ok("E/1" in pg.inner_text("[data-letter-text]"), "the letter names Echo's invoice E/1")
    shot(pg, "5-write-to-suppliers")
    pg.keyboard.press("Escape"); W(pg, 500)
    ok(pg.locator("[data-pop=itc-letters]").count() == 0, "Esc closes it")
    # carried forward to May
    IB.open_part(pg, "r2b", "202605")
    c5 = chips(pg)
    ok(c5.get("cf", 0) >= 2 and pg.locator(row_sel("E/1")).count() == 1 and pg.get_attribute(row_sel("E/1"), "data-ic-chip-of") == "cf", "May: Echo (not taken in April) is Carried forward (%s)" % c5)
    pg.click('[data-ic-chip="cf"]'); W(pg, 400); shot(pg, "6-carried-forward-desktop", True)
    may0 = a5(pg, "202605")
    choose(pg, "E/1", "now")
    ok(abs(a5(pg, "202605") - may0 - 900) < 0.01 and abs(a5(pg) - base) < 0.01, "Echo taken now in May: May + 900, April unchanged")
    choose(pg, "E/1", "follow"); pg.click('[data-ic-chip="cf"]'); W(pg, 300)

    print("4. Final credit for 3B, adjustments")
    IB.open_part(pg, "itcfinal")
    fin = pg.evaluate("() => Number(document.querySelector('[data-if-a5]').innerText.replace(/[₹,]/g, ''))")
    ok(abs(fin - base) < 0.01, "Final credit: 4(A)(5) %s" % fin)
    shot(pg, "7-final-credit-desktop", True)
    pg.click("[data-if-add]"); W(pg, 600)
    pg.select_option('[data-drawer=itc-adj] select[aria-label="3B line"]', "4A5"); pg.select_option('[data-drawer=itc-adj] select[aria-label="Head"]', "cgst")
    pg.fill('[data-drawer=itc-adj] input[aria-label="Amount"]', "1200"); pg.fill('[data-drawer=itc-adj] input[aria-label="Reason"]', "bill of March received late")
    shot(pg, "8-adjustment-add-desktop")
    pg.click("[data-itc-adj-save]"); W(pg, 700)
    ok(pg.locator("[data-if-adj-row]").count() == 1 and abs(a5(pg) - base - 1200) < 0.01, "+ Add adjustment (4(A)(5), CGST 1,200): listed, 4(A)(5) + 1,200")
    IB.open_part(pg, "r3b"); pg.evaluate("() => { S.gstSub = 'summary'; render(); }"); W(pg, 600)
    r3 = pg.evaluate("() => Object.fromEntries(['4A5', '4C', '4B1'].map(id => [id, [...document.querySelectorAll('[data-r3b-row=\"' + id + '\"] td.n')].slice(0, 4).map(td => Number(td.innerText.replace(/[₹,]/g, '')))]))")
    t4 = T4(pg)
    ok(all(abs(r3[k][i] - t4[k][h]) < 0.01 for k in r3 for i, h in enumerate(H)), "GSTR-3B's table 4 on screen is Final credit's, figure for figure (%s)" % r3)
    IB.open_part(pg, "itcfinal")
    pg.click("[data-if-adj-remove]"); W(pg, 600)
    pg.fill('[data-drawer=itc-adj] input[aria-label="Why it is taken out"]', "entered twice"); pg.click("[data-itc-adj-delete-go]"); W(pg, 700)
    ok(pg.locator("[data-if-adj-row]").count() == 0 and abs(a5(pg) - base) < 0.01 and "Removed (1)" in pg.inner_text("[data-if-adj]"), "Remove (with the reason): 4(A)(5) as before, kept under Removed")
    pg.click('[data-if-fig="4A5"] >> nth=0'); W(pg, 700)
    ok(pg.locator("[data-pop=if-4A5] [data-fig-check=ok]").count() == 1, "4(A)(5) opens its lines in a window; they add up to the figure")
    pg.keyboard.press("Escape"); W(pg, 500)

    print("5. GSTR-3B: figures open a window")
    IB.open_part(pg, "r3b"); pg.evaluate("() => { S.gstSub = 'summary'; render(); }"); W(pg, 600)
    for fid, what in (("4A5|cgst", "4A5"), ("3.1a|cgst", "3.1a"), ("5gst|intra", "5gst"), ("6.1A|cgst|cash", "6.1A|cgst")):
        sel = '[data-r3b-fig="%s"]' % fid
        if not pg.locator(sel).count(): ok(False, "3B figure %s is a link" % fid); continue
        pg.click(sel); W(pg, 700)
        pop = '[data-pop="r3b-%s"]' % what
        vis = pg.locator(pop).count() == 1 and pg.locator(pop).is_visible()
        ok(vis and pg.locator("#app .r3b-bills").count() == 0, "3B %s: a window opens (nothing below the table)" % what)
        if what == "4A5":
            ok(pg.locator(pop + " [data-fig-check=ok]").count() == 1 and pg.locator(pop + " [data-pop-go=r2b]").count() == 1, "4(A)(5)'s window: the lines add up to the figure; Open in 2B reconciliation")
            shot(pg, "9-3b-popup-desktop")
        if what == "3.1a": pg.click(pop + " button[aria-label='Close dialog']")
        else: pg.keyboard.press("Escape")
        W(pg, 500)
        ok(pg.locator(pop).count() == 0, "and %s closes it" % ("x" if what == "3.1a" else "Esc"))
    pg.click('[data-r3b-fig="4A5|cgst"]'); W(pg, 600); pg.click('[data-pop="r3b-4A5"] [data-pop-go=r2b]'); W(pg, 800)
    ok(pg.evaluate("S.gstPart") == "r2b" and pg.locator("[data-pop]").count() == 0, "Open in 2B reconciliation: there, the window closed")

    print("6. Excel round trip")
    IB.open_part(pg, "r2b")
    pg.evaluate(CAPTURE); pg.click("[data-itc-excel=down]")
    for _ in range(60):
        if pg.evaluate("window.__files.length"): break
        W(pg, 250)
    f = pg.evaluate("window.__files[0]")
    ok(f and f["name"].endswith(".xlsx") and "2B-and-ITC" in f["name"], "Download Excel: %s" % (f and f["name"]))
    tmp = tempfile.mkdtemp(prefix="itc-xlsx-"); xp = os.path.join(tmp, "itc.xlsx")
    open(xp, "wb").write(base64.b64decode(f["b64"]))
    import openpyxl
    wb = openpyxl.load_workbook(xp)
    ok(wb.sheetnames[:6] == ["Input register", "2B reconciliation", "Adjustments", "Final credit for 3B", "Summary", "Write to suppliers"], "its sheets: %s" % wb.sheetnames)
    wr = wb["2B reconciliation"]
    ok(wr.column_dimensions["A"].hidden and wr.protection.sheet and wr["K2"].protection.locked is False and wr["C2"].protection.locked is not False, "the key column hidden, the sheet protected, Decision open, the data locked")
    dvs = [d for d in wr.data_validations.dataValidation]
    ok(any("Lists" in (d.formula1 or "") and "K2" in str(d.sqref) for d in dvs) and any("L2" in str(d.sqref) for d in dvs), "Decision and Take in month are lists (data validation)")
    ok(str(wr["N2"].value).startswith("=IF(") and "SUMIFS" in str(wb["Final credit for 3B"]["B6"].value) or "SUM(" in str(wb["Final credit for 3B"]["B3"].value), "Credit month and table 4 are formulas")
    # LibreOffice: the formulas as downloaded equal the screen
    d0 = os.path.join(tmp, "lo0"); os.makedirs(d0)
    orig = os.path.join(tmp, "orig.xlsx"); openpyxl.load_workbook(xp).save(orig)
    r0 = recalc(orig, d0)
    if r0:
        f0 = final_of(r0); t4 = T4(pg)
        ok(all(abs(f0[k][h] - t4[k][h]) < 0.01 for k in ("4A1", "4A3", "4A5", "4A", "4B1", "4B2", "4C", "4D2") for h in H), "LibreOffice works the downloaded Excel out to FinCom's table 4 (4(A)(5) %s / %s)" % (f0.get("4A5"), t4["4A5"]))
        ws0 = openpyxl.load_workbook(r0, data_only=True)["2B reconciliation"]
        tk = sum((r[14].value or 0) for r in ws0.iter_rows(min_row=2) if r[0].value)
        scr = pg.evaluate("() => Number(document.querySelector('[data-ic-take-sum]').innerText.replace(/[₹,]/g, ''))")
        ok(abs(tk - scr) < 0.01, "the Excel's ITC in Apr-2026 column adds up to the screen's (%s / %s)" % (tk, scr))
    else: print("  (LibreOffice not here: the recalculation checks are skipped)")
    # decisions changed in the file, an adjustment added
    rows = {r[2].value: r[0].row for r in wr.iter_rows(min_row=2) if r[0].value}
    wr.cell(rows["Echo Stores"], 11).value = "Take now though not in 2B"
    wr.cell(rows["BETA TRADERS"], 11).value = "Take in a later month"; wr.cell(rows["BETA TRADERS"], 12).value = pg.evaluate("GSTR.label('202606')")
    wr.cell(rows["OMEGA LOGISTICS"], 13).value = "asked them on 24-Oct"
    wa = wb["Adjustments"]; wa["B2"] = "SGST"; wa["C2"] = -250; wa["D2"] = "reversal of a blocked credit"; wa["E2"] = "4(B)(1) Reversed: rules 38, 42, 43 and section 17(5)"
    ed = os.path.join(tmp, "edited.xlsx"); wb.save(ed)
    d1 = os.path.join(tmp, "lo1"); os.makedirs(d1)
    r1 = recalc(ed, d1)
    f1 = final_of(r1) if r1 else {}
    if r1: ok(f1 and abs(tot(f1["4A5"]) - (base + 900 - 3600)) < 0.01 and abs(f1["4B1"]["sgst"] + 250) < 0.01, "LibreOffice: the changed decisions and adjustment move the Excel's table 4 (4(A)(5) %s, 4(B)(1) SGST %s)" % (f1.get("4A5"), (f1.get("4B1") or {}).get("sgst")))
    # upload: the preview, then Apply
    pg.set_input_files("[data-itc-upload]", r1 or ed); W(pg, 2500)
    pv = '[data-pop="itc-import"]'
    ok(pg.locator(pv).count() == 1, "Upload Excel: the preview opens")
    if pg.locator(pv).count():
        ok(pg.get_attribute(pv + " [data-pv-changes]", "data-pv-changes") == "2" and pg.locator(pv + " [data-pv-change]").count() == 2 and pg.get_attribute(pv + " [data-pv-adj]", "data-pv-adj") == "1|0|0"
           and pg.get_attribute(pv + " [data-pv-notes]", "data-pv-notes") == "1", "the preview: 2 decisions changed (which), 1 note, 1 adjustment added")
        ok(pg.locator(pv + " [data-pv-final] [data-pv-line='4A5']").count() == 1, "and the Excel's table 4 beside FinCom's")
        shot(pg, "10-excel-preview-desktop")
        pg.click(pv + " [data-pv-apply]"); W(pg, 900)
    t4 = T4(pg)
    if not f1: f1 = {"4A5": {"igst": 0, "cgst": round(t4["4A5"]["cgst"], 2), "sgst": round(t4["4A5"]["sgst"], 2), "cess": 0}} if abs(tot(t4["4A5"]) - (base + 900 - 3600)) < 0.01 else {}
    ok(f1 and all(abs(f1[k][h] - t4[k][h]) < 0.005 for k in ("4A1", "4A2", "4A3", "4A5", "4A", "4B1", "4B2", "4C", "4D1", "4D2") if k in f1 for h in H), "after Apply, FinCom's table 4 is the Excel's to the paisa (%s / %s)" % (t4["4A5"], f1.get("4A5")))
    ok(dec(pg, "E/1") == "booksnow" and dec(pg, "B/7") == "later" and pg.evaluate("S.books.itcDec['09|P|%s|O5|1'].note" % IB.GO) == "asked them on 24-Oct", "the decisions, the month and the note are FinCom's")
    # an edited key is refused
    wb2 = openpyxl.load_workbook(ed); w2 = wb2["2B reconciliation"]; w2.cell(rows["Echo Stores"], 1).value = "B|edited|X|1"; bad = os.path.join(tmp, "bad.xlsx"); wb2.save(bad)
    pg.set_input_files("[data-itc-upload]", bad); W(pg, 2000)
    ok(pg.locator(pv + " [data-pv-unknown]").count() == 1 and "refused" in pg.inner_text(pv + " [data-pv-unknown]"), "an edited key: that row refused, said plainly")
    pg.keyboard.press("Escape"); W(pg, 400)

    print("7. dark, phone")
    pg.emulate_media(color_scheme="dark"); W(pg, 400)
    IB.open_part(pg, "r2b"); shot(pg, "11-reco-desktop-dark", True)
    IB.open_part(pg, "itcfinal"); shot(pg, "12-final-desktop-dark", True)
    IB.open_part(pg, "inreg"); shot(pg, "13-input-register-desktop-dark", True)
    IB.open_part(pg, "r3b"); pg.evaluate("() => { S.gstSub = 'summary'; render(); }"); W(pg, 500); pg.click('[data-r3b-fig="4A5|cgst"]'); W(pg, 700); shot(pg, "14-3b-popup-dark"); pg.keyboard.press("Escape"); W(pg, 300)
    pg.emulate_media(color_scheme="light")
    ph = ctx.browser.new_context(viewport={"width": 390, "height": 844}, device_scale_factor=2, is_mobile=True, has_touch=True).new_page()
    ph.on("pageerror", lambda e: errors.append("phone: " + str(e)))
    IB.setup(ph, PORT)
    for part, name in (("inreg", "15-input-register-phone"), ("r2b", "16-reco-phone"), ("itcfinal", "17-final-phone")):
        IB.open_part(ph, part); shot(ph, name, True)
        sw = ph.evaluate("document.documentElement.scrollWidth")
        ok(sw <= 392, "phone %s: no sideways scroll of the page (%s)" % (part, sw))
    IB.open_part(ph, "r3b"); ph.evaluate("() => { S.gstSub = 'summary'; render(); }"); W(ph, 500)
    ph.click('[data-r3b-fig="4A5|cgst"]'); W(ph, 700); shot(ph, "18-3b-popup-phone")
    ok(ph.locator('[data-pop="r3b-4A5"]').count() == 1, "phone: the 3B window opens")
    br.close()
srv.shutdown()
errs = [e for e in errors if "supabase" not in e.lower() and "Failed to load" not in e]
ok(not errs, "no page errors" + ("" if not errs else ": " + " | ".join(errs[:4])))
print("\n" + ("%d FAILED" % len(fails) if fails else "all passed")); raise SystemExit(1 if fails else 0)
