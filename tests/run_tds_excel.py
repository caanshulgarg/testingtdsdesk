"""python3 run_tds_excel.py - a TDS / TCS return as one Excel, down and back up, as a person uses it (next-tds-excel,
24-Oct-2026; the owner: "prepare this type of download and upload return .. in all types of TDS return.. user can add or
delete or modify the details"; "add challan, add deduction, add deductee within this smart excel.. it should be dynamic";
"challan balance should also be there"). On the React build (app/dist-test or TDSDESK_SITE) with the fixture books a year
on (tdsgst_snapshot.LOAD, CHALLANS):
  1. 26Q Q4 2025-26: the return's ⋯ has Download Excel and Upload Excel (nothing else on the page changed); the workbook:
     Deductor, Challans, Deductions, Deductees, Summary, Checks (Excel Tables, hidden keys, dropdowns, formulas)
  2. in openpyxl, as a person in Excel: a deduction's rate to nil with remark T, another's section and rate, one deleted;
     a new challan, a new deductee and a new deduction on both (only in the Excel); the deductee's PAN; the deductor's phone.
     LibreOffice works the formulas out (skipped where it is not installed): the new deduction's PAN from the new deductee,
     its TDS = amount × rate, the new challan in the pickers and its balance nil, Summary's totals
  3. the same file with one more deduction pushing the new challan over its balance (and a bad PAN): Upload: the preview
     names the errors, Apply is refused
  4. the good file: Upload: the preview names exactly those changes; Apply: the return's rows, the screen's total, the FVU
     text's lines and the Excel downloaded again equal the file; the deleted row is still in the books; Undo puts it back
  5. 24Q Q1 2026-27 (Form 138, the salary sheet): a TDS changed, a row deleted, one added: applied, the Form 138 text
  6. 27Q and 27EQ (smoke): a deduction and a collection added in the Excel, applied
Screenshots (1440 and 390) to TDSDESK_SHOTS (default tests/out/tdsx-shots)."""
import os, sys, json, threading, functools, http.server, base64, shutil, subprocess, tempfile, datetime
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from books_data import CACHE
from tdsgst_snapshot import LOAD, CHALLANS
from playwright.sync_api import sync_playwright
import openpyxl
from openpyxl.formula.translate import Translator
SITE = os.environ.get("TDSDESK_SITE", os.path.join(HERE, "..", "app", "dist-test"))
PORT = int(os.environ.get("TDSX_PORT") or 8463)
SHOTS = os.environ.get("TDSDESK_SHOTS") or os.path.join(os.environ.get("TDSDESK_OUT", os.path.join(HERE, "out")), "tdsx-shots")
WORK = tempfile.mkdtemp(prefix="tdsx-")
os.makedirs(SHOTS, exist_ok=True)
fails, errors = [], []
def ok(c, w):
    print(("  ok   " if c else "  FAIL ") + w)
    if not c: fails.append(w)
class Quiet(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(("localhost", PORT), functools.partial(Quiet, directory=SITE)); threading.Thread(target=srv.serve_forever, daemon=True).start()
books = json.load(open(CACHE))
SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")
CATCH = """() => { window.__files = []; window.saveFile = async (name, data) => { const a = new Uint8Array(await data.arrayBuffer()); let s = ''; for (let i = 0; i < a.length; i += 8192) s += String.fromCharCode.apply(null, a.subarray(i, i + 8192)); window.__files.push({name, b64: btoa(s)}); return true; }; }"""
DEDUCTOR = """() => { const co = CO(); co.tds26q = Object.assign({}, co.tds26q, {flat: '12 First Floor', road: 'MG Road', area: 'Karol Bagh', town: 'New Delhi', state: '07', pin: '110005', email: 'accounts@zz.example', phone: '9811111111',
  person: 'Ravi Kumar', personDesignation: 'Partner', personPan: 'ABCPK1234K', personFlat: '12 First Floor', personState: '07', personPin: '110005', personEmail: 'ravi@zz.example', personPhone: '9822222222', branch: 'Head office'}); co.pan = 'AAGCL4827M'; }"""

def recalc(path):
    """LibreOffice works every formula out (the file saved by openpyxl has formulas, no values)"""
    if not SOFFICE: return None
    out = os.path.join(WORK, "lo"); os.makedirs(out, exist_ok=True)
    prof = "file://" + tempfile.mkdtemp(prefix="tdsx-lo-", dir=WORK)
    subprocess.run([SOFFICE, "-env:UserInstallation=" + prof, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", out, path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return os.path.join(out, os.path.basename(path))
def lo_html(path, name):
    """LibreOffice's own rendering of the workbook (every sheet), for the screenshot of the Excel"""
    if not SOFFICE: return None
    out = os.path.join(WORK, "html"); os.makedirs(out, exist_ok=True)
    prof = "file://" + tempfile.mkdtemp(prefix="tdsx-lo-", dir=WORK)
    subprocess.run([SOFFICE, "-env:UserInstallation=" + prof, "--headless", "--convert-to", "html:HTML (StarCalc)", "--outdir", out, path], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return os.path.join(out, os.path.splitext(os.path.basename(path))[0] + ".html")
def last_row(ws, col="D", start=3):
    """the last row with a value in a column (a table's rows), and the first empty one under it"""
    r = start
    while ws[col + str(r)].value not in (None, ""): r += 1
    return r
def del_row(ws, r, upto=2400):
    """a row deleted as Excel deletes it: those under it move up, their formulas follow (openpyxl does not move them)"""
    for rr in range(r + 1, upto):
        for c in ws[rr]:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = Translator(c.value, origin=c.coordinate).translate_formula(c.column_letter + str(rr - 1))
    ws.delete_rows(r)

with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(viewport={"width": 1440, "height": 950}); pg = ctx.new_page()
    pg.on("pageerror", lambda e: errors.append(str(e)))
    E, W = pg.evaluate, pg.wait_for_timeout
    pg.goto("http://localhost:%d/" % PORT); W(2500); pg.click('button[data-act="useOffline"]'); W(2000)
    E("window.__dueToday = '2026-10-24'")
    E(LOAD, books); W(1000); E("S.books = window.__bk; render();"); W(500); E(CHALLANS); W(300); E(DEDUCTOR)
    cid = E("S.coId")
    def go(h, wait=1500):
        E("(h) => { location.hash = h; }", "#/c/%s/books/tds/%s" % (cid, h)); W(wait)
    def menu():
        pg.click("#app [data-return-card] details.more > summary"); W(250)
        return E("[...document.querySelectorAll('#app [data-return-card] .more-item')].map(b => b.textContent)")
    def download(path):
        E(CATCH); menu(); pg.click("#app [data-return-card] [data-tdsx=down]"); W(2500)
        f = E("window.__files[window.__files.length - 1] || null")
        if not f: return None
        open(path, "wb").write(base64.b64decode(f["b64"])); return f["name"]
    def upload(path):
        pg.set_input_files("#app [data-tdsx-upload]", files=[{"name": os.path.basename(path), "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": open(path, "rb").read()}]); W(3500)
        return E("""() => { const pv = (S.tdsX || {}).pv; if (!pv) return {error: (S.tdsX || {}).error || 'no preview'}; const d = pv.deductions, c = pv.challans, e = pv.deductees;
          return {fatal: pv.fatal, n: pv.n, errors: pv.errors.map(x => x.sheet + ': ' + x.words), warnings: pv.warnings.map(x => x.words), ded: [d.add.length, d.edit.length, d.out.length], ch: [c.add.length, c.edit.length, c.alloc.length],
            dee: [e.add.length, e.edit.length], dor: pv.deductor.map(x => x.k), edits: d.edit.map(x => [x.r.party, x.fields.map(f => f.k).sort()]), excel: pv.excel, after: pv.after, before: pv.before,
            bal: pv.balances}; }""")
    def rows(form, fy, q): return E("([f, y, q]) => TDS.formRows(f).filter(r => r.fy === y && r.q === q).map(r => ({id: r.id, party: r.party || r.employee, pan: r.pan, section: r.section, code: r.code, paid: r.paid, tds: r.tds, challan: r.challan}))", [form, fy, q])
    def shot(name, full=False): pg.screenshot(path=os.path.join(SHOTS, name + ".png"), full_page=full)

    print("1. 26Q Q4 2025-26: Download Excel")
    FY, Q = "2025-26", "Q4"
    go("%s/%s/26Q/deductions" % (FY, Q))
    items = menu()
    ok("Download Excel" in items and "Upload Excel" in items and not any(i.startswith("Excel to correct") or i.startswith("Import corrected") for i in items), "the return's ⋯: Download Excel and Upload Excel, in place of the old two (%s)" % items)
    shot("1-return-menu-desktop")
    pg.keyboard.press("Escape"); W(200)
    pg.set_viewport_size({"width": 390, "height": 844}); W(500); menu(); shot("1-return-menu-phone"); pg.keyboard.press("Escape"); pg.set_viewport_size({"width": 1440, "height": 950}); W(500)
    R0 = rows("26Q", FY, Q)
    f0 = os.path.join(WORK, "26Q-down.xlsx"); name = download(f0)
    ok(name and name.endswith(".xlsx"), "downloaded: %s" % name)
    wb = openpyxl.load_workbook(f0)
    ok(wb.sheetnames[:6] == ["Deductor", "Challans", "Deductions", "Deductees", "Summary", "Checks"] and wb["Lists"].sheet_state == "hidden", "the sheets: %s" % wb.sheetnames)
    wd, wc, we = wb["Deductions"], wb["Challans"], wb["Deductees"]
    ok(wd.column_dimensions["A"].hidden and wd["A3"].value == R0[0]["id"] and wd.tables and wc.tables and we.tables, "Deductions, Challans and Deductees are Excel Tables; the hidden Key first (%s)" % list(wd.tables.keys()))
    dv = {str(v.sqref).split()[0]: v.formula1 for v in wd.data_validations.dataValidation}
    ok(any("OFFSET('Challans'" in (f or "") for f in dv.values()) and any("OFFSET('Deductees'" in (f or "") for f in dv.values()) and any("Lists" in (f or "") for f in dv.values()),
       "the dropdowns: challan and deductee pickers read the Challans and Deductees sheets (OFFSET), section and remark the lists")
    ok(str(wd["J3"].value).startswith("=IF(AND(I3") and "ROUND(N(I3)*N(H3)/100,2)" in str(wd["J3"].value) and "INDEX('Deductees'" in str(wd["E3"].value) and "INDEX('Challans'" in str(wd["Q3"].value), "TDS = amount × rate, the PAN from Deductees and the challan's balance from Challans, by formulas")
    ok(str(wc["Q3"].value).startswith("=IF(V3") and "SUMIF('Deductions'" in str(wc["O3"].value) and any(cf.sqref and "Q3" in str(cf.sqref) for cf in wc.conditional_formatting), "Challans: used in this return (SUMIF), Balance (red below nil)")
    ok(str(wd["J%d" % (len(R0) + 5)].value or "").startswith("=IF(AND("), "the formulas are on the spare rows under the table too")

    print("2. changed in openpyxl as a person would, LibreOffice works it out")
    nch = last_row(wc, "B") ; nde = last_row(we, "B")
    r_n, r_j, r_h = 3, 4, 5      # Nightjar 194C, Juniper 194J, Hemant 194A
    ok([wd["D%d" % r].value for r in (r_n, r_j, r_h)] == ["Nightjar Sound & Light Co", "Juniper Legal Associates", "Hemant Zaverchand (Loan)"], "the rows as the return has them")
    wd["H%d" % r_n] = 0; wd["M%d" % r_n] = "T"                         # rate nil, remark T (a transporter's declaration)
    wd["F%d" % r_j] = "194H"; wd["H%d" % r_j] = 2                      # section and rate
    del_row(wd, r_h)                                                    # deleted: out of the return
    for r in range(3, 30):
        if (we["B%d" % r].value or "") == "Nightjar Sound & Light Co": we["D%d" % r] = "AAAFN1234C"
    wc["B%d" % nch] = "0510308"; wc["C%d" % nch] = datetime.datetime(2026, 4, 7); wc["D%d" % nch] = "00451"; wc["E%d" % nch] = "94C"; wc["F%d" % nch] = "200"; wc["G%d" % nch] = 1500; wc["T%d" % nch] = "online"
    we["B%d" % nde] = "Zenith Advisors LLP"; we["D%d" % nde] = "AAAFZ1234Q"; we["F%d" % nde] = "Non-company"
    nd = last_row(wd, "A")
    wd["B%d" % nd] = datetime.datetime(2026, 3, 20); wd["C%d" % nd] = datetime.datetime(2026, 3, 18); wd["D%d" % nd] = "Zenith Advisors LLP"; wd["F%d" % nd] = "194C"; wd["H%d" % nd] = 1; wd["I%d" % nd] = 150000
    wd["P%d" % nd] = "0510308/07-04-2026/00451"
    dor = wb["Deductor"]
    for r in range(3, 40):
        if dor["A%d" % r].value == "phone": dor["C%d" % r] = "9899999999"
    good = os.path.join(WORK, "26Q-good.xlsx"); wb.save(good)
    # the bad one: the same, and one more deduction pushing the new challan over its balance, a bad PAN on a new deductee
    wd["B%d" % (nd + 1)] = datetime.datetime(2026, 3, 21); wd["D%d" % (nd + 1)] = "Zenith Advisors LLP"; wd["F%d" % (nd + 1)] = "194C"; wd["H%d" % (nd + 1)] = 1; wd["I%d" % (nd + 1)] = 300000; wd["P%d" % (nd + 1)] = "0510308/07-04-2026/00451"
    we["B%d" % (nde + 1)] = "Bad Pan Co"; we["D%d" % (nde + 1)] = "ABC12"
    bad = os.path.join(WORK, "26Q-bad.xlsx"); wb.save(bad)
    want_tds = round(0 + 600 + 1500, 2)
    if SOFFICE:
        good2, bad2 = recalc(good), recalc(bad)
        v = openpyxl.load_workbook(good2, data_only=True)
        sm = {r[0].value: r[2].value for r in v["Summary"].iter_rows(min_row=3, max_row=20)}
        dd = v["Deductions"]
        ok(dd["E%d" % nd].value == "AAAFZ1234Q" and abs((dd["J%d" % nd].value or 0) - 1500) < 0.005 and abs((dd["Q%d" % nd].value or 0)) < 0.005, "LibreOffice: the new deduction's PAN from the new deductee, TDS 1,500 = 150,000 × 1%%, its challan's balance nil (%s %s %s)" % (dd["E%d" % nd].value, dd["J%d" % nd].value, dd["Q%d" % nd].value))
        ok(abs((sm.get("tds") or 0) - want_tds) < 0.005 and abs((sm.get("onch") or 0) - want_tds) < 0.005 and sm.get("over") == 0 and sm.get("n") == 3, "Summary: 3 deductions, TDS %s, all on challans, no challan over-used (%s)" % (want_tds, {k: sm.get(k) for k in ("n", "tds", "onch", "gap", "over", "fee")}))
        ck = v["Checks"]
        ok(ck["E1"].value == "0510308/07-04-2026/00451" and ck["F1"].value == "Bad Pan Co" or ck["F1"].value == "Zenith Advisors LLP", "the pickers list what was just added: the new challan and the new deductee (%s, %s)" % (ck["E1"].value, ck["F1"].value))
        vb = openpyxl.load_workbook(bad2, data_only=True)
        ok(str(vb["Checks"]["C1"].value).startswith("Challan over-used: 1") and (vb["Challans"]["Q%d" % nch].value or 0) < -0.005 and vb["Summary"]["C10"].value is not None, "the bad file: Checks says a challan is over-used; its Balance below nil (%s, %s)" % (vb["Checks"]["C1"].value, vb["Challans"]["Q%d" % nch].value))
        good, bad = good2, bad2
        h = lo_html(good, "26Q")
        if h:
            pv = br.new_page(viewport={"width": 1440, "height": 900}); pv.goto("file://" + h); pv.wait_for_timeout(800)
            pv.evaluate("() => { const hs = [...document.querySelectorAll('h1, a[name]')]; const t = hs.find(e => /Deductions/.test(e.textContent || e.getAttribute('name') || '')); if (t) t.scrollIntoView(); }")
            pv.screenshot(path=os.path.join(SHOTS, "2-excel-deductions-libreoffice.png"));
            pv.evaluate("() => { const hs = [...document.querySelectorAll('h1, a[name]')]; const t = hs.find(e => /Challans/.test(e.textContent || e.getAttribute('name') || '')); if (t) t.scrollIntoView(); }")
            pv.screenshot(path=os.path.join(SHOTS, "2-excel-challans-libreoffice.png")); pv.close()
    else: print("  (LibreOffice is not installed: the formulas are not worked out; the upload works the TDS out as Amount × Rate)")

    print("3. the bad file: refused")
    x = upload(bad)
    ok(any("over-used" in e for e in x["errors"]) and any("ABC12" in e for e in x["errors"]), "the preview: the challan over-used and the bad PAN are errors (%s)" % x["errors"][:4])
    ok(pg.locator("#app [data-tdsx-apply]").count() == 1 and pg.locator("#app [data-tdsx-apply]").is_disabled(), "Apply is refused while errors are left")
    shot("3-upload-preview-error-desktop")
    pg.set_viewport_size({"width": 390, "height": 844}); W(600); shot("3-upload-preview-error-phone"); pg.set_viewport_size({"width": 1440, "height": 950}); W(400)
    ok(rows("26Q", FY, Q) == R0, "nothing applied")
    pg.click("#app [data-drawer=tdsx] button:has-text('Cancel')"); W(400)

    print("4. the good file: preview, Apply, Undo")
    x = upload(good)
    ok(not x["errors"] and not x["fatal"], "no error (%s)" % x["errors"][:3])
    ok(x["ded"] == [1, 2, 1] and x["ch"][0] == 1 and x["dee"] == [1, 1] and x["dor"] == ["phone"], "the preview: 1 deduction added, 2 changed, 1 taken out; a challan added; a deductee added and one changed; the deductor's phone (%s)" % {k: x[k] for k in ("ded", "ch", "dee", "dor")})
    ed = dict((p_, f) for p_, f in x["edits"])
    ok(ed.get("Nightjar Sound & Light Co") == ["basis", "rate", "tds"] and ed.get("Juniper Legal Associates") == ["rate", "section", "tds"], "exactly the fields changed: Nightjar rate, TDS, remark; Juniper section, rate, TDS (%s)" % ed)
    if SOFFICE: ok(x["excel"] and abs(x["excel"]["tds"] - x["after"]["tds"]) < 0.005 and abs(x["excel"]["tds"] - want_tds) < 0.005, "the Excel's Summary and FinCom's after Apply agree: TDS %s" % (x["excel"] and x["excel"].get("tds")))
    if x["warnings"]: pg.check("#app [data-tdsx-tick]"); W(200)
    shot("4-upload-preview-desktop")
    pg.set_viewport_size({"width": 390, "height": 844}); W(600); shot("4-upload-preview-phone"); pg.set_viewport_size({"width": 1440, "height": 950}); W(400)
    pg.click("#app [data-tdsx-apply]"); W(1500)
    R1 = rows("26Q", FY, Q)
    z = next((r for r in R1 if r["party"] == "Zenith Advisors LLP"), None)
    newch = E("TDS.challans().find(c => c.serial === '00451') || null")
    ok(len(R1) == 3 and z and z["pan"] == "AAAFZ1234Q" and abs(z["tds"] - 1500) < 0.005 and newch and z["challan"] == newch["id"], "Apply: the new deduction on the new challan with the new deductee's PAN (%s)" % z)
    ok(abs(sum(r["tds"] for r in R1) - want_tds) < 0.005 and next(r for r in R1 if r["party"].startswith("Juniper"))["section"] == "194H" and next(r for r in R1 if r["party"].startswith("Nightjar"))["pan"] == "AAAFN1234C", "the return: TDS %s, Juniper 194H, Nightjar's PAN" % want_tds)
    ok(E("TDS26Q.remark(TDS.rows().find(r => r.party === 'Nightjar Sound & Light Co' && r.fy === '%s' && r.q === '%s'), '26Q').remark" % (FY, Q)) == "T" and E("CO().tds26q.phone") == "9899999999", "FinCom writes remark T for Nightjar; the deductor's phone kept")
    ok(E("S.books.vouchers.some(v => '%s'.startsWith(v.id + '|'))" % R0[2]["id"]) and not any(r["id"] == R0[2]["id"] for r in R1), "the deleted row: out of the return, still in the books")
    t = E("([fy, q]) => { const r = TDS26Q.build(fy, q, TDS26Q.firmDetails(), '26Q'); return r.error ? {error: r.error} : {dd: r.text.split('\\r\\n').filter(l => l.split('^')[1] === 'DD')}; }", [FY, Q])
    ok(not t.get("error") and len(t["dd"]) == 3 and any("AAAFZ1234Q" in l and "1500.00" in l for l in t["dd"]) and any("600.00" in l for l in t["dd"]), "the FVU text: three DD lines, Zenith's 1,500 with its PAN, Juniper's 600 (%s)" % (t.get("error") or len(t["dd"])))
    E("() => { S.tdsDedView = 'table'; render(); }"); W(600)
    foot = pg.inner_text("#app [data-ded-count]") if pg.locator("#app [data-ded-count]").count() else ""
    ok(E("INR.format(%s)" % want_tds) in foot, "the screen: the Deductions table's foot says TDS %s (%s)" % (want_tds, foot[:80]))
    shot("5-after-apply-desktop")
    pg.set_viewport_size({"width": 390, "height": 844}); W(600); shot("5-after-apply-phone"); pg.set_viewport_size({"width": 1440, "height": 950}); W(400)
    f2 = os.path.join(WORK, "26Q-again.xlsx"); download(f2)
    v2 = openpyxl.load_workbook(f2, data_only=True)["Deductions"]
    again = sorted((v2["D%d" % r].value, v2["F%d" % r].value, round(v2["I%d" % r].value or 0, 2), round(v2["J%d" % r].value or 0, 2), v2["P%d" % r].value or "") for r in range(3, 3 + len(R1)))
    filed = sorted((r["party"], r["section"], round(r["paid"], 2), round(r["tds"], 2)) for r in R1)
    ok([a[:4] for a in again] == filed and all(a[4] for a in again), "downloaded again: the same rows as FinCom's and the Excel's, each on its challan (%s)" % again)
    x2 = upload(f2)
    ok(x2.get("n") == 0 and not x2.get("errors"), "and uploaded again: nothing to change (%s)" % {k: x2.get(k) for k in ("n", "errors")})
    pg.click("#app [data-drawer=tdsx] button:has-text('Close')"); W(300)
    menu(); pg.click("#app [data-return-card] [data-tdsx=undo]"); W(1200)
    ok(sorted((r["id"], r["tds"], r["section"]) for r in rows("26Q", FY, Q)) == sorted((r["id"], r["tds"], r["section"]) for r in R0), "Undo: the return as it was (the deleted row back, the added one gone)")

    print("5. 24Q Q1 2026-27 (Form 138, the salary sheet)")
    FY, Q = "2026-27", "Q1"
    E("() => { const rs = TDS.formRows('24Q').filter(r => r.fy === '2026-27' && r.q === 'Q1'); S.books.challans.push({id: 'ch24', bsr: '0240020', serial: '20001', date: '20260707', tax: 60000, section: '192', minorHead: '200', from: 'upload'}); TDSReg.forget(); TDSCH.tag(rs.map(r => r.id), 'ch24', {nosave: true}); TDSReg.forget(); }")
    go("%s/%s/24Q/deductions" % (FY, Q))
    S0 = rows("24Q", FY, Q)
    f24 = os.path.join(WORK, "24Q-down.xlsx"); download(f24)
    wb = openpyxl.load_workbook(f24); wd = wb["Deductions"]
    ra = next(r for r in range(3, 12) if wd["D%d" % r].value == "Asha Verma"); rv = next(r for r in range(3, 12) if wd["D%d" % r].value == "Ravi Menon")
    wd["J%d" % ra] = 7000
    rv_tds = next(r["tds"] for r in S0 if r["id"] == wd["A%d" % rv].value)
    del_row(wd, rv)
    n24 = last_row(wd, "A")
    wd["B%d" % n24] = datetime.datetime(2026, 6, 30); wd["D%d" % n24] = "Asha Verma"; wd["F%d" % n24] = "192"; wd["I%d" % n24] = 10000; wd["J%d" % n24] = 1000; wd["P%d" % n24] = "0240020/07-07-2026/20001"
    g24 = os.path.join(WORK, "24Q-good.xlsx"); wb.save(g24)
    if SOFFICE: g24 = recalc(g24)
    x = upload(g24)
    ok(not x["errors"] and x["ded"] == [1, 1, 1], "24Q preview: one TDS changed, one row taken out, one added (%s %s)" % (x["ded"], x["errors"][:2]))
    if x["warnings"]: pg.check("#app [data-tdsx-tick]"); W(200)
    pg.click("#app [data-tdsx-apply]"); W(1500)
    S1 = rows("24Q", FY, Q); want24 = round(sum(r["tds"] for r in S0) - 6405 + 7000 - rv_tds + 1000, 2)
    ok(len(S1) == len(S0) and abs(sum(r["tds"] for r in S1) - want24) < 0.005, "24Q applied: %d rows, TDS %s" % (len(S1), want24))
    t = E("() => { const r = TDS24Q_FILE.build('2026-27', 'Q1', TDS26Q.firmDetails()); return r.error ? {error: r.error} : {dd: r.text.split('\\r\\n').filter(l => l.split('^')[1] === 'DD')}; }")
    ok(not t.get("error") and len(t["dd"]) == len(S1) and any("7000.00" in l for l in t["dd"]) and any("1000.00" in l for l in t["dd"]), "the Form 138 text: a DD line a row, the changed and the added (%s)" % (t.get("error") or len(t["dd"])))
    menu(); pg.click("#app [data-return-card] [data-tdsx=undo]"); W(1000)
    ok(round(sum(r["tds"] for r in rows("24Q", FY, Q)), 2) == round(sum(r["tds"] for r in S0), 2), "Undo: 24Q as it was")

    print("6. 27Q and 27EQ (smoke)")
    E("""() => { const b = S.books;
      b.vouchers.push({id: 'nr1', date: '20260512', type: 'Journal', no: 'NR1', party: 'Globex GmbH', ent: [{l: 'Globex GmbH', a: 80000}, {l: 'TDS 195', a: 20000}, {l: 'Fees Abroad', a: -100000}]},
        {id: 'tc1', date: '20260515', type: 'Sales', no: 'S9', party: 'Scrapco', ent: [{l: 'Scrapco', a: -101000}, {l: 'Scrap Sales', a: 100000}, {l: 'TCS Scrap', a: 1000}]});
      b.map['TDS 195'] = {n: 1, kind: 'tds_payable', what: 'tds_payable', section: '195', ok: true, byHand: true}; b.map['TCS Scrap'] = {n: 1, kind: 'tcs_payable', what: 'tcs_payable', section: '206C', ok: true, byHand: true};
      b.challans.push({id: 'chN', bsr: '0240020', serial: '30001', date: '20260607', tax: 30000, section: '195', minorHead: '200', from: 'upload'}, {id: 'chT', bsr: '0240020', serial: '30002', date: '20260607', tax: 3000, section: '206CE', minorHead: '200', from: 'upload'});
      TDSReg.forget(); render(); }""")
    for form, sec, party, lab, paid, rate in [("27Q", "195", "Globex GmbH", "0240020/07-06-2026/30001", 50000, 10), ("27EQ", "6CE", "Scrapco", "0240020/07-06-2026/30002", 50000, 1)]:
        go("2026-27/Q1/%s/deductions" % form)
        before = rows(form, "2026-27", "Q1")
        fx = os.path.join(WORK, form + "-down.xlsx"); download(fx)
        wb = openpyxl.load_workbook(fx); wd = wb["Deductions"]; n = last_row(wd, "A")
        wd["P3"] = lab
        if form == "27EQ": wd["F3"] = sec
        wd["B%d" % n] = datetime.datetime(2026, 5, 20); wd["D%d" % n] = party; wd["F%d" % n] = sec; wd["H%d" % n] = rate; wd["I%d" % n] = paid; wd["P%d" % n] = lab
        gx = os.path.join(WORK, form + "-good.xlsx"); wb.save(gx)
        if SOFFICE: gx = recalc(gx)
        x = upload(gx)
        ok(not x["errors"] and x["ded"][0] == 1, "%s preview: one added, its challan picked (%s %s)" % (form, x["ded"], x["errors"][:2]))
        if x["warnings"]: pg.check("#app [data-tdsx-tick]"); W(200)
        pg.click("#app [data-tdsx-apply]"); W(1200)
        after = rows(form, "2026-27", "Q1")
        ok(len(after) == len(before) + 1 and all(r["challan"] for r in after) and abs(sum(r["tds"] for r in after) - sum(r["tds"] for r in before) - paid * rate / 100) < 0.005, "%s applied: %d rows, each on its challan" % (form, len(after)))
    br.close()

shutil.rmtree(WORK, ignore_errors=True)
ok(not errors, "no page errors (%s)" % errors[:3])
srv.shutdown()
print("\n%d failed" % len(fails) if fails else "\nall passed")
sys.exit(1 if fails else 0)
