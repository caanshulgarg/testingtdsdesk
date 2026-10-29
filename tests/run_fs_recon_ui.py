"""python3 run_fs_recon_ui.py - matching FinCom's balance sheet with the audited one, on the Financial statements page's
Checks tab (owner, 29-Oct-2026). On the React build (app/dist-test, or TDSDESK_SITE), the made-up books of
tests/fixtures/fin-statements and the made-up audited balance sheet as at 31-Mar-2026 (audited-bs-2026.json, as an Excel):
  - the audited balance sheet is uploaded and confirmed on the Checks tab; the table shows audited · FinCom · difference
    head by head, a head opens to its note lines, a FinCom figure to its ledgers in the trial balance;
  - the proposals: Willow Labs mapped to loans to related parties; the provision for tax and the depreciation as
    adjustments dated 31-Mar-2026; the profit's difference and what is left (1,000) said; reject and take back;
  - Accept all mapping changes / Accept all adjustments: the view recomputes, two heads of 1,000 left; each change in
    the history with who and when, undone and accepted again;
  - "Download journal entries for Tally": an Excel with the two entries;
  - 2026-27: "Opening balances adjusted to the audited balance sheet (2 entries)", the comparative as audited;
  - a phone (390 wide): no sideways scroll.
Screenshots (1440 light and dark, and a phone) go to TDSDESK_SHOTS (default tests/out/recon-shots)."""
import os, sys, json, subprocess, threading, functools, http.server, shutil
os.environ.setdefault("PLAYWRIGHT_BROWSERS_PATH", "/opt/pw-browsers")
from playwright.sync_api import sync_playwright
import openpyxl
HERE = os.path.dirname(os.path.abspath(__file__))
SITE = os.environ.get("TDSDESK_SITE", os.path.join(HERE, "..", "app", "dist-test"))
OUT = os.path.join(os.environ.get("TDSDESK_OUT", os.path.join(HERE, "out")), "fs-recon"); shutil.rmtree(OUT, ignore_errors=True); os.makedirs(OUT)
SHOTS = os.environ.get("TDSDESK_SHOTS", os.path.join(HERE, "out", "recon-shots")); os.makedirs(SHOTS, exist_ok=True)
fails, errors = [], []
def ok(c, w):
    print(("  ok   " if c else "  FAIL ") + w)
    if not c: fails.append(w)
class Q(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a): pass
srv = http.server.ThreadingHTTPServer(("localhost", 0), functools.partial(Q, directory=SITE)); PORT = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()
AUD = json.load(open(os.path.join(HERE, "fixtures", "fin-statements", "audited-bs-2026.json")))
BOOKS = json.loads(subprocess.check_output(["node", "-e", r"""
const F = require('./fixtures/fin-statements/company.json'), from = '20240401', r2 = v => Math.round(v * 100) / 100;
const under = Object.assign({}, F.groups), led = {};
Object.entries(F.o).forEach(([n, v]) => { led[n] = {open: -v, close: 0, parent: under[n]}; });
const vouchers = [], bank = 'Northbank Current A/c', shift = e => String(Number(e.slice(0, 4)) - 1) + e.slice(4);
F.additions.forEach(([l, d, a], i) => vouchers.push({id: 'a' + i, date: d, no: 'A' + i, type: 'Journal', narr: 'Purchase of ' + l, ent: [{l, a: -a}, {l: bank, a}]}));
[['20250331', F.p, F.plp, F.o], ['20260331', F.c, F.plc, F.p]].forEach(([end, bs, pl, prev]) => { const ent = [];
  Object.keys(bs).forEach(n => { if (n === bank) return; const added = F.additions.filter(([l, d]) => l === n && d <= end && d > shift(end)).reduce((s, z) => s + z[2], 0);
    const mv = r2((bs[n] || 0) - (prev[n] || 0) - added); if (mv) ent.push({l: n, a: -mv}); });
  Object.entries(pl).forEach(([n, v]) => { if (v) ent.push({l: n, a: -v}); });
  ent.push({l: bank, a: r2(-ent.reduce((s, e) => s + e.a, 0))}); vouchers.push({id: 'm' + end, date: end, no: 'M' + end, type: 'Journal', narr: 'Movements (fixture)', ent}); });
vouchers.push({id: 'n1', date: '20260630', no: 'S1', type: 'Sales', narr: 'Services, April-June 2026', ent: [{l: 'Software Services - Export', a: 500000}, {l: bank, a: -500000}]});
process.stdout.write(JSON.stringify({under, groups: {}, vouchers, meta: {from, to: '20270331'}, tb: {from, to: '20270331', at: '2027-04-30T10:00:00Z', src: 'copy', led}}));
"""], cwd=HERE))
SETUP = """(bk) => {
  S.firm = Object.assign({}, S.firm, {firmName: "Made-up & Associates", frn: "000000S", partner: "A. Sample", mno: "000001", place: "Bengaluru"});
  const c = newCompany({name: "Larkspur Analytics Private Limited (made-up)", gstin: "29AABCL1234F1Z5"}); c.pan = "AABCL1234F"; c.cin = "U72900KA2015PTC000001";
  S.companies[c.id] = c; Store.saveCompany(c); S.coId = c.id; S.view = "company"; S.tab = "books"; S.loadingCo = false;
  const b = Object.assign({loading: false, cid: c.id, challans: [], alloc: {}, ledInfo: {}, gstins: {}, pans: {}, states: {}, misCfg: {freq: "off"}, auditCfg: {freq: "off"},
    msme: {"Bluebell Print Works": "Micro"}, fs: {taxDep: {"2024": {c40: 220000, pm15: 140000, f10: 80000}}, inputs: {unit: 1, dec: 2}}}, bk);
  b.map = Books.mapLedgers(b.vouchers, {});
  S.books = b; window.__bk = b; S.fsFy = "2025"; S.fin = null; saveBooks(); goClient('books:fs'); return c.id; }"""
OPEN = "() => new Promise((r) => { const go = () => (typeof doAct === 'function' && typeof signInNeeded === 'function' && typeof S === 'object' && S && S.firm) ? (() => { try { if (signInNeeded()) doAct('useOffline'); } finally { r(); } })() : setTimeout(go, 50); go(); })"

xl = os.path.join(OUT, "audited-bs-2025-26.xlsx")
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Balance Sheet"
for r in [[AUD["title"]], [AUD["what"]], [AUD["unit"]], AUD["head"]] + AUD["rows"]: ws.append(r)
wb.save(xl)

with sync_playwright() as p:
    br = p.chromium.launch()
    for scheme in ("light", "dark"):
        ctx = br.new_context(viewport={"width": 1440, "height": 900}, color_scheme=scheme, accept_downloads=True); pg = ctx.new_page()
        pg.on("pageerror", lambda e: errors.append(str(e)))
        pg.goto("http://localhost:%d/" % PORT); pg.wait_for_timeout(2500); pg.evaluate(OPEN); pg.wait_for_timeout(1000)
        pg.evaluate(SETUP, BOOKS); pg.wait_for_timeout(2500)
        tab = lambda k: (pg.click('#app nav.fin-tabs [data-fin-tab="%s"]' % k), pg.wait_for_timeout(700))
        tab("checks")
        R = "#app [data-fin-recon]"
        if scheme == "dark":
            pg.set_input_files(R + " input[data-fin-recon-file]", xl); pg.wait_for_timeout(2500); pg.click(R + " [data-fin-recon-save]"); pg.wait_for_timeout(1500)
            pg.evaluate("() => { document.querySelector('#app [data-fin-recon]').scrollIntoView(); }"); pg.wait_for_timeout(400)
            pg.screenshot(path=os.path.join(SHOTS, "recon-desktop-dark.png")); ctx.close(); continue
        print("Upload the audited balance sheet as at 31-Mar-2026")
        ok(pg.locator(R + " h4").inner_text().strip() == "Match with the audited balance sheet as at 31-Mar-2026", "the Checks tab: " + pg.locator(R + " h4").inner_text())
        pg.set_input_files(R + " input[data-fin-recon-file]", xl); pg.wait_for_timeout(2500)
        ok(pg.locator(R + " [data-fin-recon-confirm]").count() == 1, "read and shown to confirm")
        pg.click(R + " [data-fin-recon-save]"); pg.wait_for_timeout(1500)
        ok(pg.evaluate("S.books.fs.pyBs['2026'].versions.length") == 1, "kept (as 2026-27's last year's balance sheet too)")
        head = lambda i: pg.locator(R + ' tr[data-fin-rhead="%s"] td' % i)
        diff = lambda i: head(i).nth(3).inner_text().strip()
        ok("6 heads differ" in pg.inner_text(R + " [data-fin-recon-state]"), "6 heads differ")
        ok(diff("tr") == "(90,000.00)" and diff("stla") == "90,000.00" and diff("ocl") == "6,83,940.00" and diff("ppe") == "(25,000.00)" and diff("tpo") == "1,000.00" and diff("rs") == "(7,09,940.00)" and diff("cash") == "0",
           "audited · FinCom · difference by head: %s" % [diff(i) for i in ("tr", "stla", "ocl", "ppe", "tpo", "rs")])
        pg.click(R + ' tr[data-fin-rhead="ocl"] button'); pg.wait_for_timeout(500)
        ok(pg.locator(R + ' tr[data-fin-rcode="7D"] td').nth(3).inner_text().strip() == "6,83,940.00", "a head opens to its note lines: provision for income tax 6,83,940.00")
        pg.click(R + ' tr[data-fin-rhead="tr"] button'); pg.wait_for_timeout(500)
        pg.click(R + ' [data-fin-rled="10A"]'); pg.wait_for_timeout(900)
        keys = pg.locator('#app table[data-fin-table="tb"] tbody tr[data-key]').evaluate_all("xs => xs.map(x => x.dataset.key)")
        ok("Willow Labs" in keys and "Maple Holdings Inc" in keys, "a FinCom figure opens its ledgers in the trial balance: %s" % keys)
        tab("checks")
        print("\nThe proposals")
        props = pg.locator(R + " [data-fin-prop]").evaluate_all("xs => xs.map(x => x.dataset.finProp)")
        ok(props == ["map:Willow Labs>12H", "adj:tax:7D", "adj:dep:8A"], "a mapping change, then two adjustments: %s" % props)
        t = pg.inner_text(R + ' [data-fin-prop="adj:tax:7D"]')
        ok("Dr Income Tax - Current Year / Cr Provision for Income Tax FY 2025-26 · 6,83,940.00" in t and "To match audited balance sheet: Other current liabilities" in t, "the tax provision: " + t.replace("\n", " "))
        ok("by amount" in pg.inner_text(R + ' [data-fin-prop="map:Willow Labs>12H"]'), "the mapping change with its reason")
        pr = pg.inner_text(R + " [data-fin-rprofit]")
        ok("(7,09,940.00)" in pr and "explain (7,08,940.00)" in pr and "(1,000.00) is not explained" in pr, "the profit: " + pr)
        left = pg.locator(R + " [data-fin-left]").evaluate_all("xs => xs.map(x => x.dataset.finLeft)")
        ok(sorted(left) == ["rs", "tpo"], "left once accepted, never forced: %s" % left)
        pg.screenshot(path=os.path.join(SHOTS, "recon-desktop-light-before.png"), full_page=True)
        pg.click(R + ' [data-fin-reject="map:Willow Labs>12H"]'); pg.wait_for_timeout(900)
        ok(pg.locator(R + ' [data-fin-prop="map:Willow Labs>12H"]').count() == 0 and "1 rejected" in pg.inner_text(R + " [data-fin-rrejected]"), "rejected: gone from the list")
        pg.click(R + " [data-fin-rlog] summary"); pg.wait_for_timeout(300)
        pg.click(R + " [data-fin-rlog] [data-fin-undo]"); pg.wait_for_timeout(900)
        ok(pg.locator(R + ' [data-fin-prop="map:Willow Labs>12H"]').count() == 1, "taken back from the history")
        print("\nAccept all")
        pg.click(R + ' [data-fin-acceptall="map"]'); pg.wait_for_timeout(1200)
        pg.click(R + ' [data-fin-acceptall="adj"]'); pg.wait_for_timeout(1500)
        ok("2 heads differ" in pg.inner_text(R + " [data-fin-recon-state]") and diff("tpo") == "1,000.00" and diff("rs") == "(1,000.00)" and diff("tr") == "0" and diff("ocl") == "0" and diff("ppe") == "0",
           "recomputed: the unexplained 1,000 left (%s)" % pg.inner_text(R + " [data-fin-recon-state]"))
        ok(pg.locator(R + " [data-fin-prop]").count() == 0 and sorted(pg.locator(R + " [data-fin-left]").evaluate_all("xs => xs.map(x => x.dataset.finLeft)")) == ["rs", "tpo"], "no proposal left; what is left listed")
        if not pg.locator(R + " [data-fin-rlog][open]").count(): pg.click(R + " [data-fin-rlog] summary"); pg.wait_for_timeout(300)
        hist = pg.inner_text(R + " [data-fin-rlog]")
        ok(hist.count("Accepted:") == 3 and "A. Sample" in hist, "each accepted change in the history, with who and when")
        dep = pg.locator(R + ' [data-fin-rlog] li:has-text("Accumulated Depreciation - Computers") [data-fin-undo]')
        dep.click(); pg.wait_for_timeout(1200)
        ok(diff("ppe") == "(25,000.00)" and pg.locator(R + ' [data-fin-prop="adj:dep:8A"]').count() == 1, "undone: the depreciation entry gone, proposed again")
        pg.click(R + ' [data-fin-accept="adj:dep:8A"]'); pg.wait_for_timeout(1200)
        ok(diff("ppe") == "0", "accepted again")
        tab("tb")
        ok(pg.locator('#app [data-fin-adjrow]').count() == 2 and "matches the audited balance sheet" in pg.locator("#app [data-fin-adj]").text_content(), "the entries in FinCom's adjustments list, marked")
        tab("checks")
        with pg.expect_download(timeout=60000) as dl:
            pg.click(R + " [data-fin-journal]")
        jp = os.path.join(OUT, "journal.xlsx"); dl.value.save_as(jp)
        jw = openpyxl.load_workbook(jp).active; jr = [[c.value for c in r] for r in jw.iter_rows(min_row=3) if r[0].value]
        ok(len(jr) == 2 and all(r[0] == "31-Mar-2026" and r[1] == "Journal" for r in jr) and ["Income Tax - Current Year", "Provision for Income Tax FY 2025-26", 683940] == jr[0][2:5] and jr[1][2:5] == ["Depreciation", "Accumulated Depreciation - Computers", 25000],
           "Download journal entries for Tally: %s (%s)" % (dl.value.suggested_filename, [r[2:5] for r in jr]))
        pg.evaluate("() => { document.querySelector('#app [data-fin-recon]').scrollIntoView(); }"); pg.wait_for_timeout(400)
        pg.screenshot(path=os.path.join(SHOTS, "recon-desktop-light-after.png"))
        print("\nNext year's opening")
        pg.evaluate("() => { S.fsFy = '2026'; FinSt.periodSet({preset: 'h1', cmp: 'fy'}); render(); }"); pg.wait_for_timeout(1500); tab("bs")
        ok("Opening balances adjusted to the audited balance sheet (2 entries)" in pg.inner_text("#app [data-fin-carried]"), "2026-27: " + pg.inner_text("#app [data-fin-carried]"))
        ok(pg.locator('#app table[data-fin-table="bs"] tr[data-fin-total="II"] td').nth(3).inner_text().strip() == "85,62,800.55", "the comparative as at 31-Mar-2026: 85,62,800.55, the audited total")
        pg.screenshot(path=os.path.join(SHOTS, "recon-next-year-light.png"))
        ctx.close()
    ph = br.new_context(viewport={"width": 390, "height": 844}, device_scale_factor=2, is_mobile=True, has_touch=True); pp = ph.new_page()
    pp.on("pageerror", lambda e: errors.append(str(e)))
    pp.goto("http://localhost:%d/" % PORT); pp.wait_for_timeout(2500); pp.evaluate(OPEN); pp.wait_for_timeout(1000)
    pp.evaluate(SETUP, BOOKS); pp.wait_for_timeout(2500)
    pp.evaluate("(r) => { const v = FinPyBs.prepare(FinPyBs.fromRows([{name: 'Balance Sheet', rows: r}]), 'co'); FinPyBs.save('2026', v); S.fin.tab = 'checks'; render(); }", [[AUD["title"]], [AUD["what"]], [AUD["unit"]], AUD["head"]] + AUD["rows"]); pp.wait_for_timeout(1500)
    w = pp.evaluate("document.documentElement.scrollWidth - document.documentElement.clientWidth")
    ok(w <= 1, "a phone: no sideways scroll (%d)" % w)
    pp.evaluate("() => { document.querySelector('#app [data-fin-recon]').scrollIntoView(); }"); pp.wait_for_timeout(400)
    pp.screenshot(path=os.path.join(SHOTS, "recon-phone.png"))
    ph.close(); br.close()
errs = [e for e in errors if "supabase" not in e.lower() and "Failed to load" not in e]
ok(not errs, "no page errors" + ("" if not errs else ": " + " | ".join(errs[:4])))
print("\n" + ("%d FAILED" % len(fails) if fails else "all passed")); sys.exit(1 if fails else 0)
